from __future__ import annotations

import io
import sqlite3
import zipfile
from pathlib import Path
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from erp.core.database import SessionLocal
from erp.core.jsonutil import dumps
from erp.core.errors import AppError
from erp.models import Asset, CatalogProduct, CatalogProductAsset, PublicationStore
from erp.serializers import publication_store_summary_dict
from erp.services.jobs import execute_job, get_job
from erp.services.catalog import _pod_supplier_from_sku
from erp.integrations.testing import PublicationTestHarness
from erp.services.publications import apply_ozon_merge_plan, build_ozon_merge_plan, normalize_ozon_dictionary_attributes, product_snapshot
from erp.modules.ozon_follow.legacy.services.ozon_listing import OzonListingDraft
from erp.services.ozon_categories import save_category_snapshot
import pytest


def test_ozon_image_review_list_returns_exact_candidate_version() -> None:
    review = {'revision': 'a' * 64, 'sha256': 'b' * 64, 'url': 'https://oss.test/candidate.jpg',
              'candidate_hash': 'c' * 64, 'human_approval': 'PENDING'}
    child = SimpleNamespace(id='child', store_id='store', store=SimpleNamespace(name='ZYH', enabled=True),
        offer_id='ZYH-S2B-001876-01', status='image_review', content_json='{}',
        assets_json=dumps({'personal_review': review}), pricing_json='{}', package_json='{}',
        external_task_id='', external_product_id='', stock_status='pending', promotion_status='pending',
        last_error='', updated_at=datetime.now(timezone.utc))
    assert publication_store_summary_dict(child)['assets']['personal_review'] == review


def test_ozon_merge_plan_prefers_real_color_then_builds_unique_multi_value() -> None:
    definitions=[{"id":10096,"name":"颜色","dictionary_id":1,"is_aspect":True,"is_collection":True,"max_value_count":2}]
    values={10096:[{"id":1,"value":"黑色"},{"id":2,"value":"白色"},{"id":3,"value":"红色"}]}
    facts=[{"sku":"A","color":"黑色"},{"sku":"B","color":"黑色"},{"sku":"C","color":"红色"}]
    plan=build_ozon_merge_plan(facts,definitions,values,model_name="同组型号",variant_attribute_ids=[10096])
    assert plan["A"]["attributes"][0]["values"]==[{"dictionary_value_id":1,"value":"黑色"}]
    assert plan["B"]["attributes"][0]["values"]==[{"dictionary_value_id":1,"value":"黑色"},{"dictionary_value_id":2,"value":"白色"}]
    assert plan["B"]["attributes"][0]["source"]=="ozon_unique_fallback"
    assert plan["C"]["attributes"][0]["values"]==[{"dictionary_value_id":3,"value":"红色"}]


def test_ozon_merge_plan_only_virtualizes_color_when_size_is_also_selected() -> None:
    definitions=[
        {"id":200,"name":"尺寸，厘米","dictionary_id":2,"is_aspect":True},
        {"id":10096,"name":"商品颜色","dictionary_id":1,"is_aspect":True,"is_collection":True,"max_value_count":2},
    ]
    values={
        200:[{"id":21,"value":"80x80"},{"id":22,"value":"100x120"},{"id":23,"value":"100x1000"}],
        10096:[{"id":1,"value":"绿色"},{"id":2,"value":"白色"},{"id":3,"value":"粉色"}],
    }
    facts=[{"sku":"A","size":"80x80","color":"绿色"},{"sku":"B","size":"80x80","color":"绿色"},{"sku":"C","size":"未知","color":""}]
    plan=build_ozon_merge_plan(facts,definitions,values,model_name="同组型号",variant_attribute_ids=[200,10096])
    assert plan["A"]["attributes"][0]["values"]==[{"dictionary_value_id":21,"value":"80x80"}]
    assert plan["B"]["attributes"][0]["values"]==[{"dictionary_value_id":21,"value":"80x80"}]
    assert plan["B"]["attributes"][1]["source"]=="ozon_unique_fallback"
    assert plan["C"]["attributes"][0]["values"]==[]
    assert plan["C"]["attributes"][0]["source"]=="missing_real_value"
    assert plan["C"]["attributes"][1]["source"]=="ozon_fallback"


def test_ozon_merge_plan_uses_first_variant_as_fallback_when_color_is_unavailable() -> None:
    definitions=[{"id":200,"name":"尺寸，厘米","dictionary_id":2,"is_aspect":True}]
    values={200:[{"id":21,"value":"80x80"},{"id":22,"value":"100x120"}]}
    facts=[{"sku":"A","size":""},{"sku":"B","size":""}]
    plan=build_ozon_merge_plan(facts,definitions,values,model_name="同组型号",variant_attribute_ids=[200])
    assert plan["A"]["attributes"][0]["values"]==[{"dictionary_value_id":21,"value":"80x80"}]
    assert plan["B"]["attributes"][0]["values"]==[{"dictionary_value_id":22,"value":"100x120"}]


def test_apply_ozon_merge_plan_overrides_ai_values() -> None:
    definitions=[{"id":9000,"name":"型号名称（针对合并为一张商品卡片）"},{"id":10096,"name":"颜色","dictionary_id":1}]
    draft=OzonListingDraft(source_sku="A",offer_id="A",source_ozon_url="",title="A",exported_category="A",attributes=[{"id":10096,"values":[{"dictionary_value_id":99,"value":"旧值"}]}])
    apply_ozon_merge_plan(draft,definitions,{"mode":"custom_group","model_name":"固定型号","attributes":[{"attribute_id":10096,"values":[{"dictionary_value_id":1,"value":"黑色"}]}]})
    assert next(item for item in draft.attributes if item["id"]==9000)["values"]==[{"value":"固定型号"}]
    assert next(item for item in draft.attributes if item["id"]==10096)["values"]==[{"dictionary_value_id":1,"value":"黑色"}]


def test_apply_ozon_merge_plan_overrides_russian_model_name_default() -> None:
    definitions=[{"id":9048,"name":"Название модели (для объединения в одну карточку)"}]
    draft=OzonListingDraft(source_sku="A",offer_id="STORE-A",source_ozon_url="",title="A",exported_category="A",attributes=[{"id":9048,"values":[{"value":"STORE-A"}]}])
    apply_ozon_merge_plan(draft,definitions,{"mode":"custom_group","model_name":"USER-GROUP","attributes":[]})
    assert draft.attributes[0]["values"]==[{"value":"USER-GROUP"}]


def test_dictionary_attributes_are_normalized_and_invalid_values_removed() -> None:
    class DictionaryClient:
        def category_attribute_values(self,*_args,**_kwargs):
            return [{"id":7,"value":"Хлопок"}]
    definitions=[{"id":10,"name":"Материал","dictionary_id":5}]
    draft=OzonListingDraft(source_sku="A",offer_id="A",source_ozon_url="",title="A",exported_category="A",description_category_id=1,type_id=2,attributes=[{"id":10,"values":[{"value":"Хлопок"}]},{"id":99,"values":[{"value":"unchanged"}]}])
    removed=normalize_ozon_dictionary_attributes(draft,definitions,DictionaryClient())
    assert removed==[]
    assert draft.attributes[0]["values"]==[{"dictionary_value_id":7,"value":"Хлопок"}]
    draft.attributes[0]["values"]=[{"dictionary_value_id":999,"value":"not allowed"}]
    assert normalize_ozon_dictionary_attributes(draft,definitions,DictionaryClient())==["Материал"]
    assert draft.attributes[0]["values"]==[]


def test_ozon_merge_plan_rejects_normal_product_attributes() -> None:
    with pytest.raises(AppError,match="普通商品属性"):
        build_ozon_merge_plan([{"sku":"A","color":"棉"}],[{"id":7,"name":"材质","dictionary_id":1,"is_aspect":False}],{7:[{"id":1,"value":"棉"}]},model_name="同组型号",variant_attribute_ids=[7])


def _run(job_id: str) -> str:
    with SessionLocal() as session:
        execute_job(session, get_job(session, job_id))
        return get_job(session, job_id).status


def _store(client, name: str, external: str, prefix: str = "") -> str:
    return client.post("/api/v1/stores", json={"platform":"ozon","name":name,"external_account_id":external,"offer_prefix":prefix,"secrets":{"api_key":"test"}}).json()["id"]


def _import_category(client) -> dict[str, str]:
    store_id = _store(client, "导入类目店", "category-import")
    save_category_snapshot(store_id, "导入类目店", "ZH_HANS", [{
        "description_category_id": 10, "category_name": "箱包",
        "children": [{"description_category_id": 10, "category_name": "收纳包", "type_id": 20, "type_name": "化妆包"}],
    }])
    return {"store_id": store_id, "description_category_id": "10", "type_id": "20"}


def _catalog_product(client, source_type: str = "self_selected", source_url: str = "https://detail.1688.com/item/1") -> str:
    product=client.post("/api/v1/catalog/products",json={"sku":f"SKU-{source_type}","group_code":f"GROUP-{source_type}","source_type":source_type,"name":"测试产品","source_url":source_url,"supplier":"1688","purchase_price_cny":20,"package_length_mm":300,"package_width_mm":200,"package_height_mm":100,"package_weight_g":500,"description_category_id":10,"type_id":20,"category_display":"家居 / 测试","variants":[{"sku":f"SKU-{source_type}-M","color":"黑色","size":"M"}]}).json()
    with SessionLocal() as session:
        main=Asset(kind="image",name="main.jpg",source="test",oss_url="https://oss.test/main.jpg",sha256=f"{source_type:0<64}"[:64],mime_type="image/jpeg")
        extra=Asset(kind="image",name="extra.jpg",source="test",oss_url="https://oss.test/extra.jpg",mime_type="image/jpeg")
        session.add_all([main,extra]);session.flush()
        session.add_all([CatalogProductAsset(product_id=product["id"],asset_id=main.id,role="main",position=0),CatalogProductAsset(product_id=product["id"],asset_id=extra.id,role="additional",position=0)])
        session.commit()
    return product["id"]


def test_catalog_product_keeps_exact_facts_and_current_gallery(client):
    product_id=_catalog_product(client)
    detail=client.get(f"/api/v1/catalog/products/{product_id}").json()
    assert detail["source_type"] == "self_selected"
    assert detail["package_weight_g"] == 500
    assert [item["role"] for item in detail["assets"]] == ["main","additional"]
    assert detail["variants"][0]["sku"] == "SKU-self_selected-M"
    listing=client.get("/api/v1/catalog/products",params={"source_type":"self_selected","category":"家居 / 测试","image_status":"with_images"}).json()
    assert listing["total"] == 1
    assert listing["items"][0]["thumbnail"]["asset"]["oss_url"] == "https://oss.test/main.jpg"
    assert listing["category_options"] == [{"label":"家居 / 测试","count":1}]


def test_publication_preview_explains_manual_price_and_checks_materials(client):
    product_id=_catalog_product(client)
    store_id=_store(client,"预览店","preview-store","PV-")
    response=client.post("/api/v1/publications/preview",json={
        "product_ids":[product_id],"store_ids":[store_id],"requirements":{"price":999,"old_price":1299},
    })
    assert response.status_code == 200
    row=response.json()["items"][0]
    assert row["pricing"]["source"] == "manual_requirement"
    assert row["pricing"]["price"] == "999"
    assert row["offer_id"].startswith("PV-")
    assert row["image_count"] == 2
    assert row["missing"] == []


def test_follow_publication_uses_virtual_dimensions_without_changing_catalog():
    product=CatalogProduct(source_type="follow",sku="FOLLOW-1",name="跟卖商品",package_length_mm=None,package_width_mm=None,package_height_mm=None,package_weight_g=280)
    facts=product_snapshot(product,None)
    assert (facts["package_length_mm"],facts["package_width_mm"],facts["package_height_mm"]) == (300.0,200.0,100.0)
    assert facts["package_dimensions_source"] == "virtual_follow_30x20x10cm"
    assert product.package_length_mm is None


def test_duplicates_are_only_suggested_across_sources(client):
    first=_catalog_product(client,"follow","https://detail.1688.com/item/shared")
    second=_catalog_product(client,"pod","https://detail.1688.com/item/shared")
    suggestions=client.get("/api/v1/catalog/duplicates").json()["items"]
    assert suggestions
    assert {suggestions[0]["product_id"],suggestions[0]["candidate_product_id"]} == {first,second}
    assert "货源链接相同" in suggestions[0]["reasons"]
    assert client.get("/api/v1/catalog/summary").json()["total_products"] == 2
    follow_summary=client.get("/api/v1/catalog/summary?source_type=follow").json()
    pod_summary=client.get("/api/v1/catalog/summary?source_type=pod").json()
    assert follow_summary["total_products"] == 1
    assert follow_summary["variants"] == 1
    assert follow_summary["current_assets"] == 2
    assert pod_summary["total_products"] == 1


def test_self_selected_import_validates_then_imports_without_ai(client, tmp_path: Path):
    workbook=Workbook();sheet=workbook.active
    sheet.append(["商品信息",None,None,None,None,None,None,None,None,None,None,"包装尺寸",None,None,None,"生成内容",None,None])
    sheet.append(["来源链接","中文名称","SKU","规格型号","颜色","尺码","材质","拿货价","长","宽","高","长(cm)","宽(cm)","高（cm)","重量(g)","主图链接","附图链接","商品名称","备注"])
    sheet.append(["https://detail.1688.com/item/1","USB 收纳包","G-1-BLACK","黑色 M","黑色","M","涤纶",12.5,18,9,3,20,10,4,300,"https://oss.test/g1-main.jpg","https://oss.test/g1-1.jpg\nhttps://oss.test/g1-2.jpg","Чехол для хранения USB","仅内部核对：黑色款"])
    stream=io.BytesIO();workbook.save(stream)
    response=client.post("/api/v1/catalog/imports/self-selected",files={"file":("products.xlsx",stream.getvalue(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},data={"dry_run":"false",**_import_category(client)})
    assert response.status_code == 202
    assert _run(response.json()["job"]["id"]) == "completed"
    listing=client.get("/api/v1/catalog/products?source_type=self_selected").json()
    assert listing["total"] == 1
    detail=client.get(f"/api/v1/catalog/products/{listing['items'][0]['id']}").json()
    assert len(detail["assets"]) == 3
    assert detail["variants"][0]["sku"] == "G-1-BLACK"
    assert detail["package_length_mm"] == 200
    assert detail["description_category_id"] == 10
    assert detail["type_id"] == 20
    assert detail["category_display"] == "箱包 / 收纳包 / 化妆包"
    assert detail["production"]["copywriting"]["title_ru"] == "Чехол для хранения USB"
    assert detail["notes"] == "仅内部核对：黑色款"


def test_self_selected_download_returns_complete_standard_workbook(client):
    response=client.get("/api/v1/catalog/imports/self-selected/template")
    assert response.status_code == 200
    workbook=load_workbook(io.BytesIO(response.content),read_only=True,data_only=False)
    assert workbook.sheetnames == ["产品SKU信息","Sheet1","Sheet2","MP Tree Tarifs CN","Sheet3"]
    assert workbook["产品SKU信息"]["A1"].value == "商品信息"
    assert workbook["产品SKU信息"]["A2"].value == "编号(fsku)"
    assert "备注" in [cell.value for cell in workbook["产品SKU信息"][2]]
    workbook.close()


def test_self_selected_bundle_syncs_local_images_by_stable_group_code(client):
    category = _import_category(client)
    workbook=Workbook();sheet=workbook.active;sheet.title="产品导入"
    sheet.append(["SKU","来源链接","中文名称","拿货价","长(cm)","宽(cm)","高（cm)","重量(g)","主图链接","附图链接","产品名称","规格型号","颜色","尺码","材质","长","宽","高"])
    sheet.append(["LOCAL-1-M","https://detail.1688.com/item/local","本地图片商品",18,20,12,5,320,"https://oss.test/local-main.jpg","https://oss.test/local-extra.jpg","Локальный товар","","","","",None,None,None])
    manifest=workbook.create_sheet("图片清单");manifest.append(["产品组编号","SKU（可留空）","图片相对路径","图片类型","排序"]);manifest.append(["LOCAL-1-M","","images/LOCAL-1/main.jpg","main",0])
    excel=io.BytesIO();workbook.save(excel)
    bundle=io.BytesIO()
    with zipfile.ZipFile(bundle,"w") as archive:
        archive.writestr("products.xlsx",excel.getvalue());archive.writestr("images/LOCAL-1/main.jpg",b"local-image-v1")
    response=client.post("/api/v1/catalog/imports/self-selected/bundle",files={"file":("products.zip",bundle.getvalue(),"application/zip")},data={"dry_run":"false",**category})
    assert response.status_code == 202
    assert _run(response.json()["job"]["id"]) == "completed"
    listing=client.get("/api/v1/catalog/products?source_type=self_selected").json()
    assert listing["total"] == 1
    detail=client.get(f"/api/v1/catalog/products/{listing['items'][0]['id']}").json()
    assert detail["assets"][0]["role"] == "main"
    local_path=next(item["asset"]["local_path"] for item in detail["assets"] if item["asset"]["local_path"])
    assert local_path.endswith("images\\LOCAL-1\\main.jpg") or local_path.endswith("images/LOCAL-1/main.jpg")

    sheet["C2"]="本地图片商品（更新）";excel=io.BytesIO();workbook.save(excel);bundle=io.BytesIO()
    with zipfile.ZipFile(bundle,"w") as archive:
        archive.writestr("products.xlsx",excel.getvalue());archive.writestr("images/LOCAL-1/main.jpg",b"local-image-v2")
    response=client.post("/api/v1/catalog/imports/self-selected/bundle",files={"file":("products-v2.zip",bundle.getvalue(),"application/zip")},data={"dry_run":"false",**category})
    assert _run(response.json()["job"]["id"]) == "completed"
    listing=client.get("/api/v1/catalog/products?source_type=self_selected").json()
    assert listing["total"] == 1
    assert listing["items"][0]["name"] == "本地图片商品（更新）"


def test_pod_import_dry_run_reads_legacy_database_without_writing_products(client, tmp_path: Path):
    path=tmp_path/"product_library.db";connection=sqlite3.connect(path)
    connection.executescript("""
    create table products(id text,sku text,family_no integer,variant_no integer,name text,status text,tags_json text,notes text,current_version_id text,legacy_import integer,created_at text,updated_at text,category_id text);
    create table product_versions(id text,product_id text,version_no integer,job_id text,job_status text,product_type_id integer,default_color_id integer,default_view_id integer,product_code text,product_code_status text,output_folder text,recipe_json text,assets_json text,created_at text,updated_at text);
    create table product_variants(id text,product_id text,sku text,variant_no integer,color_id integer,color_name text,size_id integer,size_name text,stock_item_id integer,stock_code text,package_length real,package_width real,package_height real,package_volume real,weight real,size_attributes_json text,prices_json text,image_group_id text,selected integer,active integer,created_at text,updated_at text);
    create table product_oss_assets(id text,version_id text,asset_id text,asset_type text,color_id integer,color_name text,view_id integer,local_path text,oss_bucket text,oss_key text,oss_uri text,oss_url text,content_type text,size integer,sha256 text,etag text,crc64 text,oss_version_id text,uploaded_at text,updated_at text);
    create table product_categories(id text,name text,sort_order integer,is_system integer,created_at text);
    """)
    connection.execute("insert into products values(?,?,?,?,?,?,?,?,?,?,?,?,?)",("p1","POD-1",1,1,"杯子","ready","[]","","v1",0,"now","now","c1"))
    connection.execute("insert into product_versions values(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",("v1","p1",1,"j1","done",99,1,1,"PC-1","ready","D:/output","{}","{}","now","now"))
    connection.execute("insert into product_categories values(?,?,?,?,?)",("c1","杯子",1,0,"now"));connection.commit();connection.close()
    response=client.post("/api/v1/catalog/imports/pod",json={"dry_run":True,"source_database":str(path)}).json()
    assert _run(response["job"]["id"]) == "completed"
    assert client.get("/api/v1/catalog/products").json()["total"] == 0
    record=client.get("/api/v1/catalog/imports").json()["items"][0]
    assert record["counts"]["products"] == 1


def test_pod_supplier_is_derived_from_sku_prefix():
    assert _pod_supplier_from_sku("S2B-001234-01") == "S2B"
    assert _pod_supplier_from_sku("HC-000123-01") == "HICUSTOM"
    assert _pod_supplier_from_sku("POD-OTHER") == ""


def test_publication_isolated_per_store_and_excel_reuses_same_package(client):
    product_id=_catalog_product(client)
    first=_store(client,"店铺 A","501","A-");second=_store(client,"店铺 B","502","B-")

    def fake_prepare(context, publication, child, snapshot):
        title=f"俄文-{child.store.name}"
        item={"offer_id":child.offer_id,"name":title,"description_category_id":10,"type_id":20,"price":"999","old_price":"1299","currency_code":"RUB","depth":300,"width":200,"height":100,"weight":500,"dimension_unit":"mm","weight_unit":"g","primary_image":f"https://oss.test/{child.store_id}/main.jpg","images":["https://oss.test/additional.jpg"],"color_image":f"https://oss.test/{child.store_id}/main.jpg","attributes":[],"complex_attributes":[],"barcode":"","vat":"0","pdf_list":[]}
        child.content_json=dumps({"title":title,"description_ru":f"描述-{child.store.name}","tags":["#test"]});child.assets_json=dumps({"primary_image":item["primary_image"],"images":item["images"],"package_ready":True,"import_item":item});context.session.commit()
        return {"item":item,"definitions":[]}

    response=client.post("/api/v1/publications",json={"product_ids":[product_id],"store_ids":[first,second],"requirements":{},"image_policy":"reuse","submit_via_api":False}).json()
    publication=response["items"][0]
    with patch("erp.services.publications.prepare_store_package",fake_prepare): assert _run(publication["job"]["id"]) == "completed"
    detail=client.get(f"/api/v1/publications/{publication['publication']['id']}").json()
    assert {row["content"]["title"] for row in detail["stores"]} == {"俄文-店铺 A","俄文-店铺 B"}
    assert {row["offer_id"] for row in detail["stores"]} == {"A-SKU-self_selected-M","B-SKU-self_selected-M"}
    exported=client.get(f"/api/v1/publications/{detail['id']}/export/ozon.xlsx")
    assert exported.status_code == 200
    workbook=load_workbook(io.BytesIO(exported.content),read_only=True,data_only=True);rows=list(workbook["OZON发布包"].iter_rows(values_only=True));workbook.close()
    assert {row[2] for row in rows[1:]} == {"俄文-店铺 A","俄文-店铺 B"}


def test_generated_images_are_not_written_back_unless_selected(client):
    product_id=_catalog_product(client);store=_store(client,"图库店","601")
    original=client.get(f"/api/v1/catalog/products/{product_id}").json()["current_asset_count"]

    def fake_prepare(context, publication, child, snapshot):
        item={"offer_id":child.offer_id,"name":"fresh","description_category_id":10,"type_id":20,"price":"999","old_price":"1299","currency_code":"RUB","depth":300,"width":200,"height":100,"weight":500,"dimension_unit":"mm","weight_unit":"g","primary_image":"https://oss.test/generated-main.jpg","images":["https://oss.test/generated-extra.jpg"],"color_image":"https://oss.test/generated-main.jpg","attributes":[],"complex_attributes":[],"barcode":"","vat":"0","pdf_list":[]}
        child.content_json=dumps({"title":"fresh"});child.assets_json=dumps({"primary_image":item["primary_image"],"images":item["images"],"package_ready":True,"import_item":item});context.session.commit();return {"item":item,"definitions":[]}

    first=client.post("/api/v1/publications",json={"product_ids":[product_id],"store_ids":[store],"image_policy":"regenerate_all","submit_via_api":False,"save_generated_to_catalog":False}).json()["items"][0]
    with patch("erp.services.publications.prepare_store_package",fake_prepare): _run(first["job"]["id"])
    assert client.get(f"/api/v1/catalog/products/{product_id}").json()["current_asset_count"] == original
    second=client.post("/api/v1/publications",json={"product_ids":[product_id],"store_ids":[store],"image_policy":"regenerate_all","submit_via_api":False,"save_generated_to_catalog":True}).json()["items"][0]
    with patch("erp.services.publications.prepare_store_package",fake_prepare): _run(second["job"]["id"])
    detail=client.get(f"/api/v1/catalog/products/{product_id}").json()
    assert {link["asset"]["oss_url"] for link in detail["assets"]} == {"https://oss.test/generated-main.jpg","https://oss.test/generated-extra.jpg"}


def test_one_blocked_store_does_not_stop_others_and_can_retry_alone(client):
    product_id=_catalog_product(client);first=_store(client,"失败店","701");second=_store(client,"成功店","702")
    response=client.post("/api/v1/publications",json={"product_ids":[product_id],"store_ids":[first,second],"image_policy":"reuse","submit_via_api":False}).json()["items"][0]
    failed_once={"value":False}

    def fake_prepare(context,publication,child,snapshot):
        if child.store_id==first and not failed_once["value"]:
            failed_once["value"]=True
            raise AppError("需要补充准确资料",code="publication_facts_missing")
        item={"offer_id":child.offer_id,"name":f"fresh-{child.store_id}","description_category_id":10,"type_id":20,"price":"999","old_price":"1299","currency_code":"RUB","depth":300,"width":200,"height":100,"weight":500,"dimension_unit":"mm","weight_unit":"g","primary_image":"https://oss.test/main.jpg","images":[],"color_image":"https://oss.test/main.jpg","attributes":[],"complex_attributes":[],"barcode":"","vat":"0","pdf_list":[]}
        child.content_json=dumps({"title":item["name"]});child.assets_json=dumps({"primary_image":item["primary_image"],"images":[],"package_ready":True,"import_item":item});context.session.commit();return {"item":item,"definitions":[]}

    with patch("erp.services.publications.prepare_store_package",fake_prepare): assert _run(response["job"]["id"]) == "paused"
    publication_id=response["publication"]["id"];detail=client.get(f"/api/v1/publications/{publication_id}").json();states={row["store_id"]:row["status"] for row in detail["stores"]}
    assert states[first]=="blocked" and states[second]=="ready_for_export"
    assert client.post(f"/api/v1/publications/{publication_id}/stores/{second}/retry").status_code==409
    retry=client.post(f"/api/v1/publications/{publication_id}/stores/{first}/retry").json()
    with patch("erp.services.publications.prepare_store_package",fake_prepare): assert _run(retry["job"]["id"]) == "completed"
    detail=client.get(f"/api/v1/publications/{publication_id}").json()
    assert detail["status"]=="completed" and all(row["status"]=="ready_for_export" for row in detail["stores"])


def test_api_publication_uses_separate_stock_and_promotion_queues(client):
    product_id=_catalog_product(client);store=_store(client,"分阶段店铺","801","Q-")

    def fake_prepare(context,publication,child,snapshot):
        item={"offer_id":child.offer_id,"name":"fresh","description_category_id":10,"type_id":20,"price":"999","old_price":"1299","currency_code":"RUB","depth":300,"width":200,"height":100,"weight":500,"dimension_unit":"mm","weight_unit":"g","primary_image":"https://oss.test/main.jpg","images":[],"color_image":"https://oss.test/main.jpg","attributes":[],"complex_attributes":[],"barcode":"","vat":"0","pdf_list":[]}
        child.content_json=dumps({"title":"fresh"});child.pricing_json=dumps({"action_price":"899"});child.assets_json=dumps({"primary_image":item["primary_image"],"images":[],"package_ready":True,"import_item":item});context.session.commit();return {"item":item,"definitions":[]}

    class FakeAdapter:
        stock_calls=[];promotion_calls=[]
        def __init__(self,*_args,**_kwargs): pass
        def import_product(self,_item): return 901
        def wait_import(self,_task_id): return {"result":{"items":[{"status":"imported","product_id":902,"errors":[]}]}}
        def update_stock(self,offer_id,warehouse_id,stock): self.stock_calls.append((offer_id,warehouse_id,stock));return {}
        def find_promotion_candidate(self,action_id,product_id): return {"id":product_id,"action_id":action_id}
        def activate_promotion(self,action_id,product_id,action_price,stock): self.promotion_calls.append((action_id,product_id,action_price,stock));return {}

    created=client.post("/api/v1/publications",json={"product_ids":[product_id],"store_ids":[store],"image_policy":"reuse","submit_via_api":True}).json()["items"][0]
    assert created["job"]["max_attempts"]==3
    from erp.services.jobs import job_types_for_queue
    assert created["job"]["type"] == "catalog_ozon_prepare"
    assert created["job"]["type"] in job_types_for_queue("ai")
    with patch("erp.services.publications.prepare_store_package",fake_prepare),patch("erp.services.publications.OzonSellerAdapter",FakeAdapter):
        assert _run(created["job"]["id"])=="action_required"
    publication_id=created["publication"]["id"]
    detail=client.get(f"/api/v1/publications/{publication_id}").json();child=detail["stores"][0]
    assert child["status"]=="ready_for_submit" and not FakeAdapter.stock_calls
    waiting_job=client.get(f"/api/v1/jobs/{created['job']['id']}").json()
    assert waiting_job["progress"]<100
    assert [step["status"] for step in waiting_job["steps"][-2:]]==["pending","pending"]

    submit_job=client.post(f"/api/v1/publications/{publication_id}/stores/{store}/submit").json()["job"]
    assert submit_job["type"] == "catalog_ozon_publication"
    assert submit_job["type"] in job_types_for_queue("ozon")
    with patch("erp.services.publications.OzonSellerAdapter",FakeAdapter):
        assert _run(submit_job["id"])=="action_required"
    assert client.get(f"/api/v1/publications/{publication_id}").json()["stores"][0]["status"]=="stock_pending"

    with patch("erp.api.routes.publications.OzonSellerAdapter",FakeAdapter):
        response=client.post(f"/api/v1/publications/{publication_id}/stores/{store}/stock",json={"warehouse_id":12,"stock":7})
        assert response.status_code==200 and response.json()["status"]=="promotion_pending"
        quantities=client.get(f"/api/v1/publications/stock-quantities?child_ids={child['id']}")
        assert quantities.status_code==200 and quantities.json()['quantities'][child['id']]==7
        response=client.post(f"/api/v1/publications/{publication_id}/stores/{store}/promotion",json={"action_id":33,"action_price":899})
        assert response.status_code==200 and response.json()["status"]=="completed"
    assert FakeAdapter.stock_calls==[("Q-SKU-self_selected-M",12,7)]
    assert FakeAdapter.promotion_calls==[(33,902,899.0,None)]
    assert client.get(f"/api/v1/publications/{publication_id}").json()["status"]=="completed"
    completed_job=client.get(f"/api/v1/jobs/{submit_job['id']}").json()
    assert completed_job["status"]=="completed" and completed_job["progress"]==100


def test_product_exact_facts_and_url_gallery_can_be_maintained(client):
    product_id=_catalog_product(client)
    updated=client.patch(f"/api/v1/catalog/products/{product_id}",json={"package_weight_g":650,"notes":"已人工核对"}).json()
    assert updated["package_weight_g"]==650 and updated["notes"]=="已人工核对"
    changed=client.post(f"/api/v1/catalog/products/{product_id}/assets/url",json={"oss_url":"https://oss.test/new-main.jpg","role":"main","position":0,"set_current":True}).json()
    assert any(row["role"]=="main" and row["asset"]["oss_url"]=="https://oss.test/new-main.jpg" for row in changed["assets"])
    reread=client.get(f"/api/v1/catalog/products/{product_id}").json()
    assert [(row["role"],row["asset"]["oss_url"]) for row in reread["assets"] if row["role"]=="main"] == [("main","https://oss.test/new-main.jpg")]


@pytest.mark.parametrize("source_type", ["pod", "follow", "self_selected"])
def test_fake_supplier_ai_image_oss_and_ozon_close_each_product_line(source_type):
    harness=PublicationTestHarness();packages=harness.run(source_type,"SKU-1",["store-a","store-b"],"regenerate_main")
    assert len(packages)==2
    assert packages[0]["name"] != packages[1]["name"]
    assert all(package["primary_image"].startswith("https://fake-oss.invalid/") for package in packages)
    assert len(harness.supplier.calls)==1 and len(harness.ai.calls)==2 and len(harness.image.calls)==2 and len(harness.ozon.calls)==2
    original_calls=len(harness.ozon.calls);harness.ozon.submit(packages[0],f"{source_type}:SKU-1:store-a")
    assert len(harness.ozon.calls)==original_calls


def test_ozon_pod_batch_image_order_and_rates_are_frozen_in_publication(client):
    from erp.models import Publication
    from sqlalchemy import select
    from erp.core.jsonutil import loads
    product_id=_catalog_product(client,'pod')
    store_id=_store(client,'图片排序店','ordered-store','ORD-')
    detail=client.get(f'/api/v1/catalog/products/{product_id}').json()
    ids=[x['asset']['id'] for x in reversed(detail['assets'])]
    requirements={'price':100,'image_selection':{product_id:ids},'pricing_overrides':{'pricing_advertising_rate':12,'pricing_return_rate':18}}
    preview=client.post('/api/v1/publications/preview',json={'product_ids':[product_id],'store_ids':[store_id],'requirements':requirements})
    assert preview.status_code==200 and preview.json()['items'][0]['pricing_profile']['advertising_rate']==12
    created=client.post('/api/v1/publications',json={'product_ids':[product_id],'store_ids':[store_id],'requirements':requirements,'image_policy':'reuse','dry_run':True})
    assert created.status_code==202,created.text
    with SessionLocal() as db:
        publication=db.scalar(select(Publication))
        facts=loads(publication.product_snapshot_json,{})
        assert [x['asset_id'] for x in facts['assets']]==ids and facts['assets'][0]['role']=='main'
        assert loads(publication.requirements_json,{})['pricing_overrides']==requirements['pricing_overrides']
    invalid=client.post('/api/v1/publications',json={'product_ids':[product_id],'store_ids':[store_id],'requirements':{'image_selection':{product_id:['foreign-image']}}})
    assert invalid.status_code==409


def test_pod_filename_template_is_materialized_before_enqueue(client):
    from erp.models import Publication
    from erp.core.jsonutil import loads
    product_id=_catalog_product(client,'pod')
    store_id=_store(client,'模板店','template-store','T-')
    detail=client.get(f'/api/v1/catalog/products/{product_id}').json()
    keys=[x['asset']['name'].replace(chr(92),'/').split('/')[-1] for x in reversed(detail['assets'])]
    ids=[x['asset']['id'] for x in reversed(detail['assets'])]
    request={'product_ids':[product_id],'store_ids':[store_id],'requirements':{'price':100,'pod_image_template':keys},'image_policy':'reuse','dry_run':True}
    created=client.post('/api/v1/publications',json=request)
    assert created.status_code==202,created.text
    publication=created.json()['items'][0]['publication']
    with SessionLocal() as db:
        record=db.get(Publication,publication['id'])
        requirements=loads(record.requirements_json,{})
        assert 'pod_image_template' not in requirements
        assert requirements['image_selection'][product_id]==ids
        assert [x['asset_id'] for x in loads(record.product_snapshot_json,{})['assets']]==ids


def test_ozon_selected_variants_limit_preview_and_created_jobs(client):
    product_id=_catalog_product(client)
    response=client.put(f'/api/v1/catalog/products/{product_id}/variants',json=[
        {'sku':'SELECT-M','size':'M','color':'黑色'},
        {'sku':'SELECT-L','size':'L','color':'黑色'}])
    assert response.status_code==200
    variants=client.get(f'/api/v1/catalog/products/{product_id}').json()['variants']
    chosen=next(v['id'] for v in variants if v['sku']=='SELECT-L')
    store=_store(client,'规格店','selected-variant-store','SV-')
    body={'product_ids':[product_id],'variant_ids':[chosen],'store_ids':[store]}
    preview=client.post('/api/v1/publications/preview',json=body)
    assert preview.status_code==200
    assert [x['sku'] for x in preview.json()['items']]==['SELECT-L']
    created=client.post('/api/v1/publications',json={**body,'image_policy':'reuse'})
    assert created.status_code==202
    assert created.json()['publication_count']==1
    assert created.json()['items'][0]['publication']['variant_id']==chosen
    assert client.post('/api/v1/publications',json={**body,'variant_ids':['foreign-variant']}).status_code==409
    assert client.post('/api/v1/publications/preview',json={**body,'variant_ids':[]}).status_code==422
    listing=client.get('/api/v1/catalog/products?include_variants=true').json()
    assert len(listing['items'][0]['variants'])==2
