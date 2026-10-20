from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
import pytest
from sqlalchemy import select

from erp.core.database import SessionLocal
from erp.core.jsonutil import dumps
from erp.core.security import vault
from erp.domains.pod import wb_export
from erp.domains.pod import wb_history_transport
from erp.models import Asset, CatalogProduct, CatalogProductAsset, CatalogVariant, IntegrationConnection, Job, Store, WbPublicationItem
from erp.services.jobs import execute_job, get_job


def _template() -> bytes:
    workbook = Workbook(); sheet = workbook.active; sheet.title = "商品"
    names = {
        "static_groupID_test": "A", "static_vendorCode_test": "B", "static_subject_test": "C",
        "static_photo_test": "D", "static_price_test": "E", "characteristic_15000000_1": "F",
        "characteristic_14177446_1": "G", "characteristic_14177452_1": "H",
        "characteristic_90849_1": "I", "characteristic_90745_1": "J", "characteristic_90846_1": "K",
        "characteristic_244146_1": "L", "characteristic_90673_1": "M", "characteristic_90630_1": "N",
    }
    for name, column in names.items():
        workbook.defined_names.add(DefinedName(name, attr_text=f"'商品'!${column}$3"))
    stream = BytesIO(); workbook.save(stream); return stream.getvalue()


def _template_with_wb_full_column_validation() -> bytes:
    content = _template()
    source, target = BytesIO(content), BytesIO()
    from zipfile import ZIP_DEFLATED, ZipFile
    with ZipFile(source, "r") as incoming, ZipFile(target, "w", ZIP_DEFLATED) as outgoing:
        for info in incoming.infolist():
            value = incoming.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                text = value.decode("utf-8")
                validation = '<dataValidations count="1"><dataValidation sqref="E:E" type="decimal" operator="greaterThan"><formula1>0</formula1></dataValidation></dataValidations>'
                text = text.replace("</worksheet>", validation + "</worksheet>")
                value = text.encode("utf-8")
            outgoing.writestr(info, value)
    return target.getvalue()


def test_pod_wb_excel_background_export(client, monkeypatch):
    template = client.post("/api/v1/pod/wb/templates", files={"file": ("wb.xlsx", _template(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, data={"subject_name": "Пледы", "subject_id": "321"})
    assert template.status_code == 201
    with SessionLocal() as session:
        store = Store(platform="wb", name="WB 测试店", external_account_id="wb-test", offer_prefix="W-", currency="CNY", enabled=True)
        ai = IntegrationConnection(provider="ai_chat", name="micuapi", base_url="https://ai.invalid", config_json=dumps({"default_model": "gpt-5.6-terra"}), secrets_encrypted=vault.encrypt({"api_key": "token"}), enabled=True)
        product = CatalogProduct(sku="HC-000001", group_code="HC-000001", source_type="pod", name="测试毛毯", supplier="HICUSTOM", package_length_mm=300, package_width_mm=200, package_height_mm=100, package_weight_g=500, status="ready")
        session.add_all([store, ai, product]); session.flush()
        selected_variant = CatalogVariant(product_id=product.id, sku="HC-000001-01", color="蓝", size="125x150 cm", properties_json=dumps({"color_id": 1}), active=True)
        skipped_variant = CatalogVariant(product_id=product.id, sku="HC-000001-02", color="红", size="150x200 cm", properties_json=dumps({"color_id": 2}), active=True)
        session.add_all([selected_variant, skipped_variant]); session.flush()
        blue = Asset(kind="rendering", name="blue.jpg", source="test", oss_url="https://oss.example/blue.jpg", mime_type="image/jpeg")
        red = Asset(kind="image", name="red.jpg", source="test", oss_url="https://oss.example/red.jpg", mime_type="image/jpeg")
        session.add_all([blue, red]); session.flush()
        session.add_all([CatalogProductAsset(product_id=product.id, variant_id=selected_variant.id, asset_id=blue.id, role="main", position=0, is_current=True), CatalogProductAsset(product_id=product.id, variant_id=skipped_variant.id, asset_id=red.id, role="main", position=0, is_current=True)]); session.commit()
        store_id, product_id, selected_variant_id = store.id, product.id, selected_variant.id
    preview = client.post("/api/v1/pod/wb/exports/preview", json={"product_ids": [product_id], "subject_id": 321, "variant_selection": "selected", "variant_ids": [selected_variant_id]})
    assert preview.status_code == 200
    assert preview.json()["product_count"] == 1
    assert preview.json()["variant_count"] == 1
    assert preview.json()["card_count"] == 1
    assert preview.json()["group_count"] == 1
    queued = client.post("/api/v1/pod/wb/exports", json={"product_ids": [product_id], "store_id": store_id, "template_id": template.json()["template"]["id"], "brand": "No brand", "price_cny": 99.5, "variant_selection": "selected", "variant_ids": [selected_variant_id], "image_rules": [{"key": "blue.jpg", "action": "main", "position": 0}], "requirement": "简洁"})
    assert queued.status_code == 202
    monkeypatch.setattr(wb_export, "_generate_copy", lambda *_: {"title": "Плед тест", "description": "Описание теста", "characteristics": []})
    with SessionLocal() as session:
        execute_job(session, get_job(session, queued.json()["job"]["id"]))
    download = client.get(f"/api/v1/pod/wb/exports/{queued.json()['job']['id']}/download")
    assert download.status_code == 200
    workbook = load_workbook(BytesIO(download.content), data_only=True); sheet = workbook[workbook.sheetnames[0]]
    assert sheet["B5"].value == "W-HC-000001-01"
    assert sheet["C5"].value == "Пледы"
    assert sheet["E5"].value == 99.5
    assert sheet["F5"].value == "Плед тест"
    assert sheet["D5"].value == "https://oss.example/blue.jpg"
    assert sheet["L5"].value == "125x150 см"
    assert sheet["M5"].value == 125
    assert sheet["N5"].value == 150
    availability = client.post("/api/v1/pod/wb/availability", json={"store_id": store_id, "product_ids": [product_id]})
    assert availability.status_code == 200
    assert availability.json()["products"][product_id]["status"] == "exported_pending_confirmation"
    duplicate = client.post("/api/v1/pod/wb/exports", json={"product_ids": [product_id], "store_id": store_id, "template_id": template.json()["template"]["id"], "brand": "No brand", "price_cny": 99.5, "variant_selection": "selected", "variant_ids": [selected_variant_id]})
    assert duplicate.status_code == 400
    confirmed = client.post(f"/api/v1/pod/wb/exports/{queued.json()['job']['id']}/confirm-upload")
    assert confirmed.status_code == 200
    assert confirmed.json()["status"] == "submitted"
    assert sheet["B6"].value is None
    history = client.get("/api/v1/pod/wb/exports").json()["items"]
    assert history[0]["status"] == "completed" and history[0]["download_ready"] is True
    assert history[0]["publication_status"] == "submitted"
    assert client.get("/api/v1/pod/wb/publications?stage=processing").json()["total"] == 1
    # Older versions marked Excel uploads complete without a WB product ID.
    with SessionLocal() as session:
        ledger = session.scalar(select(WbPublicationItem).where(WbPublicationItem.job_id == queued.json()["job"]["id"]))
        ledger.status = "completed"
        session.commit()
    assert client.get("/api/v1/pod/wb/exports").json()["items"][0]["publication_status"] == "submitted"
    class FakeContentClient:
        def __init__(self, _config): pass
        def find_cards_exact(self, offers):
            return {offer: {"vendorCode": offer, "nmID": 9001} for offer in offers}
        def close(self): pass
    monkeypatch.setattr(wb_export, "_wb_store_snapshot", lambda session, identity: (session.get(Store, identity), vault.encrypt({"store": {"secrets": {"token": "test"}}})))
    monkeypatch.setattr(wb_export, "_WbContentClient", FakeContentClient)
    reconciled = client.post(f"/api/v1/pod/wb/exports/{queued.json()['job']['id']}/reconcile-upload")
    assert reconciled.status_code == 200
    assert reconciled.json()["matched"] == 1
    stock_queue = client.get("/api/v1/pod/wb/publications?stage=pending_stock").json()
    assert stock_queue["total"] == 1 and stock_queue["items"][0]["nm_id"] == "9001"


def test_wb_copy_failure_skips_one_product_and_exports_the_other(client, monkeypatch):
    template = client.post("/api/v1/pod/wb/templates", files={"file": ("wb.xlsx", _template(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, data={"subject_name": "Пледы", "subject_id": "321"})
    assert template.status_code == 201
    with SessionLocal() as session:
        store = Store(platform="wb", name="WB 部分成功测试", external_account_id="wb-partial", offer_prefix="P-", currency="CNY", enabled=True)
        ai = IntegrationConnection(provider="ai_chat", name="文案接口", base_url="https://ai.invalid", config_json=dumps({"default_model": "test"}), secrets_encrypted=vault.encrypt({"api_key": "token"}), enabled=True)
        session.add_all([store, ai]); session.flush()
        products = []
        for index in range(2):
            product = CatalogProduct(sku=f"COPY-{index}", group_code=f"COPY-{index}", source_type="pod", name=f"毛毯 {index}", supplier="HICUSTOM", package_length_mm=300, package_width_mm=200, package_height_mm=100, package_weight_g=500, status="ready")
            session.add(product); session.flush()
            variant = CatalogVariant(product_id=product.id, sku=f"COPY-{index}-01", size="125x150 cm", active=True)
            asset = Asset(kind="rendering", name=f"copy-{index}.jpg", source="test", oss_url=f"https://oss.example/copy-{index}.jpg", mime_type="image/jpeg")
            session.add_all([variant, asset]); session.flush()
            session.add(CatalogProductAsset(product_id=product.id, variant_id=variant.id, asset_id=asset.id, role="main", position=0, is_current=True))
            products.append((product.id, variant.id))
        session.commit()
        store_id = store.id
    queued = client.post("/api/v1/pod/wb/exports", json={"product_ids": [row[0] for row in products], "store_id": store_id, "template_id": template.json()["template"]["id"], "brand": "No brand", "price_cny": 99.5, "variant_selection": "selected", "variant_ids": [row[1] for row in products]})
    assert queued.status_code == 202
    calls = []
    def generate(_config, product, _template, _payload):
        calls.append(product.id)
        if product.id == products[0][0]: raise ValueError("upstream unavailable")
        return {"title": "Плед", "description": "Описание", "characteristics": []}
    monkeypatch.setattr(wb_export, "_generate_copy", generate)
    job_id = queued.json()["job"]["id"]
    with SessionLocal() as session:
        execute_job(session, get_job(session, job_id))
        job = get_job(session, job_id)
        assert job.status == "completed" and job.max_attempts == 1, job.error_message
        ledgers = {row.product_id: row for row in session.scalars(select(WbPublicationItem).where(WbPublicationItem.job_id == job_id)).all()}
        assert ledgers[products[0][0]].status == "failed"
        assert "upstream unavailable" in ledgers[products[0][0]].last_error
        assert ledgers[products[1][0]].status == "exported_pending_confirmation"
        assert any(item.role == "wb_copy_failures" for item in job.artifacts)
    assert calls.count(products[0][0]) == 3 and calls.count(products[1][0]) == 1
    assert client.get(f"/api/v1/pod/wb/exports/{job_id}/download").status_code == 200
    assert client.get("/api/v1/pod/wb/exports").json()["items"][0]["copy_failed_items"] == 1
    confirmed = client.post(f"/api/v1/pod/wb/exports/{job_id}/confirm-upload")
    assert confirmed.status_code == 200 and confirmed.json()["items"] == 1


def test_wb_template_requires_official_columns(client):
    workbook = Workbook(); stream = BytesIO(); workbook.save(stream)
    response = client.post("/api/v1/pod/wb/templates", files={"file": ("bad.xlsx", stream.getvalue())}, data={"subject_name": "Пледы", "subject_id": "123"})
    assert response.status_code == 400
    assert response.json()["code"] == "wb_template_invalid"


def test_wb_official_template_full_column_validations_are_normalized(client):
    response = client.post(
        "/api/v1/pod/wb/templates",
        files={"file": ("official.xlsx", _template_with_wb_full_column_validation())},
        data={"subject_name": "毛毯", "subject_id": "321"},
    )
    assert response.status_code == 201
    item = response.json()["template"]
    assert item["schema"]["sheet"] == "商品"
    stored = wb_export._template(item["id"])
    workbook = load_workbook(stored["path"], read_only=False)
    assert str(workbook.active.data_validations.dataValidation[0].sqref) == "E1:E1048576"


def test_wb_excel_rows_are_naturally_sorted_by_sku(tmp_path):
    template = tmp_path / "template.xlsx"
    template.write_bytes(_template())
    output = tmp_path / "sorted.xlsx"

    def card(sku: str) -> dict:
        return {
            "vendorCode": f"STORE-{sku}", "title": sku, "description": sku,
            "brand": "No brand", "price": 100,
            "dimensions": {"length": 30, "width": 20, "height": 10, "weightBrutto": .5},
            "characteristics": [], "images": [f"https://oss.example/{sku}.jpg"],
            "sizes": [{"sku": sku}], "product_id": sku, "product_sku": sku,
        }

    groups = [
        {"variants": [card("SKU-10"), card("SKU-1")]},
        {"variants": [card("SKU-2")]},
    ]
    assert wb_export._write_excel(template, output, groups, "Пледы") == 3
    sheet = load_workbook(output, data_only=True).active
    assert [sheet[f"B{row}"].value for row in range(5, 8)] == [
        "STORE-SKU-1", "STORE-SKU-2", "STORE-SKU-10",
    ]


def test_wb_excel_obeys_template_characteristic_value_counts(tmp_path):
    template = tmp_path / "template.xlsx"
    workbook = load_workbook(BytesIO(_template()))
    workbook.defined_names.add(DefinedName("characteristic_12_1", attr_text="'商品'!$O$3"))
    workbook.defined_names.add(DefinedName("characteristic_377_1", attr_text="'商品'!$P$3"))
    workbook.active["O4"] = "值的最大数量: 1"
    workbook.active["P4"] = "值的最大数量: 3"
    workbook.save(template)
    output = tmp_path / "export.xlsx"
    card = {
        "vendorCode": "SKU-1", "title": "桌布", "description": "描述", "brand": "No brand", "price": 100,
        "dimensions": {"length": 30, "width": 20, "height": 10, "weightBrutto": .5},
        "characteristics": [{"id": 12, "value": ["花卉", "植物", "叶子"]},
                            {"id": 377, "value": ["餐桌", "客厅", "厨房", "阳台"]}],
        "images": ["https://oss.example/sku.jpg"], "sizes": [], "product_id": "1", "product_sku": "SKU-1",
    }
    assert wb_export._template_value_limits(template) == {12: 1, 377: 3}
    wb_export._write_excel(template, output, [{"variants": [card]}], "桌布")
    sheet = load_workbook(output, data_only=True).active
    assert sheet["O5"].value == "花卉"
    assert sheet["P5"].value == "餐桌;客厅;厨房"


def test_personal_wb_reconcile_accepts_only_requested_offer_ids():
    from erp.domains.pod.desktop_wb import accept_reconcile_result
    original = 'a' * 32
    row1 = WbPublicationItem(job_id=original, store_id='store', offer_id='A-1', status='submitted')
    row2 = WbPublicationItem(job_id=original, store_id='store', offer_id='A-2', status='submitted')
    class FakeSession:
        def scalars(self, statement): return iter([row1, row2])
    job = SimpleNamespace(type='pod_wb_excel_reconcile', payload_json=dumps({
        'export_job_id': original, 'store_id': 'store', 'offer_ids': ['A-1', 'A-2']}))
    with pytest.raises(Exception):
        accept_reconcile_result(FakeSession(), job, {'found': {'other-store': 991}})
    assert row1.external_product_id in ('', None) and row2.external_product_id in ('', None)
    accept_reconcile_result(FakeSession(), job, {'found': {'A-1': 991}})
    assert row1.external_product_id == '991' and row1.status == 'pending_stock'
    assert row2.status == 'submitted' and row2.external_product_id in ('', None)


def test_personal_wb_stock_accepts_only_verified_selected_items():
    from erp.domains.pod.desktop_wb import accept_stock_result
    identity = 'a' * 32
    row = WbPublicationItem(id=identity, store_id='store', offer_id='SKU-1',
                            status='stock_processing', external_product_id='123')
    class FakeSession:
        def scalars(self, statement): return iter([row])
    job = SimpleNamespace(type='pod_wb_stock', payload_json=dumps({
        'ledger_ids': [identity], 'store_id': 'store', 'warehouse_id': 2073831, 'amount': 100}))
    result = {'verified': True, 'warehouse_id': 2073831, 'amount': 100,
              'items': [{'id': identity, 'nm_id': '123', 'chrt_id': 456}]}
    with pytest.raises(Exception):
        accept_stock_result(FakeSession(), job, {**result, 'warehouse_id': 999})
    assert row.status == 'stock_processing'
    accept_stock_result(FakeSession(), job, result)
    assert row.status == 'pending_price'


@pytest.mark.parametrize("card_preexists", [False, True])
def test_pod_wb_seller_api_is_idempotent_and_records_nm_id(client, monkeypatch, card_preexists):
    template = client.post("/api/v1/pod/wb/templates", files={"file": ("wb.xlsx", _template())}, data={"subject_name": "Пледы", "subject_id": "321"}).json()["template"]
    with SessionLocal() as session:
        store=Store(platform="wb",name="WB API",external_account_id="wb-api",offer_prefix="API",currency="RUB",enabled=True,secrets_encrypted=vault.encrypt({"token":"secret"}))
        ai=IntegrationConnection(provider="ai_chat",name="AI",base_url="https://ai.invalid",config_json=dumps({}),secrets_encrypted=vault.encrypt({"api_key":"token"}),enabled=True)
        product=CatalogProduct(sku="POD-1",group_code="POD-1",source_type="pod",name="毛毯",supplier="S2B",package_length_mm=300,package_width_mm=200,package_height_mm=100,package_weight_g=500,status="ready")
        session.add_all([store,ai,product]);session.flush()
        variant=CatalogVariant(product_id=product.id,sku="POD-1-01",size="125x150 cm",purchase_price_cny=20,active=True)
        session.add(variant);session.flush();asset=Asset(kind="image",name="01.jpg",source="test",oss_url="https://oss.example/01.jpg");session.add(asset);session.flush()
        session.add(CatalogProductAsset(product_id=product.id,variant_id=variant.id,asset_id=asset.id,role="main",position=0,is_current=True));session.commit()
        store_id,product_id,variant_id=store.id,product.id,variant.id

    created:dict[str,dict]={};uploads=[]
    class FakeWb:
        def __init__(self,config):pass
        def close(self):pass
        def generate_barcodes(self,count):return [f"BC{index}" for index in range(count)]
        def find_card(self,vendor,aliases=None):return created.get(vendor)
        def find_cards_exact(self,vendors):return {vendor:created[vendor] for vendor in vendors if vendor in created}
        def create_cards(self,groups):
            for group in groups:
                for card in group["variants"]:created[card["vendorCode"]]={"vendorCode":card["vendorCode"],"nmID":9001}
        def errors(self):return {}
        def save_media(self,nm_id,urls):uploads.append((nm_id,urls))
    monkeypatch.setattr(wb_export,"_WbContentClient",FakeWb)
    monkeypatch.setattr(wb_export,"_wb_characteristics",lambda *_:[])
    monkeypatch.setattr(wb_export,"_generate_copy",lambda *_:{"title":"Плед","description":"Описание","characteristics":[]})
    queued=client.post("/api/v1/pod/wb/exports",json={"product_ids":[product_id],"store_id":store_id,"template_id":template["id"],"brand":"Нет бренда","price_cny":999,"mode":"api","variant_selection":"selected","variant_ids":[variant_id]})
    assert queued.status_code==202
    if card_preexists:
        created["API-POD-1-01"]={"vendorCode":"API-POD-1-01","nmID":9001}
    with SessionLocal() as session:
        job=get_job(session,queued.json()["job"]["id"]);execute_job(session,job)
        row=session.scalar(select(WbPublicationItem).where(WbPublicationItem.job_id==job.id))
        assert job.status=="completed",job.error_message;assert row.status=="pending_stock";assert row.external_product_id=="9001"
        execute_job(session,job)
    assert uploads==[(9001,["https://oss.example/01.jpg"])]
    queue=client.get("/api/v1/pod/wb/publications?stage=pending_stock")
    assert queue.status_code==200
    assert queue.json()["page"]==1 and queue.json()["page_size"]==30 and queue.json()["total"]>=1
    queued_rows=[item for item in queue.json()["items"] if item["job_id"]==queued.json()["job"]["id"]]
    assert len(queued_rows)==1
    assert queued_rows[0]["nm_id"]=="9001"
    assert queued_rows[0]["price"]==999
    for term in ("POD-1", "API-POD-1-01", "9001"):
        found=client.get("/api/v1/pod/wb/publications",params={"stage":"pending_stock","search":term})
        assert found.status_code==200
        assert any(item["job_id"]==queued.json()["job"]["id"] for item in found.json()["items"])
    missing=client.get("/api/v1/pod/wb/publications",params={"stage":"pending_stock","search":"NO-SUCH-WB-SKU"})
    assert missing.json()["total"]==0
    availability=client.post("/api/v1/pod/wb/availability",json={"store_id":store_id,"product_ids":[product_id]}).json()["products"][product_id]
    assert availability["status"]=="pending_stock" and availability["blocked"] is True
    available=client.get(f"/api/v1/pod/library/products?page=1&page_size=30&wb_store_id={store_id}&wb_state=available").json()
    assert product_id not in {item["id"] for item in available["items"]}
    with SessionLocal() as session:
        row=session.scalar(select(WbPublicationItem).where(WbPublicationItem.product_id==product_id,WbPublicationItem.store_id==store_id));row.status="pending_price";session.commit()
    availability=client.post("/api/v1/pod/wb/availability",json={"store_id":store_id,"product_ids":[product_id]}).json()["products"][product_id]
    assert availability["status"]=="pending_price" and availability["blocked"] is True


def test_wb_legacy_offer_can_match_contained_erp_sku(monkeypatch):
    wb=object.__new__(wb_export._WbContentClient)
    cards={"API-POD-1-01":[],"POD-1":[{"vendorCode":"OLDPREFIX_POD-1_campaign","nmID":88}],"POD-1-01":[]}
    monkeypatch.setattr(wb,"_search_cards",lambda text:cards.get(text,[]))
    found=wb.find_card("API-POD-1-01",["POD-1","POD-1-01"])
    assert found and found["nmID"]==88
    assert found["_erp_match"]=="legacy_contains_erp_sku"


def test_wb_exact_cursor_lookup_recovers_card_missing_from_text_search(monkeypatch):
    wb=object.__new__(wb_export._WbContentClient)
    pages=[{"cards":[{"vendorCode":"WXZ-S2B-003239-01","nmID":1619725735}],"cursor":{}}]
    monkeypatch.setattr(wb,"request",lambda *_args,**_kwargs:pages.pop(0))
    found=wb.find_cards_exact({"WXZ-S2B-003239-01"})
    assert found["WXZ-S2B-003239-01"]["nmID"]==1619725735
    assert found["WXZ-S2B-003239-01"]["_erp_match"]=="current_offer"


def test_wb_image_rules_preserve_the_numbered_user_order():
    candidates=[{"key":f"effect-0{i}.jpg","url":f"https://oss.example/{i}.jpg"} for i in range(1,8)]
    rules=[
        *[{"key":f"effect-0{i}.jpg","action":"additional","position":i-2} for i in range(2,8)],
        {"key":"effect-01.jpg","action":"main","position":6},
    ]
    assert wb_export._apply_image_rules(candidates,rules)==[
        *(f"https://oss.example/{i}.jpg" for i in range(2,8)),
        "https://oss.example/1.jpg",
    ]


def test_wb_api_card_uses_dimensions_instead_of_legacy_characteristics():
    card={
        "vendorCode":"SHOP-POD-1","title":"Скатерть","description":"Описание","brand":"Нет бренда",
        "dimensions":{"length":30,"width":20,"height":10,"weightBrutto":0.5},
        "characteristics":[
            {"id":88952,"value":500},{"id":88953,"value":500},
            {"id":90745,"value":20},{"id":90846,"value":10},{"id":90849,"value":30},
            {"id":14177449,"value":["белый"]},
        ],
        "sizes":[{"variant_id":7,"techSize":"140x180","price":999}],
    }
    payload=wb_export._api_card(card,{"7":"4600000000001"})
    assert payload["dimensions"]==card["dimensions"]
    assert payload["characteristics"]==[{"id":14177449,"value":["белый"]}]
    assert payload["sizes"]==[{"price":999,"skus":["4600000000001"]}]


def test_wb_variant_color_uses_erp_fact_instead_of_ai_guess():
    class Variant:
        color="白色"
    assert wb_export._variant_color_characteristic(Variant())=={"id":14177449,"value":["белый"]}


def test_wb_compliance_name_fields_are_always_omitted():
    schema=[
        {"charcID":15000785,"name":"Торговое наименование","charcType":1,"required":False,"maxCount":1},
        {"charcID":62945173,"name":"Свидетельство о регистрации СГР","charcType":1,"required":False,"maxCount":1},
        {"charcID":14177449,"name":"Цвет","charcType":1,"required":False,"maxCount":1},
    ]
    values=[
        {"id":15000785,"value":"Название товара"},
        {"id":62945173,"value":"Название товара"},
        {"id":14177449,"value":["белый"]},
    ]
    assert wb_export._validated_characteristics(values,schema)==[{"id":14177449,"value":["белый"]}]


def test_wb_automatic_price_does_not_apply_currency_exchange_rate():
    class Product:
        purchase_price_cny = 20
        domestic_delivery_cny = 0
        package_weight_g = 500

    class Variant:
        purchase_price_cny = None
        weight_grams = None

    low_rate = wb_export._variant_price(
        Product(), Variant(),
        {"mode": "multiplier", "multiplier": 3, "discount_percent": 30, "rub_to_cny_rate": 0.01},
    )
    high_rate = wb_export._variant_price(
        Product(), Variant(),
        {"mode": "multiplier", "multiplier": 3, "discount_percent": 30, "rub_to_cny_rate": 1},
    )
    assert low_rate["list_price_rub"] == 86
    assert high_rate["list_price_rub"] == 86
    assert low_rate["sale_price_cny"] == 60.2


def test_wb_history_transport_requires_complete_identical_parts_before_ledger(client):
    import hashlib
    from erp.core.errors import ConflictError
    with SessionLocal() as session:
        store=Store(platform='wb',name='WB 历史传输',external_account_id='history-transport',offer_prefix='NEW',enabled=True,
                    secrets_encrypted=vault.encrypt({'token':'secret'}))
        product=CatalogProduct(sku='ERP-9900',source_type='pod',name='桌布',supplier='S2B')
        session.add_all([store,product]);session.flush()
        variant=CatalogVariant(product_id=product.id,sku='ERP-9900-01',active=True)
        session.add(variant);session.commit()
        job=Job(id='a'*32,type='pod_wb_history_reconcile',idempotency_key='history-transport-test',
                payload_json=dumps({'store_id':store.id,'store_name':store.name,'offer_prefix':'NEW','stock_workflow':True}))
        session.add(job);session.commit()
        first=[{'vendorCode':'OLD-ERP-9900-01','nmID':99001}]
        raw=dumps(first).encode()
        wb_history_transport.stage(job,0,first)
        wb_history_transport.stage(job,0,first)
        with pytest.raises(ConflictError):wb_history_transport.stage(job,0,[{'vendorCode':'OTHER','nmID':99}])
        with pytest.raises(ConflictError):wb_history_transport.accept(session,job,{'chunks':2,'cards':1,'sha256':hashlib.sha256(raw).hexdigest()})
        assert session.scalar(select(WbPublicationItem).where(WbPublicationItem.variant_id==variant.id)) is None
        result=wb_history_transport.accept(session,job,{'chunks':1,'cards':1,'sha256':hashlib.sha256(raw).hexdigest()})
        session.commit()
        row=session.scalar(select(WbPublicationItem).where(WbPublicationItem.variant_id==variant.id))
        assert result['matched_variants']==1 and row.status=='pending_stock' and row.external_product_id=='99001'


def test_history_scan_preserves_existing_active_publication(client):
    with SessionLocal() as session:
        store=Store(platform='wb',name='WB 已有任务',external_account_id='history-existing',offer_prefix='NEW',enabled=True)
        product=CatalogProduct(sku='ERP-9910',source_type='pod',name='桌布',supplier='S2B')
        session.add_all([store,product]);session.flush()
        variant=CatalogVariant(product_id=product.id,sku='ERP-9910-01',active=True)
        original=Job(
            type='pod_wb_api_publish',
            idempotency_key="test",
        )
        session.add_all([variant,original]);session.flush()
        row=WbPublicationItem(store_id=store.id,product_id=product.id,variant_id=variant.id,job_id=original.id,
                              offer_id='NEW-ERP-9910-01',mode='api',status='running',template_id='',last_error='still running')
        session.add(row);session.commit()
        scan=Job(type='pod_wb_history_reconcile',idempotency_key='scan-9910',
                 payload_json=dumps({'store_id':store.id,'store_name':store.name,'offer_prefix':'NEW','stock_workflow':False}))
        session.add(scan);session.commit()
        wb_export.apply_history_cards(session,scan,[{'vendorCode':'NEW-ERP-9910-01','nmID':99101}]);session.commit()
        session.refresh(row)
        assert row.job_id==original.id and row.status=='running' and row.mode=='api'
        assert row.external_product_id=='99101' and row.last_error=='still running'


def test_wb_history_scan_links_existing_pod_card(client,monkeypatch):
    with SessionLocal() as session:
        store=Store(platform="wb",name="历史店铺",external_account_id="history",offer_prefix="NEW",currency="RUB",enabled=True,secrets_encrypted=vault.encrypt({"token":"secret"}))
        product=CatalogProduct(sku="ERP-7788",group_code="ERP-7788",source_type="pod",name="桌布",supplier="S2B",status="ready")
        session.add_all([store,product]);session.flush();variant=CatalogVariant(product_id=product.id,sku="ERP-7788-01",size="140x180",active=True);session.add(variant);session.commit();store_id,product_id,variant_id=store.id,product.id,variant.id
    class HistoryWb:
        def __init__(self,config):pass
        def close(self):pass
        def list_all_cards(self):return [{"vendorCode":"OLD_SHOP_ERP-7788-01_2024","nmID":778899}]
    monkeypatch.setattr(wb_export,"_WbContentClient",HistoryWb)
    queued=client.post("/api/v1/pod/wb/history/reconcile",json={"store_id":store_id})
    assert queued.status_code==202
    with SessionLocal() as session:
        job=get_job(session,queued.json()["job"]["id"]);execute_job(session,job)
        row=session.scalar(select(WbPublicationItem).where(WbPublicationItem.store_id==store_id,WbPublicationItem.variant_id==variant_id))
        assert job.status=="completed",job.error_message
        assert row and row.external_product_id=="778899" and row.mode=="history"
    report=client.get(f"/api/v1/pod/wb/history/{queued.json()['job']['id']}/report")
    assert report.status_code==200 and report.json()['summary']['matched_variants']==1
    availability=client.post("/api/v1/pod/wb/availability",json={"store_id":store_id,"product_ids":[product_id]}).json()
    assert availability["products"][product_id]["blocked"] is True
    available=client.get(f"/api/v1/pod/library/products?wb_store_id={store_id}&wb_state=available")
    listed=client.get(f"/api/v1/pod/library/products?wb_store_id={store_id}&wb_state=listed")
    assert available.status_code==200 and available.json()["total"]==0
    assert listed.status_code==200 and listed.json()["total"]==1


def test_wb_sync_uploaded_recovers_unverified_excel(client,monkeypatch):
    with SessionLocal() as session:
        store=Store(platform="wb",name="Excel 旧店",external_account_id="excel-old",offer_prefix="NEW",currency="RUB",enabled=True,secrets_encrypted=vault.encrypt({"token":"secret"}))
        product=CatalogProduct(sku="ERP-8800",group_code="ERP-8800",source_type="pod",name="桌布",supplier="S2B",status="ready")
        session.add_all([store,product]);session.flush()
        variant=CatalogVariant(product_id=product.id,sku="ERP-8800-01",size="140x180",active=True);session.add(variant);session.flush()
        ledger=WbPublicationItem(store_id=store.id,product_id=product.id,variant_id=variant.id,offer_id="OLD-ERP-8800-01",mode="excel",status="completed",template_id="")
        session.add(ledger);session.commit();store_id,ledger_id=store.id,ledger.id
    class UploadedWb:
        def __init__(self,config):pass
        def close(self):pass
        def list_all_cards(self):return [{"vendorCode":"OLD-ERP-8800-01","nmID":880001}]
    monkeypatch.setattr(wb_export,"_WbContentClient",UploadedWb)
    queued=client.post("/api/v1/pod/wb/history/sync-uploaded",json={"store_id":store_id})
    assert queued.status_code==202
    with SessionLocal() as session:
        job=get_job(session,queued.json()["job"]["id"]);execute_job(session,job)
        row=session.get(WbPublicationItem,ledger_id)
        assert job.status=="completed",job.error_message
        assert row.status=="pending_stock" and row.external_product_id=="880001" and row.offer_id=="OLD-ERP-8800-01"
    assert client.get("/api/v1/pod/wb/publications?stage=pending_stock").json()["total"]==1


def test_wb_stock_resolves_batch_in_one_card_scan(client,monkeypatch):
    with SessionLocal() as session:
        store=Store(platform="wb",name="批量库存店",external_account_id="wb-stock-batch",offer_prefix="B",currency="RUB",enabled=True,secrets_encrypted=vault.encrypt({"token":"secret"}))
        product=CatalogProduct(sku="B-1",group_code="B-1",source_type="pod",name="桌布",supplier="S2B",status="ready")
        session.add_all([store,product]);session.flush()
        variants=[CatalogVariant(product_id=product.id,sku=f"B-1-{index}",active=True) for index in (1,2)]
        session.add_all(variants);session.flush()
        ledgers=[WbPublicationItem(store_id=store.id,product_id=product.id,variant_id=variant.id,offer_id=f"B-{variant.sku}",mode="excel",status="pending_stock",template_id="",external_product_id=str(100+index)) for index,variant in enumerate(variants)]
        session.add_all(ledgers);session.commit();ledger_ids=[row.id for row in ledgers]
    calls=[]
    class StockWb:
        token="secret"
        def __init__(self,config):pass
        def close(self):pass
        def find_cards_exact(self,offers):
            calls.append(set(offers))
            return {offer:{"sizes":[{"chrtID":index+1}]} for index,offer in enumerate(sorted(offers))}
    monkeypatch.setattr(wb_export,"_WbContentClient",StockWb)
    monkeypatch.setattr(wb_export.httpx,"put",lambda *args,**kwargs:SimpleNamespace(status_code=204,text=""))
    queued=client.post("/api/v1/pod/wb/publications/stock",json={"ledger_ids":ledger_ids,"warehouse_id":1,"amount":5})
    assert queued.status_code==202
    with SessionLocal() as session:
        job=get_job(session,queued.json()["job"]["id"]);execute_job(session,job)
        assert job.status=="completed",job.error_message
        assert all(session.get(WbPublicationItem,identity).status=="pending_price" for identity in ledger_ids)
    assert len(calls)==1 and len(calls[0])==2
