"""The browser's directory path belongs to its selected Windows computer."""
import pytest
from types import SimpleNamespace
from PIL import Image
import sys
import threading
from pathlib import Path
from fastapi.testclient import TestClient
import base64
import hashlib
import io
from openpyxl import Workbook

from erp.api.routes import catalog
from erp.core.config import get_settings
from erp.core.database import SessionLocal
from erp.core.errors import ConflictError
from erp.schemas import SelfListingPipelineCreate
from erp.main import app
from erp.services.accounts import create_user
from erp.core.security import vault
from erp.models import IntegrationConnection


def test_cloud_preserves_original_computer_directory_and_never_stats_it(monkeypatch):
    monkeypatch.setattr(get_settings(), 'multiuser_enabled', True)
    monkeypatch.setattr(get_settings(), 'runtime_mode', 'server')
    monkeypatch.setattr(catalog.Path, 'is_dir', lambda self: (_ for _ in ()).throw(
        AssertionError('server inspected personal directory')))
    submitted = {}

    def enqueue(_session, **kwargs):
        submitted.update(kwargs)
        raise ConflictError('execution capability is not connected')

    monkeypatch.setattr(catalog, 'enqueue_job', enqueue)
    request = SelfListingPipelineCreate(
        selected_steps=['watermark'],
        product_root=r'D:\My Pictures\Products',
        output_dir=r'D:\My Pictures\Output',
        watermark_root=r'D:\My Pictures\Watermark',
    )
    with SessionLocal() as session, pytest.raises(ConflictError):
        catalog.start_self_listing_pipeline(request, session)
    assert submitted['job_type'] == 'self_listing_pipeline'
    assert submitted['payload']['product_root'] == request.product_root
    assert submitted['payload']['output_dir'] == request.output_dir
    assert submitted['max_attempts'] == 3


def test_cloud_rejects_relative_personal_directory(monkeypatch):
    monkeypatch.setattr(get_settings(), 'multiuser_enabled', True)
    with SessionLocal() as session, pytest.raises(ConflictError):
        catalog.start_self_listing_pipeline(
            SelfListingPipelineCreate(selected_steps=['watermark'], product_root='products'), session)


def test_original_watermark_step_runs_with_personal_pc_context(tmp_path):
    from erp.services.self_listing_pipeline import run_self_listing_pipeline

    product_root = tmp_path / 'products'
    image_dir = product_root / 'item-one' / '_generated'
    image_dir.mkdir(parents=True)
    Image.new('RGB', (120, 120), (60, 100, 160)).save(image_dir / 'one.jpg')
    Image.new('RGBA', (20, 20), (255, 255, 255, 140)).save(product_root / '水印.png')
    step = SimpleNamespace(key='watermark', status='queued')

    class Context:
        payload = {'product_root': str(product_root), 'output_dir': str(tmp_path / 'out'),
                   'watermark_root': str(tmp_path / 'watermarked'),
                   'selected_steps': ['watermark']}
        job = SimpleNamespace(id='a' * 32, steps=[step])

        def start_step(self, key):
            assert key == 'watermark'
            step.status = 'running'

        def complete_step(self, key, *, details=None):
            assert key == 'watermark'
            step.status = 'completed'

        def refresh(self):
            return None

        def event(self, *_args, **_kwargs):
            return None

    run_self_listing_pipeline(Context(), snapshot={'connections': {}, 'secrets': {}},
                              output_asset=lambda *_args: None)
    assert step.status == 'completed'
    assert (tmp_path / 'watermarked' / 'item-one' / '_generated' / 'one.jpg').is_file()


def test_watermark_pipeline_uses_personal_lease_and_returns_shared_status(monkeypatch, tmp_path):
    monkeypatch.setattr(get_settings(), 'multiuser_enabled', True)
    monkeypatch.setattr(get_settings(), 'runtime_mode', 'server')
    monkeypatch.setattr(get_settings(), 'session_cookie_secure', False)
    with SessionLocal() as session:
        create_user(session, username='pipeline-owner', display_name='pipeline-owner',
                    password='isolated-pipeline-password', role='operator')
    product = tmp_path / 'products'
    pictures = product / 'item-one' / '_generated'
    pictures.mkdir(parents=True)
    Image.new('RGB', (120, 120), (60, 100, 160)).save(pictures / 'one.jpg')
    Image.new('RGBA', (20, 20), (255, 255, 255, 140)).save(product / '水印.png')
    worker_id = 'pc-' + 'd' * 32
    base = '/api/v1/execution-devices/local/' + worker_id
    with TestClient(app) as browser, TestClient(app) as worker:
        for client in (browser, worker):
            response = client.post('/api/v1/auth/login', json={
                'username': 'pipeline-owner', 'password': 'isolated-pipeline-password'})
            assert response.status_code == 200, response.text
        response = worker.post('/api/v1/execution-devices/local/register', json={
            'id': worker_id, 'name': 'personal pipeline PC', 'self_listing_version': 1})
        assert response.status_code == 200, response.text
        assert browser.post('/api/v1/execution-devices/select', json={
            'worker_id': worker_id}).status_code == 200
        queued = browser.post('/api/v1/catalog/self-listing/pipeline', json={
            'selected_steps': ['watermark'], 'product_root': str(product),
            'output_dir': str(tmp_path / 'output'),
            'watermark_root': str(tmp_path / 'watermarked')})
        assert queued.status_code == 202, queued.text
        job = worker.post(base + '/claim', json={}).json()
        assert job['id'] == queued.json()['id']
        endpoint = '/execution-devices/local/' + worker_id + '/jobs/' + job['id']

        class Adapter:
            def request(self, method, path, body=None, *, extra_headers=None):
                response = worker.request(method, '/api/v1' + path, json=body,
                                          headers=extra_headers or {})
                assert response.status_code == 200, response.text
                return response.json(), response.headers

        monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1] / 'client' / 'desktop'))
        from self_listing_pipeline_worker import run
        result = run(Adapter(), endpoint, job, tmp_path / 'client', threading.Event())
        assert result == {'schema': 1, 'completed_steps': ['watermark'], 'artifacts': []}
        assert (tmp_path / 'watermarked' / 'item-one' / '_generated' / 'one.jpg').is_file()
        completed = worker.post('/api/v1' + endpoint + '/complete', json={'result': result},
                                headers={'X-Worker-Attempt': '1'})
        assert completed.status_code == 200, completed.text
        assert completed.json()['status'] == 'completed'


def test_personal_pipeline_table_is_leased_and_idempotent(monkeypatch, tmp_path):
    monkeypatch.setattr(get_settings(), 'multiuser_enabled', True)
    monkeypatch.setattr(get_settings(), 'runtime_mode', 'server')
    monkeypatch.setattr(get_settings(), 'session_cookie_secure', False)
    monkeypatch.setattr(get_settings(), 'data_dir', tmp_path / 'server-data')
    with SessionLocal() as session:
        create_user(session, username='pipeline-artifact-owner',
                    display_name='pipeline-artifact-owner',
                    password='isolated-pipeline-password', role='operator')
        connection = IntegrationConnection(provider='ai_chat', name='pipeline-test-chat',
                                           base_url='https://chat.invalid',
                                           secrets_encrypted=vault.encrypt({'api_key': 'isolated'}),
                                           enabled=True)
        session.add(connection)
        session.commit()
        connection_id = connection.id
    worker_id = 'pc-' + 'e' * 32
    base = '/api/v1/execution-devices/local/' + worker_id
    with TestClient(app) as browser, TestClient(app) as worker:
        for client in (browser, worker):
            assert client.post('/api/v1/auth/login', json={
                'username': 'pipeline-artifact-owner',
                'password': 'isolated-pipeline-password'}).status_code == 200
        assert worker.post('/api/v1/execution-devices/local/register', json={
            'id': worker_id, 'name': 'personal pipeline PC',
            'self_listing_version': 1}).status_code == 200
        assert browser.post('/api/v1/execution-devices/select', json={
            'worker_id': worker_id}).status_code == 200
        queued = browser.post('/api/v1/catalog/self-listing/pipeline', json={
            'selected_steps': ['generate_copywriting'],
            'product_root': str(tmp_path / 'personal-products'),
            'ai_chat_connection_id': connection_id})
        assert queued.status_code == 202, queued.text
        job = worker.post(base + '/claim', json={}).json()
        endpoint = base + '/jobs/' + job['id']
        headers = {'X-Worker-Attempt': '1'}
        snapshot = worker.get(endpoint + '/self-listing-config', headers=headers)
        assert snapshot.status_code == 200
        assert snapshot.json()['secrets']['ai_chat']['api_key'] == 'isolated'
        assert snapshot.headers['cache-control'] == 'no-store'
        book = Workbook()
        book.active.append(['SKU', 'title'])
        book.active.append(['TEST', 'isolated'])
        output = io.BytesIO()
        book.save(output)
        data = output.getvalue()
        digest = hashlib.sha256(data).hexdigest()
        upload = {'role': 'copywriting_xlsx', 'sha256': digest,
                  'content': base64.b64encode(data).decode('ascii')}
        first = worker.post(endpoint + '/self-listing-artifact', json=upload, headers=headers)
        assert first.status_code == 200, first.text
        second = worker.post(endpoint + '/self-listing-artifact', json=upload, headers=headers)
        assert second.status_code == 200 and second.json() == first.json()
        bad = worker.post(endpoint + '/self-listing-artifact', json={
            **upload, 'sha256': '0' * 64}, headers=headers)
        assert bad.status_code == 409
        result = {'schema': 1, 'completed_steps': ['generate_copywriting'],
                  'artifacts': [first.json()]}
        done = worker.post(endpoint + '/complete', json={'result': result}, headers=headers)
        assert done.status_code == 200, done.text
        assert done.json()['status'] == 'completed'


def test_all_five_original_steps_run_on_personal_computer_with_shared_tables(monkeypatch, tmp_path):
    from erp.domains.self_listing.vendor_sources import image_generation, copywriting, watermark, oss_uploader

    monkeypatch.setattr(get_settings(), 'multiuser_enabled', True)
    monkeypatch.setattr(get_settings(), 'runtime_mode', 'server')
    monkeypatch.setattr(get_settings(), 'session_cookie_secure', False)
    monkeypatch.setattr(get_settings(), 'data_dir', tmp_path / 'cloud')
    with SessionLocal() as session:
        create_user(session, username='pipeline-five-owner',
                    display_name='pipeline-five-owner',
                    password='isolated-pipeline-password', role='operator')
        for provider in ('ai_chat', 'ai_image', 'oss'):
            session.add(IntegrationConnection(provider=provider,
                name='pipeline-five-' + provider, base_url='https://example.invalid',
                config_json='{}', secrets_encrypted=vault.encrypt({
                    'api_key': 'isolated', 'access_key_id': 'isolated',
                    'access_key_secret': 'isolated'}), enabled=True))
        session.commit()
        ids = {row.provider: row.id for row in session.query(IntegrationConnection).all()}
    product = tmp_path / 'products'
    product.mkdir()
    Image.new('RGBA', (20, 20), (255, 255, 255, 140)).save(product / '水印.png')
    calls = []

    def fake_images(config, event, stop_event):
        calls.append('generate_images')
        target = product / 'item-one' / '_generated'
        target.mkdir(parents=True)
        Image.new('RGB', (120, 120), (60, 100, 160)).save(target / 'one.jpg')

    def fake_copywriting(root_dir, config):
        calls.append('generate_copywriting')
        path = Path(config['output_xlsx'])
        book = Workbook()
        book.active.append(['SKU', 'error'])
        book.active.append(['TEST', ''])
        book.save(path)
        return str(path)

    original_watermark = watermark.process_watermark_directory

    def record_watermark(*args, **kwargs):
        calls.append('watermark')
        return original_watermark(*args, **kwargs)

    def fake_upload(root_dir, *, cfg, output_path, dry_run, log_func):
        calls.append('upload_original' if Path(root_dir) == product else 'upload_watermarked')
        book = Workbook()
        book.active.append(['file', 'oss_url'])
        book.active.append(['one.jpg', 'https://example.invalid/one.jpg'])
        book.save(output_path)

    monkeypatch.setattr(image_generation, 'run_generation', fake_images)
    monkeypatch.setattr(copywriting, 'run_copywriting', fake_copywriting)
    monkeypatch.setattr(watermark, 'process_watermark_directory', record_watermark)
    monkeypatch.setattr(oss_uploader, 'batch_upload_images', fake_upload)
    worker_id = 'pc-' + 'f' * 32
    base = '/api/v1/execution-devices/local/' + worker_id
    with TestClient(app) as browser, TestClient(app) as worker:
        for client in (browser, worker):
            assert client.post('/api/v1/auth/login', json={
                'username': 'pipeline-five-owner',
                'password': 'isolated-pipeline-password'}).status_code == 200
        assert worker.post('/api/v1/execution-devices/local/register', json={
            'id': worker_id, 'name': 'personal five-step PC',
            'self_listing_version': 1}).status_code == 200
        assert browser.post('/api/v1/execution-devices/select', json={
            'worker_id': worker_id}).status_code == 200
        queued = browser.post('/api/v1/catalog/self-listing/pipeline', json={
            'selected_steps': list(('generate_images', 'generate_copywriting', 'watermark',
                                    'upload_original', 'upload_watermarked')),
            'product_root': str(product),
            'output_dir': str(tmp_path / 'tables'),
            'watermark_root': str(tmp_path / 'watermarked'),
            'ai_chat_connection_id': ids['ai_chat'],
            'ai_image_connection_id': ids['ai_image'],
            'oss_connection_id': ids['oss'], 'oss_dry_run': True})
        assert queued.status_code == 202, queued.text
        job = worker.post(base + '/claim', json={}).json()
        endpoint = '/execution-devices/local/' + worker_id + '/jobs/' + job['id']

        class Adapter:
            def request(self, method, path, body=None, *, extra_headers=None):
                response = worker.request(method, '/api/v1' + path, json=body,
                                          headers=extra_headers or {})
                assert response.status_code == 200, response.text
                return response.json(), response.headers

        monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1] / 'client' / 'desktop'))
        from self_listing_pipeline_worker import run
        result = run(Adapter(), endpoint, job, tmp_path / 'client', threading.Event())
        assert calls == ['generate_images', 'generate_copywriting', 'watermark',
                         'upload_original', 'upload_watermarked']
        assert len(result['artifacts']) == 3
        assert (tmp_path / 'watermarked' / 'item-one' / '_generated' / 'one.jpg').is_file()
        done = worker.post('/api/v1' + endpoint + '/complete', json={'result': result},
                           headers={'X-Worker-Attempt': '1'})
        assert done.status_code == 200, done.text
        assert done.json()['status'] == 'completed'
