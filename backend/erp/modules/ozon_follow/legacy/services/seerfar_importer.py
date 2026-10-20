from __future__ import annotations

import math
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from erp.modules.ozon_follow.legacy.models import ProductCandidate


class SeerfarImportError(ValueError):
    pass


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    cleaned = re.sub(r"[^0-9,\.\-]", "", str(value)).replace(",", "")
    if cleaned in ("", "-", "."):
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def _integer(value: Any) -> int | None:
    number = _number(value)
    return None if number is None else int(round(number))


def _percent(value: Any) -> float | None:
    number = _number(value)
    if number is None:
        return None
    return number / 100 if "%" in str(value) else number


def _weight_grams(value: Any) -> float | None:
    number = _number(value)
    if number is None:
        return None
    text = str(value).lower()
    if "kg" in text or "кг" in text or "千克" in text:
        return number * 1000
    return number


def _unique_headers(values: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for index, value in enumerate(values, start=1):
        header = _text(value) or f"未命名列{index}"
        seen[header] = seen.get(header, 0) + 1
        result.append(header if seen[header] == 1 else f"{header}#{seen[header]}")
    return result


def _first(row: dict[str, Any], *names: str) -> Any:
    for name in names:
        value = row.get(name)
        if value not in (None, ""):
            return value
    return None


def _image_urls(row: dict[str, Any]) -> list[str]:
    """Collect every Ozon image column (主图/附图/图片), preserving order."""
    urls: list[str] = []
    for key, value in row.items():
        label = str(key).casefold()
        if not any(token in label for token in ("主图", "附图", "图片", "image", "photo")):
            continue
        for part in re.split(r"[\r\n,;|\s]+", _text(value)):
            if part.startswith(("http://", "https://")) and part not in urls:
                urls.append(part)
    return urls


def import_seerfar_workbook(path: str | Path) -> list[ProductCandidate]:
    source = Path(path)
    if not source.exists():
        raise SeerfarImportError(f"文件不存在：{source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise SeerfarImportError("目前仅支持 Seerfar 导出的 .xlsx 或 .xlsm 文件")

    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as exc:
        raise SeerfarImportError(f"无法读取 Excel：{exc}") from exc

    try:
        sheet = workbook["Data"] if "Data" in workbook.sheetnames else workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = _unique_headers(list(next(iterator)))
        except StopIteration as exc:
            raise SeerfarImportError("Excel 中没有数据") from exc

        required = {"标题", "详情页地址", "SKU"}
        missing = sorted(required.difference(headers))
        if missing:
            raise SeerfarImportError("缺少必要字段：" + "、".join(missing))

        products: list[ProductCandidate] = []
        seen_skus: set[str] = set()
        for row_number, values in enumerate(iterator, start=2):
            row = dict(zip(headers, values, strict=False))
            sku = _identifier(row.get("SKU"))
            title = _text(row.get("标题"))
            if not sku and not title:
                continue
            if not sku:
                raise SeerfarImportError(f"第 {row_number} 行缺少 SKU")
            if sku in seen_skus:
                continue
            seen_skus.add(sku)

            image_urls = _image_urls(row)
            image_url = image_urls[0] if image_urls else _text(_first(row, "主图#2", "主图URL", "图片"))
            listed_date = _text(row.get("上架时间")).splitlines()[0]
            products.append(
                ProductCandidate(
                    source_file=str(source),
                    source_row=row_number,
                    rank=_integer(row.get("排名")),
                    sku=sku,
                    title=title,
                    ozon_url=_text(row.get("详情页地址")),
                    image_url=image_url,
                    image_urls=image_urls,
                    brand=_text(row.get("品牌")),
                    category=_text(row.get("类目")),
                    sale_method=_text(row.get("销售方式")),
                    selling_price_rub=_number(row.get("售价")),
                    sales=_integer(row.get("销量")),
                    revenue_rub=_number(row.get("销售额")),
                    sales_growth=_percent(row.get("销量增长率")),
                    revenue_growth=_percent(row.get("销售额增长率")),
                    gross_margin=_percent(row.get("毛利率")),
                    rating=_number(row.get("评分")),
                    review_count=_integer(row.get("评论数")),
                    seller=_text(row.get("店铺")),
                    seller_type=_text(row.get("卖家类型")),
                    delivery_method=_text(row.get("配送方式")),
                    weight_grams=_weight_grams(row.get("重量")),
                    listed_date=listed_date,
                    raw_data={
                        **{key: _text(value) for key, value in row.items()},
                        "_ozon_image_urls": image_urls,
                    },
                )
            )
        return products
    finally:
        workbook.close()
