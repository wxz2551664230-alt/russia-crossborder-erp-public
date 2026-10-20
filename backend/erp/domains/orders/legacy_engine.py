from __future__ import annotations

import asyncio
import hashlib
import json
import sqlite3
import threading
import time
import uuid
from contextlib import closing
from datetime import UTC, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any, Callable

import httpx
from cryptography.fernet import Fernet, InvalidToken
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

from erp.domains.orders.products import ProductLibrary, ProductLibraryError


def utc_now() -> str:
    return datetime.now(UTC).isoformat()


WB_CANCELLED_STATUSES = frozenset({
    "declined_by_client", "canceled_by_client", "cancelled_by_client",
    "canceled", "cancelled", "canceled_by_carrier", "defect",
})


def is_platform_order_cancelled(order: dict[str, Any] | sqlite3.Row) -> bool:
    statuses = {str(order[key] or "").strip().casefold() for key in ("order_status", "substatus")}
    if str(order["platform"] or "").casefold() == "wb":
        return bool(statuses & WB_CANCELLED_STATUSES)
    return bool(statuses & {"canceled", "cancelled"})


_CANCELLED_SQL = "(LOWER(COALESCE(o.order_status,'')) IN ('canceled','cancelled') OR (o.platform = 'wb' AND LOWER(COALESCE(o.substatus,'')) IN ('declined_by_client','canceled_by_client','cancelled_by_client','canceled','cancelled','canceled_by_carrier','defect')))"


def official_supplier_tracking_number(supplier: str, value: Any) -> str:
    """Return only a carrier-issued tracking number suitable for downstream use.

    S2B exposes an internal temporary logistics reference before YTO assigns the
    real waybill.  Only a YT-prefixed value longer than six characters is the
    formal domestic tracking number.
    """
    tracking = str(value or "").strip()
    if str(supplier or "").casefold() != "s2b":
        return tracking
    normalized = tracking.upper()
    return normalized if normalized.startswith("YT") and len(normalized) > 6 else ""


def hicustom_datetime(value: Any) -> str:
    text = str(value or "").strip()
    if len(text) >= 19 and text[4:5] == "-" and text[7:8] == "-":
        return text[:19].replace("T", " ")
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def s2b_country_code(value: Any) -> str:
    """S2B requires CN to expose mainland carrier quotations."""
    country = str(value or "").strip()
    return "CN" if country.casefold() in {"cn", "china", "中国", "中华人民共和国"} else country


class OrderError(RuntimeError):
    def __init__(
        self,
        message: str,
        *,
        status_code: int = 422,
        outcome_unknown: bool = False,
    ) -> None:
        super().__init__(message)
        self.status_code = status_code
        self.outcome_unknown = outcome_unknown


class OzonClient:
    """Small OZON Seller API client for rFBS/FBS posting synchronization."""

    def __init__(self, settings: dict[str, Any]) -> None:
        self.client_id = str(settings.get("ozon_client_id") or "").strip()
        self.api_key = str(settings.get("ozon_api_key") or "").strip()
        if not self.client_id or not self.api_key:
            raise OrderError("请先在设置中填写 OZON Client ID 和 API Key", status_code=503)
        self.http = httpx.AsyncClient(
            base_url="https://api-seller.ozon.ru",
            headers={"Client-Id": self.client_id, "Api-Key": self.api_key},
            timeout=90,
            follow_redirects=True,
        )
        self._order_read_lock = asyncio.Lock()
        self._next_order_read = 0.0

    async def aclose(self) -> None:
        await self.http.aclose()

    async def _post(self, path: str, payload: dict[str, Any]) -> dict[str, Any]:
        # Retry only these read-only endpoints. Never replay platform writes.
        if path in {"/v3/posting/fbs/list", "/v3/posting/fbs/get"}:
            async with self._order_read_lock:
                for attempt in range(4):
                    await asyncio.sleep(max(0.0, self._next_order_read - time.monotonic()))
                    try:
                        return await self._post_once(path, payload)
                    except OrderError as error:
                        if error.status_code != 429:
                            raise
                        delay = max(2 ** (attempt + 1), getattr(error, "retry_after", 0))
                        self._next_order_read = time.monotonic() + delay
                        if attempt == 3 or delay > 60:
                            raise OrderError("OZON 订单读取触发限流，请稍后继续原任务", status_code=429) from None
                    finally:
                        self._next_order_read = max(self._next_order_read, time.monotonic() + 1.1)
        return await self._post_once(path, payload)

    async def _post_once(self, path: str, payload: dict[str, Any]) -> dict[str, Any]:
        try:
            response = await self.http.post(path, json=payload)
        except (httpx.TimeoutException, httpx.NetworkError) as error:
            raise OrderError(f"OZON 网络请求失败：{error}", status_code=502) from error
        if response.status_code == 429:
            error = OrderError("OZON 订单接口请求频率超限", status_code=429)
            retry_after = response.headers.get("Retry-After", "")
            try:
                error.retry_after = max(0, int(retry_after))
            except (TypeError, ValueError):
                try:
                    error.retry_after = max(0, (parsedate_to_datetime(retry_after) - datetime.now(UTC)).total_seconds())
                except (TypeError, ValueError, OverflowError):
                    error.retry_after = 0
            raise error
        try:
            data = response.json()
        except ValueError as error:
            raise OrderError("OZON 返回了无效响应", status_code=502) from error
        if response.status_code < 200 or response.status_code >= 300:
            message = ""
            if isinstance(data, dict):
                message = str(data.get("message") or data.get("error") or "")
            raise OrderError(
                f"OZON 接口 HTTP {response.status_code}{f'：{message}' if message else ''}",
                status_code=response.status_code,
            )
        if not isinstance(data, dict):
            raise OrderError("OZON 响应格式错误", status_code=502)
        return data

    async def test_connection(self) -> dict[str, Any]:
        payload = await self._post(
            "/v3/posting/fbs/list",
            {
                "dir": "DESC",
                "filter": {
                    "since": (datetime.now(UTC) - timedelta(days=1)).isoformat(),
                    "to": datetime.now(UTC).isoformat(),
                    "status": "awaiting_packaging",
                },
                "limit": 1,
                "offset": 0,
                "with": {"analytics_data": False, "barcodes": False, "financial_data": False},
            },
        )
        result = payload.get("result") if isinstance(payload.get("result"), dict) else {}
        return {
            "connected": True,
            "client_id": self.client_id,
            "visible_postings": len(result.get("postings") or []),
        }

    async def posting(self, posting_number: str) -> dict[str, Any]:
        payload = await self._post(
            "/v3/posting/fbs/get",
            {
                "posting_number": posting_number,
                "with": {
                    "analytics_data": True,
                    "barcodes": True,
                    "financial_data": True,
                    "product_exemplars": True,
                    "related_postings": True,
                    "translit": True,
                },
            },
        )
        result = payload.get("result")
        return result if isinstance(result, dict) else {}

    async def postings(self, *, days: int = 30, status: str = "awaiting_packaging", date_from: str | None = None, date_to: str | None = None, order_numbers: list[str] | None = None) -> list[dict[str, Any]]:
        if order_numbers:
            values = []
            for number in dict.fromkeys(order_numbers):
                detail = await self.posting(number)
                if not detail or detail.get("posting_number") != number:
                    raise OrderError("指定 OZON 发货单号未返回匹配订单", status_code=404)
                values.append(detail)
            return values
        now = datetime.now(UTC)
        offset = 0
        values: list[dict[str, Any]] = []
        while offset < 5000:
            payload = await self._post(
                "/v3/posting/fbs/list",
                {
                    "dir": "ASC",
                    "filter": {
                        "since": date_from or (now - timedelta(days=max(1, min(365, days)))).isoformat(),
                        "to": date_to or now.isoformat(),
                        "status": status,
                    },
                    "limit": 1000,
                    "offset": offset,
                    "with": {
                        "analytics_data": True,
                        "barcodes": True,
                        "financial_data": True,
                    },
                },
            )
            result = payload.get("result") if isinstance(payload.get("result"), dict) else {}
            postings = [item for item in result.get("postings") or [] if isinstance(item, dict)]
            values.extend(postings)
            has_next = result.get("has_next", len(postings) >= 1000)
            if not has_next:
                break
            if not postings:
                raise OrderError("OZON 分页未前进，未回传不完整订单", status_code=502)
            offset += len(postings)
        else:
            raise OrderError("OZON 查询结果超过单次上限，请缩小时间范围", status_code=422)
        semaphore = asyncio.Semaphore(5)

        async def enrich(posting: dict[str, Any]) -> dict[str, Any]:
            number = str(posting.get("posting_number") or "").strip()
            if not number:
                return posting
            async with semaphore:
                try:
                    detail = await self.posting(number)
                except OrderError as error:
                    if error.status_code == 429:
                        raise
                    return posting
            return detail or posting

        return await asyncio.gather(*(enrich(posting) for posting in values))

    async def finance_transactions(
        self, *, days: int = 30, date_from: str = "", date_to: str = "",
    ) -> list[dict[str, Any]]:
        """Return OZON finance operations, splitting long ranges into API-safe windows."""
        now = datetime.now(UTC)
        try:
            start = datetime.fromisoformat(date_from).replace(tzinfo=UTC) if date_from else now - timedelta(days=max(1, min(31, days)))
            end = datetime.fromisoformat(date_to).replace(tzinfo=UTC) + timedelta(days=1) - timedelta(microseconds=1) if date_to else now
        except ValueError as error:
            raise OrderError("财务同步日期格式错误") from error
        end = min(end, now)
        if start > end:
            raise OrderError("财务同步开始日期不能晚于结束日期或今天")
        values: list[dict[str, Any]] = []
        window_start = start
        while window_start <= end:
            window_end = min(window_start + timedelta(days=30) - timedelta(microseconds=1), end)
            page = 1
            while page <= 1000:
                payload = await self._post(
                    "/v3/finance/transaction/list",
                    {
                        "filter": {
                            "date": {"from": window_start.isoformat(), "to": window_end.isoformat()},
                            "operation_type": [],
                            "posting_number": "",
                            "transaction_type": "all",
                        },
                        "page": page,
                        "page_size": 1000,
                    },
                )
                result = payload.get("result") if isinstance(payload.get("result"), dict) else {}
                operations = [item for item in result.get("operations") or [] if isinstance(item, dict)]
                values.extend(operations)
                page_count = int(result.get("page_count") or 0)
                if not operations or page >= page_count:
                    break
                page += 1
            window_start = window_end + timedelta(microseconds=1)
        return values


class OzonPerformanceClient:
    def __init__(self, settings: dict[str, Any]) -> None:
        self.client_id = str(settings.get("performance_client_id") or "").strip()
        self.client_secret = str(settings.get("performance_token") or "").strip()
        if not self.client_id or not self.client_secret:
            raise OrderError("请先为该 OZON 店铺配置广告 Client ID 和 Client Secret", status_code=503)
        self.http = httpx.AsyncClient(base_url="https://api-performance.ozon.ru", timeout=90)
        self.access_token = ""

    async def aclose(self) -> None:
        await self.http.aclose()

    async def authenticate(self) -> None:
        response = await self.http.post("/api/client/token", json={
            "client_id": self.client_id, "client_secret": self.client_secret, "grant_type": "client_credentials"
        })
        data = response.json()
        if response.status_code != 200 or not data.get("access_token"):
            raise OrderError("OZON 广告 API 授权失败，请检查广告 Client ID 和 Client Secret", status_code=502)
        self.access_token = str(data["access_token"])

    async def _get(self, path: str, params: Any = None) -> dict[str, Any]:
        if not self.access_token:
            await self.authenticate()
        response = await self.http.get(path, params=params, headers={"Authorization": f"Bearer {self.access_token}"})
        if response.status_code == 401:
            self.access_token = ""
            await self.authenticate()
            response = await self.http.get(path, params=params, headers={"Authorization": f"Bearer {self.access_token}"})
        try:
            data = response.json()
        except ValueError as error:
            raise OrderError("OZON 广告 API 返回了无效响应", status_code=502) from error
        if response.status_code < 200 or response.status_code >= 300:
            message = ""
            if isinstance(data, dict):
                message = str(data.get("message") or data.get("error_description") or data.get("error") or data.get("detail") or "")
            raise OrderError(
                f"OZON 广告 API HTTP {response.status_code}{f'：{message[:500]}' if message else ''}",
                status_code=response.status_code,
            )
        return data if isinstance(data, dict) else {}

    async def test_connection(self) -> dict[str, Any]:
        data = await self._get("/api/client/campaign", {"page": 1, "pageSize": 1})
        return {"connected": True, "campaigns_visible": len(data.get("list") or [])}

    @staticmethod
    def _number(value: Any) -> Decimal:
        text = str(value or "0").replace(" ", "").replace(",", ".")
        try:
            return Decimal(text)
        except InvalidOperation:
            return Decimal("0")

    async def expense(self, date_from: str, date_to: str) -> dict[str, Any]:
        today = datetime.now(timezone(timedelta(hours=8))).date().isoformat()
        if date_from > today:
            raise OrderError("广告查询开始日期不能晚于今天")
        date_to = min(date_to, today)
        if date_from > date_to:
            raise OrderError("广告查询开始日期不能晚于结束日期")
        campaigns: list[dict[str, Any]] = []
        page = 1
        while page <= 100:
            data = await self._get("/api/client/campaign", {"page": page, "pageSize": 100})
            batch = [item for item in data.get("list") or [] if isinstance(item, dict)]
            campaigns.extend(batch)
            if len(batch) < 100:
                break
            page += 1
        ids = [str(item.get("id")) for item in campaigns if item.get("id")]
        rows: list[dict[str, Any]] = []
        for offset in range(0, len(ids), 10):
            params: list[tuple[str, str]] = [("dateFrom", date_from), ("dateTo", date_to)]
            params.extend(("campaignIds", value) for value in ids[offset:offset + 10])
            data = await self._get("/api/client/statistics/campaign/product/json", params)
            rows.extend(item for item in data.get("rows") or [] if isinstance(item, dict))
        spent = sum((self._number(item.get("moneySpent")) for item in rows), Decimal("0"))
        sales = sum((self._number(item.get("ordersMoney")) for item in rows), Decimal("0"))
        return {"campaign_count": len(campaigns), "row_count": len(rows), "amount": format(spent, ".2f"),
                "attributed_sales": format(sales, ".2f"), "rows": rows}


class TmsClient:
    def __init__(self, settings: dict[str, Any]) -> None:
        self.base_url = str(settings.get("tms_base_url") or "").rstrip("/")
        self.app_key = str(settings.get("tms_app_key") or "").strip()
        self.app_secret = str(settings.get("tms_app_secret") or "").strip()
        if not self.base_url or not self.app_key or not self.app_secret:
            raise OrderError("请先配置代打包系统 API 地址、AppKey 和 AppSecret", status_code=503)
        self.http = httpx.AsyncClient(base_url=self.base_url, timeout=90, follow_redirects=True)

    async def aclose(self) -> None:
        await self.http.aclose()

    def signed_request(self, payload: dict[str, Any]) -> tuple[bytes, dict[str, str]]:
        body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        timestamp = str(int(time.time()))
        body_md5 = hashlib.md5(body).hexdigest()
        signature = hashlib.md5(
            f"{body_md5}{timestamp}{self.app_secret}".encode("utf-8")
        ).hexdigest()
        return body, {
            "AppKey": self.app_key,
            "Signature": signature,
            "RequestTimestamp": timestamp,
            "Content-Type": "application/json",
        }

    async def post(self, path: str, payload: dict[str, Any]) -> dict[str, Any]:
        body, headers = self.signed_request(payload)
        try:
            response = await self.http.post(path, content=body, headers=headers)
        except (httpx.TimeoutException, httpx.NetworkError) as error:
            raise OrderError(f"代打包系统网络请求失败：{error}", status_code=502) from error
        try:
            data = response.json()
        except ValueError as error:
            raise OrderError("代打包系统返回了无效响应", status_code=502) from error
        if response.status_code < 200 or response.status_code >= 300:
            raise OrderError(
                f"代打包系统 HTTP {response.status_code}：{data.get('message') or '请求失败'}",
                status_code=response.status_code,
            )
        if not isinstance(data, dict) or str(data.get("result") or "") != "success":
            raise OrderError(str(data.get("message") or data.get("code") or "代打包系统请求失败"))
        return data

    async def find_wait_submit_order(self, platform_order_sn: str) -> dict[str, Any] | None:
        for page in range(1, 11):
            data = await self.post(
                "/customer/order/fetch/search_order_page",
                {"tab": "wait_submit", "page": page, "pageSize": 100},
            )
            payload = data.get("data") if isinstance(data.get("data"), dict) else {}
            orders = [item for item in payload.get("orderList") or [] if isinstance(item, dict)]
            for order in orders:
                if str(order.get("platformOrderSn") or "") == str(platform_order_sn):
                    return order
            if len(orders) < 100:
                break
        return None

    async def submit_packaging(self, payload: dict[str, Any]) -> dict[str, Any]:
        return await self.post("/customer/order/operate/submit_platform_order_package", payload)

    async def apply_waybills(self, package_ids: list[str]) -> dict[str, Any]:
        return await self.post(
            "/customer/order/operate/batch_apply_package_tracking_no",
            {"platformOrderPackageIds": package_ids},
        )


class WildberriesOrderClient:
    """Client for WB FBS assembly orders using the Marketplace API."""

    def __init__(self, settings: dict[str, Any]) -> None:
        token = str(settings.get("wildberries_token") or settings.get("token") or "").strip()
        if not token:
            raise OrderError("请先为 WB 店铺填写 API Token", status_code=503)
        headers = {"Authorization": token}
        self.http = httpx.AsyncClient(
            base_url=str(
                settings.get("wildberries_marketplace_api")
                or "https://marketplace-api.wildberries.ru"
            ).rstrip("/"),
            headers=headers,
            timeout=90,
            follow_redirects=True,
        )

    async def aclose(self) -> None:
        await self.http.aclose()

    async def _request(
        self,
        method: str,
        path: str,
        *,
        params: dict[str, Any] | None = None,
        json_body: Any = None,
    ) -> Any:
        try:
            response = await self.http.request(method, path, params=params, json=json_body)
        except (httpx.TimeoutException, httpx.NetworkError) as error:
            raise OrderError(f"WB 网络请求失败：{error}", status_code=502) from error
        if response.status_code == 204:
            return None
        try:
            payload = response.json()
        except ValueError as error:
            raise OrderError("WB 返回了无效响应", status_code=502) from error
        if response.status_code < 200 or response.status_code >= 300:
            message = ""
            if isinstance(payload, dict):
                message = str(payload.get("message") or payload.get("detail") or payload.get("code") or "")
            raise OrderError(
                f"WB 接口 HTTP {response.status_code}{f'：{message}' if message else ''}",
                status_code=response.status_code,
            )
        return payload

    async def test_connection(self) -> dict[str, Any]:
        payload = await self._request("GET", "/api/v3/orders/new")
        orders = payload.get("orders") if isinstance(payload, dict) else []
        return {"connected": True, "visible_orders": len(orders or [])}

    async def assembly_orders(self, *, days: int = 30, date_from: str | None = None, date_to: str | None = None, order_numbers: list[str] | None = None) -> list[dict[str, Any]]:
        now = datetime.now(UTC)
        scoped = bool(date_from or date_to or order_numbers)
        start = datetime.fromisoformat(date_from) if date_from else now - timedelta(days=max(1, min(30, days)))
        end = datetime.fromisoformat(date_to) if date_to else now
        requested = set(order_numbers or [])
        values: dict[str, dict[str, Any]] = {}
        cursor = 0
        for _ in range(100):
            payload = await self._request(
                "GET",
                "/api/v3/orders",
                params={"limit": 1000, "next": cursor, "dateFrom": int(start.timestamp()), "dateTo": int(end.timestamp())},
            )
            page = payload.get("orders") if isinstance(payload, dict) else []
            for item in page or []:
                if isinstance(item, dict) and item.get("id") is not None:
                    if not requested or str(item["id"]) in requested:
                        values[str(item["id"])] = dict(item)
            if requested and requested <= values.keys():
                break
            next_cursor = int(payload.get("next") or 0) if isinstance(payload, dict) else 0
            if not next_cursor:
                break
            if not page or next_cursor == cursor:
                raise OrderError("WB 分页未前进，未回传不完整订单", status_code=502)
            cursor = next_cursor
        else:
            raise OrderError("WB 查询结果超过单次上限，请缩小时间范围", status_code=422)
        if requested - values.keys():
            raise OrderError("部分 WB 订单 ID 在所选时间范围内未找到，请核对编号或扩大范围", status_code=404)
        newest = {} if scoped else await self._request("GET", "/api/v3/orders/new")
        for item in (newest.get("orders") if isinstance(newest, dict) else []) or []:
            if isinstance(item, dict) and item.get("id") is not None:
                values[str(item["id"])] = dict(item)
        ids = [int(value) for value in values if value.isdigit()]
        for offset in range(0, len(ids), 1000):
            statuses = await self._request(
                "POST", "/api/v3/orders/status", json_body={"orders": ids[offset : offset + 1000]}
            )
            for status in (statuses.get("orders") if isinstance(statuses, dict) else []) or []:
                if isinstance(status, dict):
                    target = values.get(str(status.get("id") or ""))
                    if target is not None:
                        target.update(status)
        return list(values.values())

    async def orders(self, *, days: int = 30, **scope) -> list[dict[str, Any]]:
        """Return WB FBS assembly orders that the seller can process."""
        return await self.assembly_orders(days=days, **scope)


class OrderManager:
    REQUIRED_ADDRESS_FIELDS = (
        "country",
        "region",
        "city",
        "area",
        "street",
        "postcode",
        "firstname",
        "mobile_phone",
    )

    def __init__(self, root: Path, library: ProductLibrary, cipher: Fernet) -> None:
        self.root = root
        self.path = root / "orders.db"
        self.library = library
        self.cipher = cipher
        self._lock = threading.RLock()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.path, timeout=30)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        return connection

    def initialize(self) -> None:
        self.root.mkdir(parents=True, exist_ok=True)
        with self._lock, closing(self._connect()) as connection:
            connection.executescript(
                """
                PRAGMA journal_mode = WAL;
                CREATE TABLE IF NOT EXISTS ozon_orders (
                    posting_number TEXT PRIMARY KEY,
                    order_status TEXT NOT NULL DEFAULT '',
                    substatus TEXT NOT NULL DEFAULT '',
                    fulfillment_mode TEXT NOT NULL DEFAULT 'rfbs',
                    in_process_at TEXT NOT NULL DEFAULT '',
                    shipment_date TEXT NOT NULL DEFAULT '',
                    delivery_method TEXT NOT NULL DEFAULT '',
                    recipient_enc BLOB NOT NULL,
                    match_status TEXT NOT NULL DEFAULT 'unmatched',
                    submission_status TEXT NOT NULL DEFAULT 'not_submitted',
                    hicustom_order_id TEXT NOT NULL DEFAULT '',
                    submission_error TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    synced_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS ozon_order_items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    posting_number TEXT NOT NULL REFERENCES ozon_orders(posting_number) ON DELETE CASCADE,
                    line_no INTEGER NOT NULL,
                    offer_id TEXT NOT NULL DEFAULT '',
                    ozon_sku TEXT NOT NULL DEFAULT '',
                    product_name TEXT NOT NULL DEFAULT '',
                    platform_image_url TEXT NOT NULL DEFAULT '',
                    quantity INTEGER NOT NULL DEFAULT 1,
                    price TEXT NOT NULL DEFAULT '',
                    matched_sku TEXT NOT NULL DEFAULT '',
                    product_code TEXT NOT NULL DEFAULT '',
                    stock_item_id INTEGER,
                    size_name TEXT NOT NULL DEFAULT '',
                    color_name TEXT NOT NULL DEFAULT '',
                    match_status TEXT NOT NULL DEFAULT 'unmatched',
                    match_error TEXT NOT NULL DEFAULT '',
                    UNIQUE(posting_number, line_no)
                );
                CREATE INDEX IF NOT EXISTS idx_ozon_orders_status ON ozon_orders(order_status);
                CREATE INDEX IF NOT EXISTS idx_ozon_orders_match ON ozon_orders(match_status);
                CREATE INDEX IF NOT EXISTS idx_ozon_orders_submission ON ozon_orders(submission_status);
                CREATE INDEX IF NOT EXISTS idx_ozon_items_offer ON ozon_order_items(offer_id);
                CREATE TABLE IF NOT EXISTS ozon_finance_operations (
                    store_id TEXT NOT NULL DEFAULT '',
                    operation_id TEXT NOT NULL,
                    posting_number TEXT NOT NULL DEFAULT '',
                    operation_date TEXT NOT NULL DEFAULT '',
                    operation_type TEXT NOT NULL DEFAULT '',
                    operation_type_name TEXT NOT NULL DEFAULT '',
                    transaction_type TEXT NOT NULL DEFAULT '',
                    amount TEXT NOT NULL DEFAULT '0',
                    accruals_for_sale TEXT NOT NULL DEFAULT '0',
                    sale_commission TEXT NOT NULL DEFAULT '0',
                    delivery_charge TEXT NOT NULL DEFAULT '0',
                    return_delivery_charge TEXT NOT NULL DEFAULT '0',
                    services_json TEXT NOT NULL DEFAULT '[]',
                    raw_json TEXT NOT NULL DEFAULT '{}',
                    rub_to_cny_rate TEXT NOT NULL DEFAULT '0.0783',
                    synced_at TEXT NOT NULL,
                    PRIMARY KEY(store_id, operation_id)
                );
                CREATE INDEX IF NOT EXISTS idx_ozon_finance_posting
                    ON ozon_finance_operations(store_id, posting_number);
                CREATE TABLE IF NOT EXISTS ozon_ad_costs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    store_id TEXT NOT NULL DEFAULT '',
                    posting_number TEXT NOT NULL DEFAULT '',
                    cost_date TEXT NOT NULL DEFAULT '',
                    sku TEXT NOT NULL DEFAULT '',
                    campaign_id TEXT NOT NULL DEFAULT '',
                    amount_rub TEXT NOT NULL DEFAULT '0',
                    attribution_method TEXT NOT NULL DEFAULT 'manual',
                    source_key TEXT NOT NULL,
                    raw_json TEXT NOT NULL DEFAULT '{}',
                    synced_at TEXT NOT NULL,
                    UNIQUE(store_id, source_key)
                );
                CREATE INDEX IF NOT EXISTS idx_ozon_ads_posting
                    ON ozon_ad_costs(store_id, posting_number);
                CREATE TABLE IF NOT EXISTS ozon_ad_reports (
                    store_id TEXT NOT NULL, date_from TEXT NOT NULL, date_to TEXT NOT NULL,
                    amount TEXT NOT NULL DEFAULT '0', attributed_sales TEXT NOT NULL DEFAULT '0',
                    campaign_count INTEGER NOT NULL DEFAULT 0, raw_json TEXT NOT NULL DEFAULT '{}', synced_at TEXT NOT NULL,
                    PRIMARY KEY(store_id, date_from, date_to)
                );
                CREATE TABLE IF NOT EXISTS profit_monthly_costs (
                    platform TEXT NOT NULL DEFAULT 'ozon',
                    store_id TEXT NOT NULL DEFAULT '',
                    month TEXT NOT NULL,
                    purchase_cost TEXT NOT NULL DEFAULT '0',
                    other_cost TEXT NOT NULL DEFAULT '0',
                    notes TEXT NOT NULL DEFAULT '',
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY(platform, store_id, month)
                );
                CREATE TABLE IF NOT EXISTS profit_sku_costs (
                    platform TEXT NOT NULL DEFAULT 'ozon',
                    store_id TEXT NOT NULL DEFAULT '',
                    sku TEXT NOT NULL,
                    unit_cost TEXT NOT NULL DEFAULT '0',
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY(platform, store_id, sku)
                );
                CREATE TABLE IF NOT EXISTS batch_order_groups (
                    id TEXT PRIMARY KEY,
                    batch_id TEXT NOT NULL,
                    sku TEXT NOT NULL,
                    product_name TEXT NOT NULL DEFAULT '',
                    product_code TEXT NOT NULL,
                    stock_item_id INTEGER NOT NULL,
                    size_name TEXT NOT NULL DEFAULT '',
                    color_name TEXT NOT NULL DEFAULT '',
                    quantity INTEGER NOT NULL,
                    sales_amount TEXT NOT NULL DEFAULT '0',
                    address_enc BLOB NOT NULL,
                    submission_status TEXT NOT NULL DEFAULT 'queued',
                    hicustom_order_id TEXT NOT NULL DEFAULT '',
                    submission_error TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS batch_order_sources (
                    group_id TEXT NOT NULL REFERENCES batch_order_groups(id) ON DELETE CASCADE,
                    posting_number TEXT NOT NULL REFERENCES ozon_orders(posting_number) ON DELETE CASCADE,
                    quantity INTEGER NOT NULL,
                    product_cost TEXT NOT NULL DEFAULT '0',
                    shipping_cost TEXT NOT NULL DEFAULT '0',
                    total_cost TEXT NOT NULL DEFAULT '0',
                    PRIMARY KEY(group_id, posting_number)
                );
                CREATE INDEX IF NOT EXISTS idx_batch_groups_status ON batch_order_groups(submission_status);
                CREATE INDEX IF NOT EXISTS idx_batch_sources_order ON batch_order_sources(posting_number);
                """
            )
            columns = {
                str(row["name"])
                for row in connection.execute("PRAGMA table_info(ozon_orders)").fetchall()
            }
            for name, definition in (
                ("platform", "TEXT NOT NULL DEFAULT 'ozon'"),
                ("store_id", "TEXT NOT NULL DEFAULT ''"),
                ("store_name", "TEXT NOT NULL DEFAULT ''"),
                ("external_order_id", "TEXT NOT NULL DEFAULT ''"),
                ("tms_status", "TEXT NOT NULL DEFAULT ''"),
                ("tms_platform_order_id", "TEXT NOT NULL DEFAULT ''"),
                ("tms_package_ids", "TEXT NOT NULL DEFAULT '[]'"),
                ("tms_item_ids", "TEXT NOT NULL DEFAULT '[]'"),
                ("tms_tail_tracking_numbers", "TEXT NOT NULL DEFAULT '[]'"),
                ("tms_error", "TEXT NOT NULL DEFAULT ''"),
                ("tms_submitted_at", "TEXT NOT NULL DEFAULT ''"),
                ("manual_tracking_number", "TEXT NOT NULL DEFAULT ''"),
                ("manual_tracking_company", "TEXT NOT NULL DEFAULT ''"),
                ("manual_tracking_at", "TEXT NOT NULL DEFAULT ''"),
                ("financial_data_json", "TEXT NOT NULL DEFAULT '{}'"),
            ):
                if name not in columns:
                    connection.execute(f"ALTER TABLE ozon_orders ADD COLUMN {name} {definition}")
            connection.execute(
                "UPDATE ozon_orders SET external_order_id = posting_number WHERE external_order_id = ''"
            )
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_orders_platform_store ON ozon_orders(platform, store_id)"
            )
            item_columns = {
                str(row["name"])
                for row in connection.execute("PRAGMA table_info(ozon_order_items)").fetchall()
            }
            for name, definition in (
                ("supplier", "TEXT NOT NULL DEFAULT 'hicustom'"),
                ("remote_product_id", "INTEGER NOT NULL DEFAULT 0"),
                ("color_id", "INTEGER NOT NULL DEFAULT 0"),
                ("size_id", "INTEGER NOT NULL DEFAULT 0"),
                ("weight", "TEXT NOT NULL DEFAULT ''"),
                ("length", "TEXT NOT NULL DEFAULT ''"),
                ("width", "TEXT NOT NULL DEFAULT ''"),
                ("height", "TEXT NOT NULL DEFAULT ''"),
                ("commission_percent", "TEXT NOT NULL DEFAULT '0'"),
                ("commission_amount", "TEXT NOT NULL DEFAULT '0'"),
                ("estimated_payout", "TEXT NOT NULL DEFAULT '0'"),
                ("financial_currency", "TEXT NOT NULL DEFAULT ''"),
                ("item_services_json", "TEXT NOT NULL DEFAULT '{}'"),
                ("platform_image_url", "TEXT NOT NULL DEFAULT ''"),
            ):
                if name not in item_columns:
                    connection.execute(f"ALTER TABLE ozon_order_items ADD COLUMN {name} {definition}")
            finance_columns = {str(row["name"]) for row in connection.execute("PRAGMA table_info(ozon_finance_operations)").fetchall()}
            if "rub_to_cny_rate" not in finance_columns:
                connection.execute("ALTER TABLE ozon_finance_operations ADD COLUMN rub_to_cny_rate TEXT NOT NULL DEFAULT '0.0783'")
            monthly_columns = {str(row["name"]) for row in connection.execute("PRAGMA table_info(profit_monthly_costs)").fetchall()}
            for name in ("domestic_shipping_cost", "packaging_cost", "labor_warehouse_cost", "tax_cost"):
                if name not in monthly_columns:
                    connection.execute(f"ALTER TABLE profit_monthly_costs ADD COLUMN {name} TEXT NOT NULL DEFAULT '0'")
            batch_columns = {
                str(row["name"])
                for row in connection.execute("PRAGMA table_info(batch_order_groups)").fetchall()
            }
            for name, definition in (
                ("supplier", "TEXT NOT NULL DEFAULT 'hicustom'"),
                ("remote_product_id", "INTEGER NOT NULL DEFAULT 0"),
                ("color_id", "INTEGER NOT NULL DEFAULT 0"),
                ("size_id", "INTEGER NOT NULL DEFAULT 0"),
                ("logistics_id", "INTEGER NOT NULL DEFAULT 0"),
                ("logistics_company_name", "TEXT NOT NULL DEFAULT ''"),
                ("quoted_shipping_amount", "TEXT NOT NULL DEFAULT '0'"),
                ("actual_product_amount", "TEXT NOT NULL DEFAULT '0'"),
                ("actual_shipping_amount", "TEXT NOT NULL DEFAULT '0'"),
                ("actual_grand_total", "TEXT NOT NULL DEFAULT '0'"),
                ("cost_recorded_at", "TEXT NOT NULL DEFAULT ''"),
                ("fulfillment_product_code", "TEXT NOT NULL DEFAULT ''"),
                ("fulfillment_gallery_codes", "TEXT NOT NULL DEFAULT '[]'"),
                ("product_revision", "INTEGER NOT NULL DEFAULT 0"),
                ("out_order_id", "TEXT NOT NULL DEFAULT ''"),
                ("supersedes_order_id", "TEXT NOT NULL DEFAULT ''"),
                ("product_prepared_at", "TEXT NOT NULL DEFAULT ''"),
                ("remote_status", "INTEGER"),
                ("paid_at", "TEXT NOT NULL DEFAULT ''"),
                ("tracking_number", "TEXT NOT NULL DEFAULT ''"),
                ("tracking_company", "TEXT NOT NULL DEFAULT ''"),
                ("remote_checked_at", "TEXT NOT NULL DEFAULT ''"),
                ("purchase_id", "TEXT NOT NULL DEFAULT ''"),
                ("merge_mode", "TEXT NOT NULL DEFAULT 'same_sku'"),
            ):
                if name not in batch_columns:
                    connection.execute(f"ALTER TABLE batch_order_groups ADD COLUMN {name} {definition}")
            source_columns = {
                str(row["name"])
                for row in connection.execute("PRAGMA table_info(batch_order_sources)").fetchall()
            }
            for name in ("product_cost", "shipping_cost", "total_cost"):
                if name not in source_columns:
                    connection.execute(
                        f"ALTER TABLE batch_order_sources ADD COLUMN {name} TEXT NOT NULL DEFAULT '0'"
                    )
            connection.commit()

    @staticmethod
    def _decimal(value: Any) -> Decimal:
        try:
            return Decimal(str(value or "0"))
        except (InvalidOperation, TypeError, ValueError):
            return Decimal("0")

    def import_finance_operations(self, operations: list[dict[str, Any]], *, store_id: str = "", rub_to_cny_rate: float = 0.0783) -> dict[str, int]:
        imported = 0
        linked = 0
        now = utc_now()
        with self._lock, closing(self._connect()) as connection:
            for operation in operations:
                operation_id = str(operation.get("operation_id") or "").strip()
                if not operation_id:
                    continue
                posting = operation.get("posting") if isinstance(operation.get("posting"), dict) else {}
                posting_number = str(posting.get("posting_number") or "").strip()
                existed = connection.execute(
                    "SELECT 1 FROM ozon_finance_operations WHERE store_id = ? AND operation_id = ?",
                    (store_id, operation_id),
                ).fetchone()
                connection.execute(
                    """
                    INSERT INTO ozon_finance_operations (
                        store_id, operation_id, posting_number, operation_date,
                        operation_type, operation_type_name, transaction_type, amount,
                        accruals_for_sale, sale_commission, delivery_charge,
                        return_delivery_charge, services_json, raw_json, rub_to_cny_rate, synced_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(store_id, operation_id) DO UPDATE SET
                        posting_number=excluded.posting_number, operation_date=excluded.operation_date,
                        operation_type=excluded.operation_type, operation_type_name=excluded.operation_type_name,
                        transaction_type=excluded.transaction_type, amount=excluded.amount,
                        accruals_for_sale=excluded.accruals_for_sale,
                        sale_commission=excluded.sale_commission, delivery_charge=excluded.delivery_charge,
                        return_delivery_charge=excluded.return_delivery_charge,
                        services_json=excluded.services_json, raw_json=excluded.raw_json,
                        rub_to_cny_rate=excluded.rub_to_cny_rate,
                        synced_at=excluded.synced_at
                    """,
                    (
                        store_id, operation_id, posting_number, str(operation.get("operation_date") or ""),
                        str(operation.get("operation_type") or ""), str(operation.get("operation_type_name") or ""),
                        str(operation.get("type") or ""), format(self._decimal(operation.get("amount")), "f"),
                        format(self._decimal(operation.get("accruals_for_sale")), "f"),
                        format(self._decimal(operation.get("sale_commission")), "f"),
                        format(self._decimal(operation.get("delivery_charge")), "f"),
                        format(self._decimal(operation.get("return_delivery_charge")), "f"),
                        json.dumps(operation.get("services") or [], ensure_ascii=False),
                        json.dumps(operation, ensure_ascii=False), format(self._decimal(rub_to_cny_rate), "f"), now,
                    ),
                )
                imported += int(existed is None)
                linked += int(bool(posting_number))
            connection.commit()
        return {"received": len(operations), "imported": imported, "linked": linked}

    def upsert_ad_cost(
        self, *, store_id: str, source_key: str, amount_rub: Any,
        posting_number: str = "", cost_date: str = "", sku: str = "",
        campaign_id: str = "", attribution_method: str = "manual", raw: dict[str, Any] | None = None,
    ) -> None:
        with self._lock, closing(self._connect()) as connection:
            connection.execute(
                """
                INSERT INTO ozon_ad_costs (
                    store_id, posting_number, cost_date, sku, campaign_id, amount_rub,
                    attribution_method, source_key, raw_json, synced_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(store_id, source_key) DO UPDATE SET
                    posting_number=excluded.posting_number, cost_date=excluded.cost_date,
                    sku=excluded.sku, campaign_id=excluded.campaign_id,
                    amount_rub=excluded.amount_rub, attribution_method=excluded.attribution_method,
                    raw_json=excluded.raw_json, synced_at=excluded.synced_at
                """,
                (store_id, posting_number, cost_date, sku, campaign_id,
                 format(abs(self._decimal(amount_rub)), "f"), attribution_method, source_key,
                 json.dumps(raw or {}, ensure_ascii=False), utc_now()),
            )
            connection.commit()

    def save_monthly_costs(
        self, *, store_id: str, month: str, purchase_cost: Any, other_cost: Any,
        domestic_shipping_cost: Any = 0, packaging_cost: Any = 0,
        labor_warehouse_cost: Any = 0, tax_cost: Any = 0, notes: str = "",
    ) -> dict[str, Any]:
        with self._lock, closing(self._connect()) as connection:
            connection.execute(
                """
                INSERT INTO profit_monthly_costs(
                    platform, store_id, month, purchase_cost, other_cost, domestic_shipping_cost,
                    packaging_cost, labor_warehouse_cost, tax_cost, notes, updated_at
                ) VALUES('ozon', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(platform, store_id, month) DO UPDATE SET
                    purchase_cost=excluded.purchase_cost, other_cost=excluded.other_cost,
                    domestic_shipping_cost=excluded.domestic_shipping_cost,
                    packaging_cost=excluded.packaging_cost,
                    labor_warehouse_cost=excluded.labor_warehouse_cost,
                    tax_cost=excluded.tax_cost,
                    notes=excluded.notes, updated_at=excluded.updated_at
                """,
                (store_id, month, format(self._decimal(purchase_cost), "f"),
                 format(self._decimal(other_cost), "f"), format(self._decimal(domestic_shipping_cost), "f"),
                 format(self._decimal(packaging_cost), "f"), format(self._decimal(labor_warehouse_cost), "f"),
                 format(self._decimal(tax_cost), "f"), notes, utc_now()),
            )
            connection.commit()
        return self.profit_report(store_id=store_id, date_from=f"{month}-01", date_to=f"{month}-31")

    def save_sku_cost(self, *, store_id: str, sku: str, unit_cost: Any) -> None:
        with self._lock, closing(self._connect()) as connection:
            connection.execute(
                """
                INSERT INTO profit_sku_costs(platform, store_id, sku, unit_cost, updated_at)
                VALUES('ozon', ?, ?, ?, ?)
                ON CONFLICT(platform, store_id, sku) DO UPDATE SET
                    unit_cost=excluded.unit_cost, updated_at=excluded.updated_at
                """,
                (store_id, sku, format(self._decimal(unit_cost), "f"), utc_now()),
            )
            connection.commit()

    def save_ad_report(self, *, store_id: str, date_from: str, date_to: str, report: dict[str, Any]) -> dict[str, Any]:
        with self._lock, closing(self._connect()) as connection:
            connection.execute(
                """INSERT INTO ozon_ad_reports(store_id,date_from,date_to,amount,attributed_sales,campaign_count,raw_json,synced_at)
                VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(store_id,date_from,date_to) DO UPDATE SET amount=excluded.amount,
                attributed_sales=excluded.attributed_sales,campaign_count=excluded.campaign_count,
                raw_json=excluded.raw_json,synced_at=excluded.synced_at""",
                (store_id,date_from,date_to,str(report.get("amount") or "0"),str(report.get("attributed_sales") or "0"),
                 int(report.get("campaign_count") or 0),json.dumps(report,ensure_ascii=False),utc_now()),
            )
            connection.commit()
        return {key: report.get(key) for key in ("amount","attributed_sales","campaign_count","row_count")}

    def profit_report(self, *, store_id: str, date_from: str, date_to: str) -> dict[str, Any]:
        """Aggregate OZON profit by order and SKU for a calendar period."""
        with self._lock, closing(self._connect()) as connection:
            rows = connection.execute(
                """
                SELECT * FROM ozon_orders
                WHERE platform='ozon' AND store_id=? AND substr(in_process_at, 1, 10) BETWEEN ? AND ?
                ORDER BY in_process_at DESC
                """,
                (store_id, date_from, date_to),
            ).fetchall()
            ad_report = connection.execute(
                "SELECT * FROM ozon_ad_reports WHERE store_id=? AND date_from=? AND date_to=?",
                (store_id, date_from, date_to),
            ).fetchone()
            if ad_report is None and date_from[:7] == date_to[:7] and date_from.endswith("-01"):
                ad_report = connection.execute(
                    """SELECT * FROM ozon_ad_reports
                       WHERE store_id=? AND date_from=? AND substr(date_to,1,7)=?
                       ORDER BY date_to DESC LIMIT 1""",
                    (store_id, date_from, date_from[:7]),
                ).fetchone()
            period_ad = self._decimal(ad_report["amount"]) if ad_report else Decimal("0")
            orders: list[dict[str, Any]] = []
            sku_totals: dict[str, dict[str, Any]] = {}
            for row in rows:
                item_rows = connection.execute(
                    "SELECT * FROM ozon_order_items WHERE posting_number=? ORDER BY line_no",
                    (str(row["posting_number"]),),
                ).fetchall()
                items = [{key: item[key] for key in item.keys()} for item in item_rows]
                order = self._public_order(row, items)
                orders.append(order)
                if order["profit"]["status"] == "cancelled":
                    continue
                order_gross = self._decimal(order["profit"]["gross_sales_rub"])
                order_payout = self._decimal(order["profit"]["payout_rub"])
                for item in items:
                    sku = str(item.get("offer_id") or item.get("matched_sku") or item.get("ozon_sku") or "未识别SKU")
                    quantity = max(1, int(item.get("quantity") or 1))
                    sales = self._decimal(item.get("price")) * quantity
                    ratio = sales / order_gross if order_gross else Decimal("0")
                    target = sku_totals.setdefault(sku, {"sku": sku, "name": str(item.get("product_name") or ""), "quantity": 0,
                        "orders": set(), "sales": Decimal("0"), "payout": Decimal("0")})
                    target["quantity"] += quantity
                    target["orders"].add(str(order["posting_number"]))
                    target["sales"] += sales
                    target["payout"] += order_payout * ratio
            cost_rows = connection.execute(
                "SELECT sku, unit_cost FROM profit_sku_costs WHERE platform='ozon' AND store_id=?", (store_id,)
            ).fetchall()
            costs = {str(item["sku"]): self._decimal(item["unit_cost"]) for item in cost_rows}
            sku_values = []
            for sku, item in sku_totals.items():
                unit_cost = costs.get(sku, Decimal("0"))
                total_cost = unit_cost * int(item["quantity"])
                sku_values.append({"sku": sku, "name": item["name"], "quantity": item["quantity"],
                    "order_count": len(item["orders"]), "sales": format(item["sales"], ".2f"),
                    "payout": format(item["payout"], ".2f"), "advertising": "0.00",
                    "unit_cost": format(unit_cost, ".2f"), "total_cost": format(total_cost, ".2f"),
                    "profit": format(item["payout"] - total_cost, ".2f")})
            month = date_from[:7] if date_from[:7] == date_to[:7] else ""
            monthly = connection.execute(
                "SELECT * FROM profit_monthly_costs WHERE platform='ozon' AND store_id=? AND month=?",
                (store_id, month),
            ).fetchone() if month else None
        gross = sum((self._decimal(item["profit"]["gross_sales_rub"]) for item in orders), Decimal("0"))
        payout = sum((self._decimal(item["profit"]["payout_rub"]) for item in orders), Decimal("0"))
        ads = period_ad
        refund_loss = sum((self._decimal(item["profit"].get("refund_loss_cny")) for item in orders), Decimal("0"))
        recorded_cost = sum((self._decimal(item["purchase_total_cost"]) for item in orders), Decimal("0"))
        purchase_cost = self._decimal(monthly["purchase_cost"]) if monthly else recorded_cost
        monthly_cost_breakdown = {
            "purchase_cost": purchase_cost,
            "domestic_shipping_cost": self._decimal(monthly["domestic_shipping_cost"]) if monthly else Decimal("0"),
            "packaging_cost": self._decimal(monthly["packaging_cost"]) if monthly else Decimal("0"),
            "labor_warehouse_cost": self._decimal(monthly["labor_warehouse_cost"]) if monthly else Decimal("0"),
            "tax_cost": self._decimal(monthly["tax_cost"]) if monthly else Decimal("0"),
            "other_cost": self._decimal(monthly["other_cost"]) if monthly else Decimal("0"),
        }
        operating_cost = sum(monthly_cost_breakdown.values(), Decimal("0"))
        non_purchase_cost = operating_cost - purchase_cost
        return {"date_from": date_from, "date_to": date_to, "month": month, "orders": orders, "skus": sku_values,
            "summary": {"order_count": len(orders), "settled_count": sum(item["profit"]["status"] in {"settled", "returned"} for item in orders),
                "gross_sales": format(gross, ".2f"), "payout": format(payout, ".2f"), "advertising": format(ads, ".2f"),
                "refund_loss": format(refund_loss, ".2f"),
                "recorded_purchase_cost": format(recorded_cost, ".2f"), "purchase_cost": format(purchase_cost, ".2f"),
                "other_cost": format(non_purchase_cost, ".2f"), "operating_cost": format(operating_cost, ".2f"),
                "monthly_cost_breakdown": {key: format(value, ".2f") for key, value in monthly_cost_breakdown.items()},
                "profit_before_ads": format(payout - operating_cost, ".2f"),
                "profit": format(payout - operating_cost - ads, ".2f"),
                "advertising_status": "actual" if ad_report else "estimated", "advertising_synced_at": str(ad_report["synced_at"] or "") if ad_report else "",
                "advertising_attributed_sales": str(ad_report["attributed_sales"] or "0") if ad_report else "0",
                "cost_source": "monthly_manual" if monthly else "order_recorded", "notes": str(monthly["notes"] or "") if monthly else ""}}

    def _encrypt_recipient(self, recipient: dict[str, Any]) -> bytes:
        payload = json.dumps(recipient, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        return self.cipher.encrypt(payload)

    def _decrypt_recipient(self, value: bytes | str | None) -> dict[str, str]:
        if not value:
            return {}
        encrypted = value.encode("utf-8") if isinstance(value, str) else bytes(value)
        try:
            payload = json.loads(self.cipher.decrypt(encrypted).decode("utf-8"))
        except (InvalidToken, UnicodeDecodeError, json.JSONDecodeError):
            return {}
        return {str(key): str(item or "") for key, item in payload.items()} if isinstance(payload, dict) else {}

    @staticmethod
    def _text(source: dict[str, Any], *keys: str) -> str:
        for key in keys:
            value = source.get(key)
            if value is not None and str(value).strip():
                return str(value).strip()
        return ""

    def _recipient_from_posting(self, posting: dict[str, Any]) -> dict[str, str]:
        customer = posting.get("customer") if isinstance(posting.get("customer"), dict) else {}
        addressee = posting.get("addressee") if isinstance(posting.get("addressee"), dict) else {}
        address = posting.get("address") if isinstance(posting.get("address"), dict) else {}
        nested_address = customer.get("address") if isinstance(customer.get("address"), dict) else {}
        address = {**address, **nested_address}
        full_name = self._text(addressee, "name") or self._text(customer, "name")
        first, _, last = full_name.partition(" ")
        return {
            "country": self._text(address, "country", "country_code") or "RU",
            "region": self._text(address, "region", "state", "province"),
            "city": self._text(address, "city"),
            "area": self._text(address, "area", "district"),
            "street": self._text(address, "street", "address", "full_address"),
            "postcode": self._text(address, "postcode", "zip", "postal_code"),
            "firstname": first or full_name,
            "lastname": last,
            "mobile_phone": self._text(addressee, "phone") or self._text(customer, "phone"),
            "telephone": "",
            "email": self._text(customer, "customer_email", "email"),
        }

    @staticmethod
    def _posting_products(posting: dict[str, Any]) -> list[dict[str, Any]]:
        for key in ("products", "items"):
            values = posting.get(key)
            if isinstance(values, list):
                return [item for item in values if isinstance(item, dict)]
        return []

    def _match_offer(self, offer_id: str, platform: str = "ozon") -> dict[str, Any]:
        value = str(offer_id or "").strip()
        if not value:
            return {"match_status": "missing_sku", "match_error": f"{platform.upper()} 商品缺少 SKU"}
        try:
            product = self.library.get_product(value)
        except ProductLibraryError:
            candidates = self.library.variant_skus_contained_in(value)
            if not candidates:
                return {"match_status": "not_found", "match_error": f"产品库中找不到 SKU {value}"}
            longest = len(candidates[0])
            best = [sku for sku in candidates if len(sku) == longest]
            if len(best) != 1:
                return {
                    "match_status": "ambiguous",
                    "match_error": f"平台 SKU {value} 同时包含多个产品库 SKU：{', '.join(best[:5])}",
                }
            value = best[0]
            try:
                product = self.library.get_product(value)
            except ProductLibraryError:
                return {"match_status": "not_found", "match_error": f"产品库中找不到 SKU {value}"}
        supplier = str(product.get("supplier") or "hicustom").strip().lower()
        if supplier not in {"hicustom", "s2b"}:
            return {"match_status": "supplier_unsupported", "match_error": "该商品货盘暂不支持下单"}
        variant = next(
            (
                item
                for item in product.get("variants") or []
                if str(item.get("sku") or "").casefold() == value.casefold()
            ),
            None,
        )
        if not isinstance(variant, dict):
            return {"match_status": "not_found", "match_error": f"产品库中找不到规格 SKU {value}"}
        product_code = str(product.get("product_code") or "").strip()
        stock_item_id = variant.get("stock_item_id")
        recipe = product.get("recipe") if isinstance(product.get("recipe"), dict) else {}
        remote_product_id = int(recipe.get("remote_product_id") or 0)
        if supplier == "hicustom" and not product_code:
            return {"match_status": "product_code_missing", "match_error": "产品缺少 HICUSTOM product_code"}
        if supplier == "s2b" and remote_product_id <= 0:
            return {"match_status": "product_code_missing", "match_error": "产品缺少 S2B 定制产品编号"}
        if supplier == "hicustom" and not stock_item_id:
            return {"match_status": "stock_missing", "match_error": "规格缺少 HICUSTOM 变体ID"}
        if supplier == "s2b" and (not variant.get("size_id") or not variant.get("color_id")):
            return {"match_status": "stock_missing", "match_error": "规格缺少 S2B 尺码或颜色编号"}
        return {
            "match_status": "matched",
            "match_error": "",
            "matched_sku": str(variant.get("sku") or value),
            "product_code": product_code,
            "stock_item_id": int(stock_item_id or 0),
            "supplier": supplier,
            "remote_product_id": remote_product_id,
            "color_id": int(variant.get("color_id") or 0),
            "size_id": int(variant.get("size_id") or 0),
            "size_name": str(variant.get("size_name") or ""),
            "color_name": str(variant.get("color_name") or ""),
        }

    @staticmethod
    def _order_key(platform: str, store_id: str, external_order_id: str) -> str:
        platform = str(platform or "ozon").strip().lower()
        store_id = str(store_id or "").strip()
        external_order_id = str(external_order_id or "").strip()
        return external_order_id if platform == "ozon" and not store_id else f"{platform}:{store_id}:{external_order_id}"

    @staticmethod
    def _wb_posting(order: dict[str, Any], rub_to_cny_rate: float = 0.0783) -> dict[str, Any]:
        article = str(order.get("article") or "").strip()
        skus = order.get("skus") if isinstance(order.get("skus"), list) else []
        status = str(order.get("supplierStatus") or order.get("wbStatus") or "new")
        try:
            converted_value = order.get("convertedFinalPrice")
            if converted_value is None:
                converted_value = order.get("convertedPrice")
            converted_code = int(order.get("convertedCurrencyCode") or 0)
            if converted_value is not None and converted_code == 156:
                cny_price = (Decimal(str(converted_value)) / Decimal("100")).quantize(Decimal("0.01"))
            else:
                source_value = converted_value
                source_code = converted_code
                if source_value is None:
                    source_value = order.get("finalPrice")
                    if source_value is None:
                        source_value = order.get("price") or 0
                    source_code = int(order.get("currencyCode") or 643)
                source_price = Decimal(str(source_value)) / Decimal("100")
                cny_price = (
                    source_price
                    if source_code == 156
                    else source_price * Decimal(str(rub_to_cny_rate))
                ).quantize(Decimal("0.01"))
        except (InvalidOperation, TypeError, ValueError):
            cny_price = Decimal("0")
        return {
            "posting_number": str(order.get("id") or ""),
            "status": status,
            "substatus": str(order.get("wbStatus") or ""),
            "in_process_at": str(order.get("createdAt") or ""),
            "shipment_date": str(order.get("deliveryDate") or ""),
            "delivery_method": {
                "warehouse": f"WB {str(order.get('deliveryType') or 'FBS').upper()} · 仓库 {order.get('warehouseId') or '未知'}"
            },
            "products": [
                {
                    "offer_id": article,
                    "sku": str(skus[0] if skus else ""),
                    "name": str(order.get("productName") or f"WB 商品 {order.get('nmId') or ''}").strip(),
                    "quantity": 1,
                    "price": format(cny_price, ".2f"),
                }
            ],
        }

    def create_manual_order(self, *, order_number: str, items: list[dict],
                            platform: str, store_id: str, store_name: str) -> dict:
        products = []
        seen = set()
        for item in items:
            sku = item['sku']
            if sku.casefold() in seen:
                raise OrderError("同一规格请合并数量，不要重复添加")
            seen.add(sku.casefold())
            try:
                product = self.library.get_product(sku)
            except ProductLibraryError as error:
                raise OrderError(str(error)) from None
            variants = [v for v in product.get('variants', []) if str(v.get('sku', '')).casefold() == sku.casefold()]
            if len(variants) != 1:
                raise OrderError("请选择产品库中的有效商品规格")
            match = self._match_offer(sku, platform)
            if match.get('match_status') != 'matched' or str(match.get('matched_sku', '')).casefold() != sku.casefold():
                raise OrderError(str(match.get('match_error') or '所选规格缺少采购资料，请先完善产品库'))
            products.append({'offer_id': variants[0]['sku'], 'name': product['name'],
                             'quantity': item['quantity'], 'price': ''})
        posting = {'posting_number': order_number, 'status': 'manual', 'products': products}
        # Stable identity makes a repeated save safe, and never overwrites an imported order.
        receipt = 'manual:' + hashlib.sha256(f'{platform}:{store_id}:{order_number}'.encode()).hexdigest()
        self.import_postings([posting], platform=platform, store_id=store_id, store_name=store_name,
                             receipt_key=receipt, _manual_create=True)
        return self.get_order(self._order_key(platform, store_id, order_number))

    def import_postings(
        self,
        postings: list[dict[str, Any]],
        *,
        platform: str = "ozon",
        store_id: str = "",
        store_name: str = "",
        wb_rub_to_cny_rate: float = 0.0783,
        receipt_key: str = "",
        _manual_create: bool = False,
    ) -> dict[str, int]:
        platform = str(platform or "ozon").strip().lower()
        if platform not in {"ozon", "wb"}:
            raise OrderError("不支持的订单平台")
        imported = 0
        updated = 0
        with self._lock, closing(self._connect()) as connection:
            if receipt_key:
                if len(receipt_key) > 160:
                    raise OrderError("订单回传批次编号无效")
                digest = hashlib.sha256(json.dumps([platform, store_id, store_name, wb_rub_to_cny_rate, postings],
                    sort_keys=True, ensure_ascii=False, separators=(",", ":")).encode()).hexdigest()
                connection.execute("CREATE TABLE IF NOT EXISTS erp_order_sync_receipts (receipt_key TEXT PRIMARY KEY, digest TEXT NOT NULL, summary_json TEXT NOT NULL)")
                connection.execute("BEGIN IMMEDIATE")
                receipt = connection.execute("SELECT digest,summary_json FROM erp_order_sync_receipts WHERE receipt_key=?", (receipt_key,)).fetchone()
                if receipt:
                    if receipt[0] != digest:
                        raise OrderError("同一订单回传批次的内容发生变化，请核对原电脑检查点")
                    return json.loads(receipt[1])
            for posting in postings:
                if platform == "wb" and not _manual_create:
                    posting = self._wb_posting(posting, wb_rub_to_cny_rate)
                number = str(posting.get("posting_number") or posting.get("order_number") or "").strip()
                if not number:
                    continue
                order_key = self._order_key(platform, store_id, number)
                existing = connection.execute(
                    "SELECT recipient_enc, submission_status FROM ozon_orders WHERE posting_number = ?",
                    (order_key,),
                ).fetchone()
                if _manual_create and existing:
                    raise OrderError("该店铺订单号已存在，请打开原订单核对", status_code=409)
                recipient = self._recipient_from_posting(posting)
                if existing and existing['submission_status'] != 'not_submitted':
                    # ERP must not rewrite production mappings or quantities of an
                    # already queued/submitted/uncertain purchase during a resync.
                    connection.execute('UPDATE ozon_orders SET order_status=?, substatus=?, synced_at=? WHERE posting_number=?',
                                       (self._text(posting, 'status'), self._text(posting, 'substatus'), utc_now(), order_key))
                    updated += 1
                    continue
                if existing:
                    old_recipient = self._decrypt_recipient(existing["recipient_enc"])
                    recipient = {
                        key: value or old_recipient.get(key, "") for key, value in recipient.items()
                    }
                now = utc_now()
                delivery = posting.get("delivery_method")
                financial_data = posting.get("financial_data") if isinstance(posting.get("financial_data"), dict) else {}
                financial_products = financial_data.get("products") if isinstance(financial_data.get("products"), list) else []
                finance_by_sku = {
                    str(value.get("product_id") or value.get("sku") or ""): value
                    for value in financial_products if isinstance(value, dict)
                }
                delivery_name = ""
                if isinstance(delivery, dict):
                    # OZON delivery_method.name/tpl_provider describe the shipping
                    # service. The order screen needs the seller warehouse instead.
                    delivery_name = self._text(delivery, "warehouse")
                connection.execute(
                    """
                    INSERT INTO ozon_orders (
                        posting_number, order_status, substatus, fulfillment_mode,
                        in_process_at, shipment_date, delivery_method, recipient_enc,
                        match_status, submission_status, created_at, updated_at, synced_at,
                        platform, store_id, store_name, external_order_id, financial_data_json
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'unmatched', 'not_submitted', ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(posting_number) DO UPDATE SET
                        order_status = excluded.order_status,
                        substatus = excluded.substatus,
                        in_process_at = excluded.in_process_at,
                        shipment_date = excluded.shipment_date,
                        delivery_method = excluded.delivery_method,
                        recipient_enc = excluded.recipient_enc,
                        store_name = excluded.store_name,
                        financial_data_json = excluded.financial_data_json,
                        updated_at = excluded.updated_at,
                        synced_at = excluded.synced_at
                    """,
                    (
                        order_key,
                        self._text(posting, "status"),
                        self._text(posting, "substatus"),
                        self._text(posting, "fulfillment_mode") or ("fbs" if platform == "wb" else "rfbs"),
                        self._text(posting, "in_process_at"),
                        self._text(posting, "shipment_date", "delivering_date"),
                        delivery_name,
                        self._encrypt_recipient(recipient),
                        now,
                        now,
                        now,
                        platform,
                        str(store_id or ""),
                        str(store_name or ""),
                        number,
                        json.dumps(financial_data, ensure_ascii=False),
                    ),
                )
                connection.execute("DELETE FROM ozon_order_items WHERE posting_number = ?", (order_key,))
                item_statuses: list[str] = []
                for index, item in enumerate(self._posting_products(posting), start=1):
                    offer_id = self._text(item, "offer_id", "sku")
                    item_finance = finance_by_sku.get(self._text(item, "sku"))
                    if not isinstance(item_finance, dict) and index <= len(financial_products):
                        candidate = financial_products[index - 1]
                        item_finance = candidate if isinstance(candidate, dict) else {}
                    item_finance = item_finance if isinstance(item_finance, dict) else {}
                    dimensions = item.get("dimensions") if isinstance(item.get("dimensions"), dict) else {}
                    item_services = item_finance.get("item_services") if isinstance(item_finance.get("item_services"), dict) else {}
                    match = self._match_offer(offer_id, platform)
                    item_statuses.append(str(match["match_status"]))
                    connection.execute(
                        """
                        INSERT INTO ozon_order_items (
                            posting_number, line_no, offer_id, ozon_sku, product_name,
                            quantity, price, matched_sku, product_code, stock_item_id,
                            size_name, color_name, match_status, match_error, supplier,
                            remote_product_id, color_id, size_id, weight, length, width, height,
                            commission_percent, commission_amount, estimated_payout,
                            financial_currency, item_services_json
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            order_key,
                            index,
                            offer_id,
                            self._text(item, "sku"),
                            self._text(item, "name", "product_name"),
                            max(1, int(item.get("quantity") or item.get("qty") or 1)),
                            self._text(item, "price"),
                            str(match.get("matched_sku") or ""),
                            str(match.get("product_code") or ""),
                            match.get("stock_item_id"),
                            str(match.get("size_name") or ""),
                            str(match.get("color_name") or ""),
                            str(match["match_status"]),
                            str(match.get("match_error") or ""),
                            str(match.get("supplier") or "hicustom"),
                            int(match.get("remote_product_id") or 0),
                            int(match.get("color_id") or 0),
                            int(match.get("size_id") or 0),
                            self._text(dimensions, "weight"),
                            self._text(dimensions, "length"),
                            self._text(dimensions, "width"),
                            self._text(dimensions, "height"),
                            self._text(item_finance, "commission_percent"),
                            self._text(item_finance, "commission_amount"),
                            self._text(item_finance, "payout"),
                            self._text(item_finance, "commissions_currency_code", "currency_code"),
                            json.dumps(item_services, ensure_ascii=False),
                        ),
                    )
                match_status = "matched" if item_statuses and all(value == "matched" for value in item_statuses) else "unmatched"
                connection.execute(
                    "UPDATE ozon_orders SET match_status = ? WHERE posting_number = ?",
                    (match_status, order_key),
                )
                if existing:
                    updated += 1
                else:
                    imported += 1
            result = {"imported": imported, "updated": updated, "total": imported + updated}
            if receipt_key:
                connection.execute("INSERT INTO erp_order_sync_receipts VALUES (?,?,?)", (receipt_key, digest, json.dumps(result)))
            connection.commit()
        return result

    def platform_image_offers(self, store_id: str) -> list[str]:
        with self._lock, closing(self._connect()) as connection:
            rows = connection.execute(
                "SELECT DISTINCT i.offer_id FROM ozon_order_items i JOIN ozon_orders o ON o.posting_number=i.posting_number "
                "WHERE o.store_id=? AND TRIM(i.offer_id)<>'' ORDER BY i.offer_id", (store_id,)
            ).fetchall()
        return [str(row[0]) for row in rows]

    def update_platform_images(self, store_id: str, images: dict[str, str]) -> int:
        updated = 0
        with self._lock, closing(self._connect()) as connection:
            for offer_id, url in images.items():
                if not str(url).startswith(("https://", "http://")):
                    continue
                cursor = connection.execute(
                    "UPDATE ozon_order_items SET platform_image_url=? WHERE offer_id=? AND posting_number IN "
                    "(SELECT posting_number FROM ozon_orders WHERE store_id=?)",
                    (url, offer_id, store_id),
                )
                updated += cursor.rowcount
            connection.commit()
        return updated

    def rematch(self, posting_number: str | None = None) -> dict[str, int]:
        matched = 0
        unmatched = 0
        with self._lock, closing(self._connect()) as connection:
            if posting_number:
                rows = connection.execute(
                    "SELECT * FROM ozon_order_items WHERE posting_number = ? ORDER BY line_no",
                    (posting_number,),
                ).fetchall()
            else:
                rows = connection.execute("SELECT i.* FROM ozon_order_items i JOIN ozon_orders o ON o.posting_number=i.posting_number WHERE o.submission_status='not_submitted' ORDER BY i.posting_number, i.line_no").fetchall()
            affected: set[str] = set()
            for row in rows:
                platform_row = connection.execute(
                    "SELECT platform FROM ozon_orders WHERE posting_number = ?",
                    (row["posting_number"],),
                ).fetchone()
                match = self._match_offer(
                    str(row["offer_id"]),
                    str(platform_row["platform"] if platform_row else "ozon"),
                )
                connection.execute(
                    """
                    UPDATE ozon_order_items SET matched_sku = ?, product_code = ?,
                        stock_item_id = ?, size_name = ?, color_name = ?,
                        match_status = ?, match_error = ?, supplier = ?, remote_product_id = ?,
                        color_id = ?, size_id = ? WHERE id = ?
                    """,
                    (
                        str(match.get("matched_sku") or ""),
                        str(match.get("product_code") or ""),
                        match.get("stock_item_id"),
                        str(match.get("size_name") or ""),
                        str(match.get("color_name") or ""),
                        str(match["match_status"]),
                        str(match.get("match_error") or ""),
                        str(match.get("supplier") or "hicustom"),
                        int(match.get("remote_product_id") or 0),
                        int(match.get("color_id") or 0),
                        int(match.get("size_id") or 0),
                        row["id"],
                    ),
                )
                affected.add(str(row["posting_number"]))
            for number in affected:
                statuses = [
                    str(item["match_status"])
                    for item in connection.execute(
                        "SELECT match_status FROM ozon_order_items WHERE posting_number = ?",
                        (number,),
                    ).fetchall()
                ]
                ready = bool(statuses) and all(value == "matched" for value in statuses)
                connection.execute(
                    "UPDATE ozon_orders SET match_status = ?, updated_at = ? WHERE posting_number = ?",
                    ("matched" if ready else "unmatched", utc_now(), number),
                )
                matched += int(ready)
                unmatched += int(not ready)
            connection.commit()
        return {"matched": matched, "unmatched": unmatched}

    def _public_order(self, row: sqlite3.Row, items: list[dict[str, Any]] | None = None) -> dict[str, Any]:
        recipient = self._decrypt_recipient(row["recipient_enc"])
        missing = [key for key in self.REQUIRED_ADDRESS_FIELDS if not recipient.get(key)]
        result = {key: row[key] for key in row.keys() if key != "recipient_enc"}
        result["order_key"] = str(row["posting_number"])
        result["posting_number"] = str(row["external_order_id"] or row["posting_number"])
        for item in items or []:
            item["image_url"] = ""
            sku = str(item.get("matched_sku") or "").strip()
            if not sku:
                continue
            try:
                product = self.library.get_product(sku)
            except ProductLibraryError:
                continue
            assets = [
                asset for asset in product.get("oss_assets") or []
                if str(asset.get("oss_url") or "").startswith("https://")
            ]
            color_id = int(item.get("color_id") or 0)
            chosen = next(
                (
                    asset for asset in assets
                    if asset.get("asset_type") == "rendering"
                    and int(asset.get("color_id") or 0) == color_id
                ),
                None,
            )
            chosen = (
                chosen
                or next((asset for asset in assets if asset.get("asset_type") == "rendering"), None)
                or (assets[0] if assets else None)
            )
            if chosen:
                item["image_url"] = str(chosen.get("oss_url") or "")
        result.update(
            {
                "recipient_name": " ".join(
                    value for value in (recipient.get("firstname", ""), recipient.get("lastname", "")) if value
                ),
                "recipient_ready": not missing,
                "recipient_missing": missing,
                "item_count": len(items or []),
                "items": items or [],
            }
        )
        total = Decimal("0")
        for item in items or []:
            try:
                total += Decimal(str(item.get("price") or "0")) * max(1, int(item.get("quantity") or 1))
            except (InvalidOperation, TypeError, ValueError):
                continue
        result.update(
            {
                "total_price": format(total, "f"),
                "currency": "CNY",
                "received_at": str(row["in_process_at"] or ""),
                "last_ship_at": str(row["shipment_date"] or ""),
                "warehouse": str(row["delivery_method"] or ""),
            }
        )
        with self._lock, closing(self._connect()) as connection:
            costs = connection.execute(
                """
                SELECT COALESCE(SUM(CAST(s.product_cost AS REAL)), 0) AS product_cost,
                       COALESCE(SUM(CAST(s.shipping_cost AS REAL)), 0) AS shipping_cost,
                       COALESCE(SUM(CAST(s.total_cost AS REAL)), 0) AS total_cost
                FROM batch_order_sources s
                WHERE s.posting_number = ?
                """,
                (str(row["posting_number"]),),
            ).fetchone()
            fulfillment = connection.execute(
                """
                SELECT GROUP_CONCAT(DISTINCT g.tracking_number) AS tracking_number,
                       GROUP_CONCAT(DISTINCT g.tracking_company) AS tracking_company,
                       MAX(g.paid_at) AS paid_at,
                       MAX(g.remote_status) AS remote_status
                FROM batch_order_groups g
                JOIN batch_order_sources s ON s.group_id = g.id
                WHERE s.posting_number = ? AND g.submission_status = 'submitted'
                """,
                (str(row["posting_number"]),),
            ).fetchone()
            item_cost_rows = connection.execute(
                """
                SELECT g.sku,
                       COALESCE(SUM(CAST(s.product_cost AS REAL)), 0) AS product_cost,
                       COALESCE(SUM(s.quantity), 0) AS quantity,
                       GROUP_CONCAT(DISTINCT g.tracking_number) AS tracking_number,
                       GROUP_CONCAT(DISTINCT g.tracking_company) AS tracking_company
                FROM batch_order_groups g
                JOIN batch_order_sources s ON s.group_id = g.id
                WHERE s.posting_number = ? AND g.submission_status = 'submitted'
                GROUP BY g.sku
                """,
                (str(row["posting_number"]),),
            ).fetchall()
            finance_rows = connection.execute(
                """
                SELECT * FROM ozon_finance_operations
                WHERE store_id = ? AND posting_number = ?
                ORDER BY operation_date, operation_id
                """,
                (str(row["store_id"] or ""), str(row["external_order_id"] or row["posting_number"])),
            ).fetchall()
            ad_row = connection.execute(
                """
                SELECT COALESCE(SUM(CAST(amount_rub AS REAL)), 0) AS amount
                FROM ozon_ad_costs WHERE store_id = ? AND posting_number = ?
                """,
                (str(row["store_id"] or ""), str(row["external_order_id"] or row["posting_number"])),
            ).fetchone()
            offer_ids = [str(item.get("offer_id") or "") for item in result["items"] if str(item.get("offer_id") or "")]
            history_ratio = Decimal("0")
            history_commission_percent = Decimal("0")
            history_basis = ""
            if len(offer_ids) == 1:
                commission_history = connection.execute(
                    """SELECT AVG(CAST(previous_item.commission_percent AS REAL)) AS percent
                       FROM ozon_order_items previous_item
                       JOIN ozon_orders previous ON previous.posting_number=previous_item.posting_number
                       WHERE previous.store_id=? AND previous_item.offer_id=?
                         AND CAST(previous_item.commission_percent AS REAL)>0""",
                    (str(row["store_id"] or ""), offer_ids[0]),
                ).fetchone()
                history_commission_percent = self._decimal(commission_history["percent"] if commission_history else 0)
                history = connection.execute(
                    """
                    SELECT AVG(net_cny / gross_cny) AS ratio FROM (
                        SELECT previous.posting_number,
                               (SELECT SUM(CAST(ii.price AS REAL) * ii.quantity)
                                FROM ozon_order_items ii WHERE ii.posting_number=previous.posting_number) AS gross_cny,
                               SUM(CAST(finance.amount AS REAL) * CAST(finance.rub_to_cny_rate AS REAL)) AS net_cny
                        FROM ozon_orders previous
                        JOIN ozon_finance_operations finance
                          ON finance.store_id=previous.store_id AND finance.posting_number=previous.external_order_id
                        WHERE previous.store_id=?
                          AND EXISTS (SELECT 1 FROM ozon_order_items wanted
                                      WHERE wanted.posting_number=previous.posting_number AND wanted.offer_id=?)
                          AND lower(replace(finance.operation_type, '_', '')) NOT LIKE '%costperclick%'
                          AND lower(finance.operation_type) NOT LIKE '%subscription%'
                        GROUP BY previous.posting_number
                        HAVING SUM(CAST(finance.accruals_for_sale AS REAL))>0
                    ) WHERE net_cny>0 AND net_cny/gross_cny BETWEEN 0.05 AND 1.2
                    """,
                    (str(row["store_id"] or ""), offer_ids[0]),
                ).fetchone()
                history_ratio = self._decimal(history["ratio"] if history else 0)
                history_basis = "sku_history" if history_ratio else ""
            if not history_ratio:
                if not history_commission_percent:
                    commission_history = connection.execute(
                        """SELECT AVG(CAST(previous_item.commission_percent AS REAL)) AS percent
                           FROM ozon_order_items previous_item
                           JOIN ozon_orders previous ON previous.posting_number=previous_item.posting_number
                           WHERE previous.store_id=? AND CAST(previous_item.commission_percent AS REAL)>0""",
                        (str(row["store_id"] or ""),),
                    ).fetchone()
                    history_commission_percent = self._decimal(commission_history["percent"] if commission_history else 0)
                history = connection.execute(
                    """
                    SELECT AVG(net_cny / gross_cny) AS ratio FROM (
                        SELECT previous.posting_number,
                               (SELECT SUM(CAST(ii.price AS REAL) * ii.quantity)
                                FROM ozon_order_items ii WHERE ii.posting_number=previous.posting_number) AS gross_cny,
                               SUM(CAST(finance.amount AS REAL) * CAST(finance.rub_to_cny_rate AS REAL)) AS net_cny
                        FROM ozon_orders previous
                        JOIN ozon_finance_operations finance
                          ON finance.store_id=previous.store_id AND finance.posting_number=previous.external_order_id
                        WHERE previous.store_id=?
                          AND lower(replace(finance.operation_type, '_', '')) NOT LIKE '%costperclick%'
                          AND lower(finance.operation_type) NOT LIKE '%subscription%'
                        GROUP BY previous.posting_number
                        HAVING SUM(CAST(finance.accruals_for_sale AS REAL))>0
                    ) WHERE net_cny>0 AND net_cny/gross_cny BETWEEN 0.05 AND 1.2
                    """,
                    (str(row["store_id"] or ""),),
                ).fetchone()
                history_ratio = self._decimal(history["ratio"] if history else 0)
                history_basis = "store_history" if history_ratio else ""
        result["purchase_product_cost"] = format(Decimal(str(costs["product_cost"] or 0)), ".2f")
        result["purchase_shipping_cost"] = format(Decimal(str(costs["shipping_cost"] or 0)), ".2f")
        result["purchase_total_cost"] = format(Decimal(str(costs["total_cost"] or 0)), ".2f")
        result["purchase_paid_at"] = str(fulfillment["paid_at"] or "")
        result["purchase_remote_status"] = fulfillment["remote_status"]
        result["tracking_number"] = str(fulfillment["tracking_number"] or row["manual_tracking_number"] or "")
        result["tracking_company"] = str(fulfillment["tracking_company"] or row["manual_tracking_company"] or "")
        item_costs = {str(item["sku"]): item for item in item_cost_rows}
        for item in result["items"]:
            cost = item_costs.get(str(item.get("matched_sku") or ""))
            total_cost = Decimal(str(cost["product_cost"] or 0)) if cost else Decimal("0")
            cost_quantity = max(1, int(cost["quantity"] or 1)) if cost else 1
            item["purchase_unit_cost"] = format(total_cost / cost_quantity, ".2f")
            item["purchase_product_cost"] = format(total_cost, ".2f")
            item["tracking_number"] = str(cost["tracking_number"] or "") if cost else str(row["manual_tracking_number"] or "")
            item["tracking_company"] = str(cost["tracking_company"] or "") if cost else str(row["manual_tracking_company"] or "")
        gross_rub = total
        # CPC is accounted for exclusively from Performance API at period level.
        # It must never reduce an order/SKU payout through Seller Finance.
        order_finance_rows = [
            item for item in finance_rows
            if "costperclick" not in str(item["operation_type"] or "").replace("_", "").lower()
        ]
        finance_amount_rub = sum((self._decimal(item["amount"]) for item in order_finance_rows), Decimal("0"))
        sale_accrual_rub = sum((self._decimal(item["accruals_for_sale"]) for item in order_finance_rows), Decimal("0"))
        commission_rub = sum((self._decimal(item["sale_commission"]) for item in order_finance_rows), Decimal("0"))
        finance_rate = self._decimal(order_finance_rows[0]["rub_to_cny_rate"]) if order_finance_rows else Decimal("0.0783")
        finance_amount = finance_amount_rub * finance_rate
        sale_accrual = sale_accrual_rub * finance_rate
        commission = commission_rub * finance_rate
        delivery = sum(
            ((self._decimal(item["delivery_charge"]) + self._decimal(item["return_delivery_charge"])) * self._decimal(item["rub_to_cny_rate"]) for item in order_finance_rows),
            Decimal("0"),
        )
        services = Decimal("0")
        finance_details: list[dict[str, Any]] = []
        for item in order_finance_rows:
            try:
                service_items = json.loads(str(item["services_json"] or "[]"))
            except (TypeError, ValueError, json.JSONDecodeError):
                service_items = []
            service_total = sum(
                (self._decimal(service.get("price")) for service in service_items if isinstance(service, dict)),
                Decimal("0"),
            )
            services += service_total * self._decimal(item["rub_to_cny_rate"])
            finance_details.append({
                "operation_id": str(item["operation_id"]),
                "date": str(item["operation_date"]),
                "type": str(item["operation_type"]),
                "name": str(item["operation_type_name"]),
                "amount_rub": format(self._decimal(item["amount"]), ".2f"),
                "amount_cny": format(self._decimal(item["amount"]) * self._decimal(item["rub_to_cny_rate"]), ".2f"),
                "rub_to_cny_rate": str(item["rub_to_cny_rate"]),
                "services": service_items,
            })
        has_settlement = bool(order_finance_rows) and (
            sale_accrual != 0 or any(str(item["transaction_type"]) in {"orders", "returns"} for item in order_finance_rows)
        )
        status_text = f"{result.get('order_status', '')} {result.get('substatus', '')}".lower()
        return_operations = {"clientreturnagentoperation", "operationitemreturn", "operationreturngoodsfbsofrms"}
        has_return = any(
            str(item["transaction_type"] or "").lower() == "returns"
            or str(item["operation_type"] or "").replace("_", "").lower() in return_operations
            for item in order_finance_rows
        )
        is_cancelled = any(value in status_text for value in ("cancel", "not_accepted", "canceled", "отмен"))
        api_estimated_payout = sum((
            self._decimal(item.get("estimated_payout")) * (finance_rate if str(item.get("financial_currency") or "").upper() == "RUB" else Decimal("1"))
            for item in result["items"]
        ), Decimal("0"))
        api_estimated_commission = sum((
            self._decimal(item.get("commission_amount")) * (finance_rate if str(item.get("financial_currency") or "").upper() == "RUB" else Decimal("1"))
            for item in result["items"]
        ), Decimal("0"))
        api_estimated_services = Decimal("0")
        for order_item in result["items"]:
            try:
                estimated_services = json.loads(str(order_item.get("item_services_json") or "{}"))
            except (TypeError, ValueError, json.JSONDecodeError):
                estimated_services = {}
            if isinstance(estimated_services, dict):
                service_total = sum((self._decimal(value) for value in estimated_services.values()), Decimal("0"))
                api_estimated_services += service_total * (finance_rate if str(order_item.get("financial_currency") or "").upper() == "RUB" else Decimal("1"))
        if api_estimated_payout:
            estimated_payout = api_estimated_payout
            estimation_basis = "ozon_payout"
            estimation_confidence = "high"
        elif api_estimated_commission or api_estimated_services:
            estimated_payout = gross_rub + api_estimated_commission + api_estimated_services
            estimation_basis = "ozon_fees"
            estimation_confidence = "medium"
        elif history_ratio:
            estimated_payout = gross_rub * history_ratio
            api_estimated_commission = -(gross_rub * history_commission_percent / Decimal("100")) if history_commission_percent else Decimal("0")
            estimation_basis = history_basis
            estimation_confidence = "medium" if history_basis == "sku_history" else "low"
        else:
            api_estimated_commission = -(gross_rub * Decimal("0.20"))
            estimated_payout = gross_rub + api_estimated_commission
            estimation_basis = "fallback_rate"
            estimation_confidence = "low"
        payout_rub = finance_amount if has_settlement else estimated_payout
        purchase_cost = self._decimal(result["purchase_total_cost"])
        if is_cancelled:
            estimated_payout = Decimal("0")
            payout_rub = finance_amount if order_finance_rows else Decimal("0")
            api_estimated_commission = Decimal("0")
            api_estimated_services = Decimal("0")
            estimation_basis = "cancelled"
            estimation_confidence = "actual"
        profit_cny = payout_rub - purchase_cost
        estimated_other_fees = Decimal("0")
        if not has_settlement and not is_cancelled:
            estimated_other_fees = estimated_payout - gross_rub - api_estimated_commission - api_estimated_services
        estimate_details = [
            {"key": "sales", "label": "取消订单金额（不计收入）" if is_cancelled else "商品成交收入", "amount_cny": format(Decimal("0") if is_cancelled else gross_rub, ".2f"), "source": "order"},
            {"key": "commission", "label": "预计销售佣金", "amount_cny": format(api_estimated_commission, ".2f"),
             "source": "ozon" if any(self._decimal(item.get("commission_amount")) for item in result["items"]) else history_basis or "fallback"},
        ]
        if api_estimated_services:
            estimate_details.append({"key": "services", "label": "OZON预计物流及服务费", "amount_cny": format(api_estimated_services, ".2f"), "source": "ozon"})
        if estimated_other_fees:
            estimate_details.append({"key": "other_platform_fees", "label": "物流及其他平台费（综合估算）", "amount_cny": format(estimated_other_fees, ".2f"), "source": history_basis or "fallback"})
        estimate_details.extend([
            {"key": "payout", "label": "预计订单到账", "amount_cny": format(estimated_payout, ".2f"), "source": estimation_basis},
            {"key": "purchase", "label": "采购及国内履约成本", "amount_cny": format(-purchase_cost, ".2f"), "source": "purchase"},
            {"key": "profit", "label": "预计订单利润（不含广告）", "amount_cny": format(estimated_payout - purchase_cost, ".2f"), "source": estimation_basis},
        ])
        return_net = sum((
            self._decimal(item["amount"]) * self._decimal(item["rub_to_cny_rate"])
            for item in order_finance_rows
            if str(item["transaction_type"] or "").lower() == "returns"
               or "return" in str(item["operation_type"] or "").lower()
        ), Decimal("0"))
        refund_loss = max(Decimal("0"), -return_net)
        if is_cancelled:
            refund_loss = max(Decimal("0"), -finance_amount) + purchase_cost
        profit_status = "returned" if has_return else ("cancelled" if is_cancelled else ("settled" if has_settlement else "estimated"))
        result["profit"] = {
            "status": profit_status,
            "currency": "CNY",
            "gross_sales_rub": format(Decimal("0") if is_cancelled else gross_rub, ".2f"),
            "original_order_amount_cny": format(gross_rub, ".2f"),
            "settlement_amount_rub": format(finance_amount, ".2f") if has_settlement else "",
            "sale_accrual_rub": format(sale_accrual, ".2f"),
            "sale_accrual_original_rub": format(sale_accrual_rub, ".2f"),
            "commission_rub": format(commission, ".2f") if has_settlement else format(api_estimated_commission, ".2f"),
            "commission_original_rub": format(commission_rub, ".2f"),
            "delivery_rub": format(delivery, ".2f"),
            "services_rub": format(services, ".2f"),
            "advertising_rub": "0.00",
            "advertising_status": "period_only",
            "payout_rub": format(payout_rub, ".2f"),
            "rub_to_cny_rate": format(finance_rate, "f"),
            "purchase_cost_cny": format(purchase_cost, ".2f"),
            "profit_cny": format(profit_cny, ".2f"),
            "estimation_basis": estimation_basis if not has_settlement else "finance_operations",
            "estimation_confidence": estimation_confidence if not has_settlement else "actual",
            "historical_payout_ratio": format(history_ratio, ".4f") if history_ratio else "",
            "historical_commission_percent": format(history_commission_percent, ".2f") if history_commission_percent else "",
            "estimate_details": estimate_details if not has_settlement and not order_finance_rows else [],
            "has_return": has_return,
            "refund_loss_cny": format(refund_loss, ".2f"),
            "excluded_cpc_operations": len(finance_rows) - len(order_finance_rows),
            "operations": finance_details,
        }
        result["profit"].update({
            "gross_sales_cny": result["profit"]["gross_sales_rub"],
            "settlement_amount_cny": result["profit"]["settlement_amount_rub"],
            "sale_accrual_cny": result["profit"]["sale_accrual_rub"],
            "commission_cny": result["profit"]["commission_rub"],
            "delivery_cny": result["profit"]["delivery_rub"],
            "services_cny": result["profit"]["services_rub"],
            "advertising_cny": result["profit"]["advertising_rub"],
            "payout_cny": result["profit"]["payout_rub"],
        })
        if is_platform_order_cancelled(result):
            result["workflow_status"] = "cancelled_purchased" if result.get("submission_status") == "submitted" else "cancelled"
        elif str(result.get("tms_status") or "") == "manual_tracking":
            result["workflow_status"] = "manual_tracking"
        elif str(result.get("tms_status") or "") in {"submitted", "completed"}:
            result["workflow_status"] = "completed"
        elif str(result.get("tms_status") or "") == "failed":
            result["workflow_status"] = "tms_error"
        elif str(result.get("match_status") or "") != "matched":
            result["workflow_status"] = "pending_match"
        elif result["purchase_paid_at"] and result["tracking_number"]:
            result["workflow_status"] = (
                "purchased_in_transit"
                if str(result.get("platform") or "") == "wb"
                else "purchased_pending_tms"
            )
        elif result["purchase_paid_at"]:
            result["workflow_status"] = "waiting_official_tracking"
        elif str(result.get("submission_status") or "") == "submitted":
            result["workflow_status"] = "supplier_processing"
        else:
            result["workflow_status"] = "pending_purchase"
        return result

    def update_tms_status(self, posting_number: str, **changes: Any) -> dict[str, Any]:
        allowed = {
            "tms_status", "tms_platform_order_id", "tms_package_ids", "tms_item_ids",
            "tms_tail_tracking_numbers", "tms_error", "tms_submitted_at",
        }
        values = {key: value for key, value in changes.items() if key in allowed}
        if not values:
            return self.get_order(posting_number)
        assignments = ", ".join(f"{key} = ?" for key in values)
        with self._lock, closing(self._connect()) as connection:
            connection.execute(
                f"UPDATE ozon_orders SET {assignments}, updated_at = ? WHERE posting_number = ?",
                (*values.values(), utc_now(), posting_number),
            )
            connection.commit()
        return self.get_order(posting_number)

    def record_manual_tracking(self, posting_number: str, tracking_number: str, tracking_company: str) -> dict[str, Any]:
        number = str(tracking_number or "").strip()
        company = str(tracking_company or "").strip()
        if not number or not company:
            raise OrderError("请填写物流公司和物流单号")
        if len(number) > 100 or len(company) > 100 or any(char.isspace() for char in number):
            raise OrderError("物流公司或物流单号格式无效")
        with self._lock, closing(self._connect()) as connection:
            row = connection.execute("SELECT * FROM ozon_orders WHERE posting_number = ?", (posting_number,)).fetchone()
            if row is None:
                raise OrderError("订单不存在", status_code=404)
            if is_platform_order_cancelled(row):
                raise OrderError("平台已取消的订单不能录入正常物流单号", status_code=409)
            if row["match_status"] != "matched" or row["submission_status"] != "not_submitted":
                raise OrderError("仅待创建采购、且尚未提交供应商的订单可人工录入物流单号", status_code=409)
            if row["tms_status"] not in ("", "manual_tracking"):
                raise OrderError("订单已进入其他处理流程，不能覆盖物流单号", status_code=409)
            now = utc_now()
            connection.execute(
                """UPDATE ozon_orders SET manual_tracking_number = ?, manual_tracking_company = ?,
                   manual_tracking_at = ?, tms_status = 'manual_tracking', updated_at = ?
                   WHERE posting_number = ?""",
                (number, company, now, now, posting_number),
            )
            connection.commit()
        return self.get_order(posting_number)

    def tms_packaging_payload(
        self, posting_number: str, remote_order: dict[str, Any], warehouse_id: str
    ) -> tuple[dict[str, Any], list[str], list[str]]:
        order = self.get_order(posting_number)
        if str(order.get("platform") or "") != "ozon":
            raise OrderError("目前只支持提交 OZON 代打包订单")
        if order.get("workflow_status") not in {"purchased_pending_tms", "tms_error"}:
            raise OrderError("订单尚未完成采购并取得国内快递单号")
        local_items = list(order.get("items") or [])
        packages_payload: list[dict[str, Any]] = []
        package_ids: list[str] = []
        item_ids: list[str] = []
        matched_local_ids: set[int] = set()
        for package in remote_order.get("packages") or []:
            if not isinstance(package, dict):
                continue
            package_id = str(package.get("platformOrderPackageId") or "").strip()
            remote_items_payload: list[dict[str, Any]] = []
            for remote_item in package.get("items") or []:
                if not isinstance(remote_item, dict):
                    continue
                remote_sku = str(remote_item.get("platformSkuId") or "").strip()
                remote_outer_sku = str(remote_item.get("platformOuterSkuId") or "").strip()
                local_item = next(
                    (
                        item for item in local_items
                        if id(item) not in matched_local_ids and (
                            (remote_sku and str(item.get("ozon_sku") or "") == remote_sku)
                            or (remote_outer_sku and str(item.get("offer_id") or "") == remote_outer_sku)
                        )
                    ),
                    None,
                )
                if local_item is None:
                    raise OrderError(
                        f"TMS 商品 {remote_sku or remote_outer_sku or '未知'} 无法匹配本地 OZON 商品"
                    )
                tracking_number = str(local_item.get("tracking_number") or "").split(",")[0].strip()
                tracking_company = str(local_item.get("tracking_company") or "").split(",")[0].strip()
                if not tracking_number or not tracking_company:
                    raise OrderError(f"SKU {local_item.get('matched_sku')} 缺少国内物流信息")
                item_id = str(remote_item.get("platformOrderItemId") or "").strip()
                remote_items_payload.append({
                    "platformOrderItemId": item_id,
                    "headPackages": [{
                        "trackingNo": tracking_number,
                        "logisticsCompany": tracking_company,
                    }],
                    "platformOuterSkuId": remote_outer_sku,
                })
                matched_local_ids.add(id(local_item))
                item_ids.append(item_id)
            if remote_items_payload:
                if not package_id:
                    raise OrderError("TMS 订单响应缺少包裹 ID")
                packages_payload.append({
                    "platformOrderPackageId": package_id,
                    "items": remote_items_payload,
                })
                package_ids.append(package_id)
        if not packages_payload or len(matched_local_ids) != len(local_items):
            raise OrderError("TMS 包裹商品与本地 OZON 订单商品不完整")
        return ({
            "platformOrderId": str(remote_order.get("platformOrderId") or ""),
            "customerNote": f"程序自动提交；采购订单 {order.get('hicustom_order_id') or ''}"[:500],
            "isSubmitAlibbServiceOrder": 0,
            "packages": packages_payload,
            "warehouseId": str(warehouse_id),
            "logisticsId": "",
            "fmTn": "",
            "shopeeWarehouseId": "",
            "shopeeWarehouseType": "",
        }, package_ids, item_ids)

    def export_ozon_purchases(self, purchase_date: str) -> tuple[Path, int]:
        try:
            normalized_date = datetime.strptime(str(purchase_date), "%Y-%m-%d").date().isoformat()
        except ValueError as error:
            raise OrderError("采购日期格式无效") from error
        with self._lock, closing(self._connect()) as connection:
            rows = connection.execute(
                """
                SELECT o.external_order_id AS order_number,
                       g.tracking_company, g.tracking_number,
                       i.ozon_sku, g.sku AS local_sku,
                       s.quantity, s.product_cost, g.paid_at
                FROM batch_order_groups g
                JOIN batch_order_sources s ON s.group_id = g.id
                JOIN ozon_orders o ON o.posting_number = s.posting_number
                JOIN ozon_order_items i
                  ON i.posting_number = o.posting_number AND i.matched_sku = g.sku
                WHERE o.platform = 'ozon'
                  AND g.submission_status = 'submitted'
                  AND SUBSTR(g.paid_at, 1, 10) = ?
                ORDER BY g.paid_at, o.external_order_id, i.line_no
                """,
                (normalized_date,),
            ).fetchall()
        if not rows:
            raise OrderError(f"{normalized_date} 没有已支付的 OZON 采购订单", status_code=404)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "OZON采购订单"
        headers = ["订单编号", "国内物流公司", "国内快递单号", "OZON SKU ID", "SKU成本（人民币）"]
        sheet.append(headers)
        for row in rows:
            quantity = max(1, int(row["quantity"] or 1))
            unit_cost = Decimal(str(row["product_cost"] or 0)) / quantity
            sheet.append([
                str(row["order_number"] or ""),
                str(row["tracking_company"] or ""),
                str(row["tracking_number"] or ""),
                str(row["ozon_sku"] or ""),
                float(unit_cost.quantize(Decimal("0.01"))),
            ])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:E{sheet.max_row}"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column, width in {"A": 24, "B": 18, "C": 24, "D": 18, "E": 20}.items():
            sheet.column_dimensions[column].width = width
        for cell in sheet["E"][1:]:
            cell.number_format = '¥0.00'
        export_root = self.root / "order_exports"
        export_root.mkdir(parents=True, exist_ok=True)
        path = export_root / f"ozon-purchases-{normalized_date}-{datetime.now().strftime('%H%M%S')}.xlsx"
        workbook.save(path)
        return path, len(rows)

    def export_tms_tracking(self, order_keys: list[str]) -> tuple[Path, int]:
        keys = list(dict.fromkeys(str(value or "").strip() for value in order_keys if str(value or "").strip()))
        if not keys:
            raise OrderError("请先勾选已采购订单")
        rows: list[list[Any]] = []
        for key in keys:
            order = self.get_order(key)
            if is_platform_order_cancelled(order):
                raise OrderError(f"订单 {order.get('posting_number') or key} 已被平台取消，请先人工核对供应商采购，不能按正常订单完成")
            if order.get("workflow_status") not in {"purchased_pending_tms", "purchased_in_transit", "manual_tracking", "tms_error"}:
                raise OrderError(f"订单 {order.get('posting_number') or key} 不是已采购待处理状态")
            order_number = str(order.get("posting_number") or "").strip()
            purchase_amount = Decimal(str(order.get("purchase_total_cost") or 0)).quantize(Decimal("0.01"))
            for index, item in enumerate(order.get("items") or []):
                tracking_number = str(item.get("tracking_number") or order.get("tracking_number") or "").split(",")[0].strip()
                tracking_company = str(item.get("tracking_company") or order.get("tracking_company") or "").split(",")[0].strip()
                platform_sku = str(item.get("offer_id") or item.get("matched_sku") or item.get("ozon_sku") or "").strip()
                if not platform_sku or not tracking_number or not tracking_company:
                    raise OrderError(f"订单 {order_number} 缺少平台 SKU、快递单号或物流公司")
                rows.append([
                    order_number,
                    platform_sku,
                    float(purchase_amount) if index == 0 else None,
                    tracking_number,
                    tracking_company,
                ])
        if not rows:
            raise OrderError("勾选的订单没有可导出的商品")

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet.merge_cells("A1:E1")
        sheet["A1"] = (
            "注意：\n"
            "一、红色*表头部分为必填项。\n"
            "二、产品推荐填写平台SKU，若没有则依次用SKUID、主货号进行匹配。\n"
            "三、同一订单同一产品填写多行快递信息会关联到同一产品上。（同一订单同一产品不支持关联重复快递单号）\n"
            "四、多行采购金额仅识别并关联第一行，其它自动过滤，采购金额最多可保留小数点后两位超出将会自动忽略。"
        )
        headers = ["*订单号", "*平台SKU/SKUID/主货号", "采购金额\n（CNY）", "*快递单号", "*物流公司"]
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        thin = Side(style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        sheet["A1"].alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        sheet["A1"].font = Font(name="Microsoft YaHei", size=11, italic=True, bold=True)
        sheet.row_dimensions[1].height = 120
        for cell in sheet[2]:
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            cell.font = Font(name="Microsoft YaHei", size=11, italic=True, bold=True, color="FF0000" if cell.column != 3 else "000000")
            cell.border = border
        sheet.row_dimensions[2].height = 42
        for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row, min_col=1, max_col=5):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border
        for cell in sheet["C"][2:]:
            cell.number_format = "0.00"
        for column, width in {"A": 26, "B": 34, "C": 22, "D": 26, "E": 22}.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A3"
        export_root = self.root / "order_exports"
        export_root.mkdir(parents=True, exist_ok=True)
        path = export_root / f"提交代打包-快递单号-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
        workbook.save(path)
        return path, len(rows)

    def mark_orders_completed(self, order_keys: list[str]) -> dict[str, int]:
        keys = list(dict.fromkeys(str(value or "").strip() for value in order_keys if str(value or "").strip()))
        if not keys:
            raise OrderError("请先勾选已采购订单")
        for key in keys:
            order = self.get_order(key)
            if is_platform_order_cancelled(order):
                raise OrderError(f"订单 {order['posting_number']} 已被平台取消，请先人工核对供应商采购，不能按正常订单完成", status_code=409)
            # Local-only archival is also allowed after domestic TMS data was
            # written, even when a retired SKU no longer matches the library.
            can_archive_written_tms = str(order.get("tms_status") or "") == "info_written"
            if not can_archive_written_tms and order.get("workflow_status") not in {"purchased_pending_tms", "purchased_in_transit", "manual_tracking", "tms_error"}:
                raise OrderError(f"订单 {order.get('posting_number') or key} 不是已采购待处理状态")
        completed = 0
        with self._lock, closing(self._connect()) as connection:
            for key in keys:
                row = connection.execute(
                    "SELECT tms_status FROM ozon_orders WHERE posting_number = ?", (key,)
                ).fetchone()
                if row is None:
                    raise OrderError(f"订单 {key} 不存在", status_code=404)
                connection.execute(
                    """
                    UPDATE ozon_orders
                    SET tms_status = 'completed', tms_error = '', tms_submitted_at = ?, updated_at = ?
                    WHERE posting_number = ?
                    """,
                    (utc_now(), utc_now(), key),
                )
                completed += 1
            connection.commit()
        return {"completed": completed}

    def product_options(self, search: str = "", limit: int = 30) -> list[dict[str, Any]]:
        payload = self.library.list_products(
            search=search,
            supplier="",
            limit=max(1, min(50, limit)),
            offset=0,
        )
        options: list[dict[str, Any]] = []
        for summary in payload.get("products") or []:
            try:
                product = self.library.get_product(str(summary.get("sku") or ""))
            except ProductLibraryError:
                continue
            for variant in product.get("variants") or []:
                sku = str(variant.get("sku") or "").strip()
                if not sku:
                    continue
                options.append(
                    {
                        "sku": sku,
                        "product_sku": str(product.get("sku") or ""),
                        "name": str(product.get("name") or ""),
                        "color_name": str(variant.get("color_name") or ""),
                        "size_name": str(variant.get("size_name") or ""),
                        "product_code": str(product.get("product_code") or ""),
                        "supplier": str(variant.get("supplier") or product.get("supplier") or "hicustom"),
                        "supplier_name": "S2B 货盘" if str(variant.get("supplier") or product.get("supplier") or "hicustom") == "s2b" else "指纹科技",
                        "thumbnail_url": str(summary.get("thumbnail_url") or ""),
                    }
                )
        return options[: max(1, min(200, limit * 10))]

    def manual_match(self, posting_number: str, line_no: int, sku: str) -> dict[str, Any]:
        order = self.get_order(posting_number)
        if order['submission_status'] != 'not_submitted':
            raise OrderError('已采购、排队中或结果不确定的订单不能改绑 SKU', status_code=409)
        match = self._match_offer(sku, str(order.get("platform") or "ozon"))
        if match.get("match_status") != "matched":
            raise OrderError(str(match.get("match_error") or "所选 SKU 无法用于订单"))
        with self._lock, closing(self._connect()) as connection:
            cursor = connection.execute(
                """
                UPDATE ozon_order_items SET matched_sku = ?, product_code = ?,
                    stock_item_id = ?, size_name = ?, color_name = ?,
                    match_status = 'matched', match_error = '', supplier = ?,
                    remote_product_id = ?, color_id = ?, size_id = ?
                WHERE posting_number = ? AND line_no = ?
                """,
                (
                    str(match.get("matched_sku") or sku),
                    str(match.get("product_code") or ""),
                    match.get("stock_item_id"),
                    str(match.get("size_name") or ""),
                    str(match.get("color_name") or ""),
                    str(match.get("supplier") or "hicustom"),
                    int(match.get("remote_product_id") or 0),
                    int(match.get("color_id") or 0),
                    int(match.get("size_id") or 0),
                    posting_number,
                    int(line_no),
                ),
            )
            if cursor.rowcount <= 0:
                raise OrderError("订单商品不存在", status_code=404)
            statuses = [
                str(item["match_status"])
                for item in connection.execute(
                    "SELECT match_status FROM ozon_order_items WHERE posting_number = ?",
                    (posting_number,),
                ).fetchall()
            ]
            ready = bool(statuses) and all(value == "matched" for value in statuses)
            connection.execute(
                "UPDATE ozon_orders SET match_status = ?, updated_at = ? WHERE posting_number = ?",
                ("matched" if ready else "unmatched", utc_now(), posting_number),
            )
            connection.commit()
        return self.get_order(posting_number)

    def preview_batch(self, order_keys: list[str], merge_mode: str = "smart") -> dict[str, Any]:
        keys = list(dict.fromkeys(str(value or "").strip() for value in order_keys if str(value or "").strip()))
        if not keys:
            raise OrderError("请至少选择一个订单")
        groups: dict[str, dict[str, Any]] = {}
        selected: list[dict[str, Any]] = []
        for key in keys:
            order = self.get_order(key)
            if is_platform_order_cancelled(order):
                raise OrderError(f"订单 {order['posting_number']} 已被平台取消，不能创建或重新提交采购", status_code=409)
            if order.get("tms_status") == "manual_tracking":
                raise OrderError(f"订单 {order['posting_number']} 已人工录入物流单号，不能重复创建供应商采购", status_code=409)
            if order["submission_status"] != "not_submitted":
                raise OrderError(f"订单 {order['posting_number']} 已下单或正在处理，不能重复加入", status_code=409)
            if order["match_status"] != "matched":
                raise OrderError(f"订单 {order['posting_number']} 仍有商品未匹配产品库")
            selected.append(order)
            for item in order.get("items") or []:
                sku = str(item.get("matched_sku") or "").strip()
                supplier = str(item.get("supplier") or "hicustom")
                if not sku or (supplier == "hicustom" and (not item.get("product_code") or not item.get("stock_item_id"))) or (supplier == "s2b" and not item.get("remote_product_id")):
                    raise OrderError(f"订单 {order['posting_number']} 的 SKU 无法用于对应货盘下单")
                merge_mode = merge_mode if merge_mode in {"smart", "same_sku", "separate"} else "smart"
                grouping_key = f"{supplier}:{sku.casefold()}" if merge_mode != "separate" else f"{supplier}:{key}:{item.get('line_no')}"
                group = groups.setdefault(
                    grouping_key,
                    {
                        "group_key": grouping_key,
                        "sku": sku,
                        "supplier": supplier,
                        "product_name": str(item.get("product_name") or ""),
                        "product_code": str(item.get("product_code") or ""),
                        "stock_item_id": int(item["stock_item_id"] or 0),
                        "remote_product_id": int(item.get("remote_product_id") or 0),
                        "color_id": int(item.get("color_id") or 0),
                        "size_id": int(item.get("size_id") or 0),
                        "size_name": str(item.get("size_name") or ""),
                        "color_name": str(item.get("color_name") or ""),
                        "quantity": 0,
                        "sales_amount": Decimal("0"),
                        "sources": {},
                    },
                )
                quantity = max(1, int(item.get("quantity") or 1))
                group["quantity"] += quantity
                try:
                    group["sales_amount"] += Decimal(str(item.get("price") or "0")) * quantity
                except InvalidOperation:
                    pass
                group["sources"][key] = int(group["sources"].get(key, 0)) + quantity
        values = []
        for group in groups.values():
            values.append(
                {
                    **{key: value for key, value in group.items() if key != "sources"},
                    "sales_amount": format(group["sales_amount"], "f"),
                    "source_order_count": len(group["sources"]),
                    "source_order_keys": list(group["sources"]),
                    "source_quantities": dict(group["sources"]),
                }
            )
        values.sort(key=lambda item: str(item["sku"]).casefold())
        return {
            "order_count": len(selected),
            "group_count": len(values),
            "total_quantity": sum(int(item["quantity"]) for item in values),
            "groups": values,
            "merge_mode": merge_mode,
            "estimated_purchase_count": (
                (1 if any(item["supplier"] == "hicustom" for item in values) else 0)
                + sum(1 for item in values if item["supplier"] == "s2b")
                if merge_mode == "smart" else len(values)
            ),
        }

    def queue_batch(
        self,
        order_keys: list[str],
        address: dict[str, Any],
        logistics: list[dict[str, Any]],
        merge_mode: str = "smart",
    ) -> dict[str, Any]:
        clean_address = {
            key: str(address.get(key) or "").strip()
            for key in (
                "country", "region", "city", "area", "street", "postcode",
                "firstname", "lastname", "mobile_phone", "telephone", "email",
            )
        }
        missing = [key for key in self.REQUIRED_ADDRESS_FIELDS if not clean_address.get(key)]
        if missing:
            raise OrderError(f"固定收货地址不完整：缺少 {', '.join(missing)}")
        preview = self.preview_batch(order_keys, merge_mode)
        choices = {
            str(item.get("group_key") or item.get("sku") or "").strip().casefold(): {
                "company_name": str(item.get("company_name") or "").strip(),
                "shipping_amount": self._money(item.get("shipping_amount")),
                "logistics_id": int(item.get("logistics_id") or 0),
            }
            for item in logistics
            if str(item.get("sku") or "").strip()
        }
        for group in preview["groups"]:
            choice = choices.get(str(group["group_key"]).casefold()) or choices.get(str(group["sku"]).casefold())
            if not choice or not choice["company_name"]:
                raise OrderError(f"SKU {group['sku']} 还没有选择物流方式")
            if group.get("supplier") == "s2b" and choice["logistics_id"] <= 0:
                raise OrderError(f"S2B SKU {group['sku']} 必须填写物流渠道编号")
        batch_id = uuid.uuid4().hex
        now = utc_now()
        smart_purchase_ids: dict[tuple[str, str], str] = {}
        with self._lock, closing(self._connect()) as connection:
            for group in preview["groups"]:
                group_id = uuid.uuid4().hex
                choice = choices.get(str(group["group_key"]).casefold()) or choices[str(group["sku"]).casefold()]
                purchase_id = (
                    smart_purchase_ids.setdefault((str(group["supplier"]), str(choice["company_name"]).casefold()), uuid.uuid4().hex)
                    if merge_mode == "smart" and group["supplier"] == "hicustom" else uuid.uuid4().hex
                )
                connection.execute(
                    """
                    INSERT INTO batch_order_groups (
                        id, batch_id, sku, product_name, product_code, stock_item_id,
                        size_name, color_name, quantity, sales_amount, address_enc,
                        logistics_company_name, quoted_shipping_amount, supplier,
                        remote_product_id, color_id, size_id, logistics_id,
                        purchase_id, merge_mode, submission_status, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'queued', ?, ?)
                    """,
                    (
                        group_id, batch_id, group["sku"], group["product_name"],
                        group["product_code"], group["stock_item_id"], group["size_name"],
                        group["color_name"], group["quantity"], group["sales_amount"],
                        self._encrypt_recipient(clean_address),
                        choice["company_name"], choice["shipping_amount"],
                        group["supplier"], group["remote_product_id"], group["color_id"],
                        group["size_id"], choice["logistics_id"], purchase_id, merge_mode,
                        now, now,
                    ),
                )
                source_items = list(group["source_quantities"].items())
                quoted_parts = self._split_money(
                    Decimal(choice["shipping_amount"]),
                    [int(quantity) for _, quantity in source_items],
                    evenly=True,
                )
                for (order_key, quantity), quoted_part in zip(source_items, quoted_parts):
                    connection.execute(
                        """
                        INSERT INTO batch_order_sources (
                            group_id, posting_number, quantity, shipping_cost, total_cost
                        ) VALUES (?, ?, ?, ?, ?)
                        """,
                        (
                            group_id, order_key, int(quantity),
                            format(quoted_part, ".2f"), format(quoted_part, ".2f"),
                        ),
                    )
            placeholders = ",".join("?" for _ in order_keys)
            connection.execute(
                f"UPDATE ozon_orders SET submission_status = 'queued', submission_error = '', updated_at = ? WHERE posting_number IN ({placeholders})",
                (now, *order_keys),
            )
            connection.commit()
        return {"batch_id": batch_id, **preview, "groups": self.list_batch_groups(batch_id=batch_id)}

    def _public_batch_group(self, row: sqlite3.Row, sources: list[sqlite3.Row]) -> dict[str, Any]:
        return {
            **{key: row[key] for key in row.keys() if key != "address_enc"},
            "source_order_count": len(sources),
            "source_orders": [
                {
                    "order_key": str(item["posting_number"]),
                    "quantity": int(item["quantity"]),
                    "product_cost": str(item["product_cost"]),
                    "shipping_cost": str(item["shipping_cost"]),
                    "total_cost": str(item["total_cost"]),
                }
                for item in sources
            ],
        }

    def list_batch_groups(self, *, batch_id: str = "", limit: int = 100) -> list[dict[str, Any]]:
        where = "WHERE batch_id = ?" if batch_id else ""
        params: tuple[Any, ...] = (batch_id,) if batch_id else ()
        with self._lock, closing(self._connect()) as connection:
            rows = connection.execute(
                f"SELECT * FROM batch_order_groups {where} ORDER BY created_at DESC LIMIT ?",
                (*params, max(1, min(500, limit))),
            ).fetchall()
            return [
                self._public_batch_group(
                    row,
                    connection.execute(
                        "SELECT posting_number, quantity, product_cost, shipping_cost, total_cost FROM batch_order_sources WHERE group_id = ? ORDER BY posting_number",
                        (row["id"],),
                    ).fetchall(),
                )
                for row in rows
            ]

    def get_batch_group(self, group_id: str) -> dict[str, Any]:
        with self._lock, closing(self._connect()) as connection:
            row = connection.execute("SELECT * FROM batch_order_groups WHERE id = ?", (group_id,)).fetchone()
            if row is None:
                raise OrderError("批量下单任务不存在", status_code=404)
            sources = connection.execute(
                "SELECT posting_number, quantity, product_cost, shipping_cost, total_cost FROM batch_order_sources WHERE group_id = ? ORDER BY posting_number",
                (group_id,),
            ).fetchall()
            result = self._public_batch_group(row, sources)
            result["address"] = self._decrypt_recipient(row["address_enc"])
            return result

    def assert_purchase_sources_active(self, group_ids: list[str]) -> None:
        for group_id in group_ids:
            group = self.get_batch_group(group_id)
            for source in group["source_orders"]:
                order = self.get_order(source["order_key"])
                if is_platform_order_cancelled(order):
                    raise OrderError(f"平台订单 {order['posting_number']} 已取消；供应商采购已拦截，请人工核对", status_code=409)
                if order.get("tms_status") == "manual_tracking":
                    raise OrderError(f"平台订单 {order['posting_number']} 已人工录入物流单号，供应商采购已拦截", status_code=409)

    def reusable_hicustom_product(self, sku: str, *, exclude_group_id: str = "") -> dict[str, Any] | None:
        """Return the newest prepared production product for the same SKU."""
        with self._lock, closing(self._connect()) as connection:
            row = connection.execute(
                """
                SELECT id FROM batch_order_groups
                WHERE sku = ? AND supplier = 'hicustom' AND id <> ?
                  AND fulfillment_product_code <> ''
                  AND fulfillment_gallery_codes <> '[]'
                ORDER BY product_prepared_at DESC, updated_at DESC
                LIMIT 1
                """,
                (str(sku), str(exclude_group_id)),
            ).fetchone()
        return self.get_batch_group(str(row["id"])) if row else None

    def reusable_s2b_product(self, sku: str, *, exclude_group_id: str = "") -> dict[str, Any] | None:
        """Return the newest prepared S2B product for the exact same ERP SKU."""
        with self._lock, closing(self._connect()) as connection:
            row = connection.execute(
                """
                SELECT id FROM batch_order_groups
                WHERE sku = ? AND supplier = 's2b' AND id <> ?
                  AND remote_product_id > 0
                  AND fulfillment_product_code <> ''
                  AND fulfillment_gallery_codes <> '[]'
                ORDER BY product_prepared_at DESC, updated_at DESC
                LIMIT 1
                """,
                (str(sku), str(exclude_group_id)),
            ).fetchone()
        return self.get_batch_group(str(row["id"])) if row else None

    def list_purchase_groups(self, purchase_id: str) -> list[dict[str, Any]]:
        if not purchase_id:
            return []
        with self._lock, closing(self._connect()) as connection:
            rows = connection.execute(
                "SELECT id FROM batch_order_groups WHERE purchase_id = ? ORDER BY created_at, id",
                (purchase_id,),
            ).fetchall()
        return [self.get_batch_group(str(row["id"])) for row in rows]

    def update_purchase_groups(self, purchase_id: str, **changes: Any) -> list[dict[str, Any]]:
        groups = self.list_purchase_groups(purchase_id)
        return [self.update_batch_group(str(group["id"]), **changes) for group in groups]

    def next_batch_queued(self) -> dict[str, Any] | None:
        with self._lock, closing(self._connect()) as connection:
            row = connection.execute(
                "SELECT id FROM batch_order_groups WHERE submission_status = 'queued' ORDER BY created_at LIMIT 1"
            ).fetchone()
        return self.get_batch_group(str(row["id"])) if row else None

    def _refresh_batch_source_statuses(self, connection: sqlite3.Connection, group_id: str) -> None:
        sources = connection.execute(
            "SELECT posting_number FROM batch_order_sources WHERE group_id = ?", (group_id,)
        ).fetchall()
        for source in sources:
            key = str(source["posting_number"])
            groups = connection.execute(
                """
                SELECT g.submission_status, g.hicustom_order_id, g.submission_error
                FROM batch_order_groups g JOIN batch_order_sources s ON s.group_id = g.id
                WHERE s.posting_number = ?
                """,
                (key,),
            ).fetchall()
            statuses = {str(item["submission_status"]) for item in groups}
            if statuses and statuses <= {"submitted"}:
                status = "submitted"
            elif statuses & {"failed", "unknown"}:
                status = "failed" if "failed" in statuses else "unknown"
            elif statuses & {"preparing_product", "checking", "submitting"}:
                status = "submitting"
            else:
                status = "queued"
            ids = ",".join(str(item["hicustom_order_id"]) for item in groups if item["hicustom_order_id"])
            errors = "；".join(str(item["submission_error"]) for item in groups if item["submission_error"])
            connection.execute(
                "UPDATE ozon_orders SET submission_status = ?, hicustom_order_id = ?, submission_error = ?, updated_at = ? WHERE posting_number = ?",
                (status, ids, errors[:1000], utc_now(), key),
            )

    def update_batch_group(self, group_id: str, **changes: Any) -> dict[str, Any]:
        allowed = {
            "submission_status", "hicustom_order_id", "submission_error",
            "actual_product_amount", "actual_shipping_amount", "actual_grand_total",
            "cost_recorded_at",
            "fulfillment_product_code", "fulfillment_gallery_codes", "product_revision",
            "out_order_id", "supersedes_order_id", "product_prepared_at",
            "remote_status", "paid_at", "tracking_number", "tracking_company",
            "remote_checked_at",
            "remote_product_id", "logistics_id", "logistics_company_name",
            "quoted_shipping_amount",
        }
        values = {key: value for key, value in changes.items() if key in allowed}
        if not values:
            return self.get_batch_group(group_id)
        assignments = ", ".join(f"{key} = ?" for key in values)
        with self._lock, closing(self._connect()) as connection:
            connection.execute(
                f"UPDATE batch_order_groups SET {assignments}, updated_at = ? WHERE id = ?",
                (*values.values(), utc_now(), group_id),
            )
            self._refresh_batch_source_statuses(connection, group_id)
            connection.commit()
        return self.get_batch_group(group_id)

    def retry_failed_order(self, posting_number: str) -> dict[str, Any]:
        """Requeue failed supplier groups after proving no remote order was created."""
        order = self.get_order(posting_number)
        if str(order.get("submission_status") or "") != "failed":
            raise OrderError("只有下单失败的订单可以人工重新下单", status_code=409)
        with self._lock, closing(self._connect()) as connection:
            groups = connection.execute(
                """
                SELECT g.* FROM batch_order_groups g
                JOIN batch_order_sources s ON s.group_id = g.id
                WHERE s.posting_number = ?
                ORDER BY g.created_at, g.id
                """,
                (posting_number,),
            ).fetchall()
            if not groups:
                raise OrderError("没有找到这笔订单的采购任务，请重新加入批量下单", status_code=404)
            if any(str(row["submission_status"] or "") != "failed" for row in groups):
                raise OrderError("订单包含并非失败状态的采购任务，不能人工重试", status_code=409)
            if any(str(row["hicustom_order_id"] or "").strip() for row in groups):
                raise OrderError("供应商已返回订单号，请先核实，避免重复下单", status_code=409)
            if any(str(row["out_order_id"] or "").strip() for row in groups):
                raise OrderError("任务已有商户订单号，请先同步供应商订单，避免重复下单", status_code=409)
            now = utc_now()
            group_ids = [str(row["id"]) for row in groups]
            placeholders = ",".join("?" for _ in group_ids)
            connection.execute(
                f"""
                UPDATE batch_order_groups
                SET submission_status = 'queued', submission_error = '',
                    hicustom_order_id = '', fulfillment_product_code = '',
                    fulfillment_gallery_codes = '[]', remote_product_id = 0,
                    product_revision = 0, out_order_id = '',
                    product_prepared_at = '', updated_at = ?
                WHERE id IN ({placeholders})
                """,
                (now, *group_ids),
            )
            connection.execute(
                """
                UPDATE ozon_orders
                SET submission_status = 'queued', submission_error = '',
                    hicustom_order_id = '', updated_at = ?
                WHERE posting_number = ?
                """,
                (now, posting_number),
            )
            connection.commit()
        return self.get_order(posting_number)

    @staticmethod
    def _money(value: Any) -> str:
        try:
            amount = Decimal(str(value or "0")).quantize(Decimal("0.01"))
        except (InvalidOperation, TypeError, ValueError):
            amount = Decimal("0.00")
        return format(max(Decimal("0.00"), amount), ".2f")

    @staticmethod
    def _split_money(amount: Decimal, weights: list[int], *, evenly: bool = False) -> list[Decimal]:
        if not weights:
            return []
        cents = int(max(Decimal("0"), amount).quantize(Decimal("0.01")) * 100)
        effective = [1 for _ in weights] if evenly else [max(0, int(value)) for value in weights]
        denominator = sum(effective) or len(effective)
        allocated = [(cents * value) // denominator for value in effective]
        for index in range(cents - sum(allocated)):
            allocated[index % len(allocated)] += 1
        return [Decimal(value) / 100 for value in allocated]

    def record_batch_costs(self, group_id: str, response: dict[str, Any]) -> dict[str, Any]:
        group = self.get_batch_group(group_id)
        product_amount = Decimal(self._money(response.get("subtotal") or response.get("product_amount")))
        if product_amount <= 0:
            product_amount = Decimal(self._money(response.get("platform_subtotal")))
        if product_amount <= 0:
            item_total = Decimal("0")
            for item in response.get("item") or response.get("order_items") or []:
                if not isinstance(item, dict):
                    continue
                total = Decimal(self._money(item.get("total")))
                if total <= 0:
                    total = Decimal(self._money(item.get("price"))) * max(1, int(item.get("qty") or 1))
                item_total += total
            product_amount = item_total.quantize(Decimal("0.01"))
        shipping_amount = Decimal(self._money(response.get("shipping_amount")))
        if shipping_amount <= 0:
            shipping_amount = Decimal(self._money(response.get("platform_shipping_amount")))
        allocated_shipping_amount = shipping_amount
        if allocated_shipping_amount <= 0:
            allocated_shipping_amount = Decimal(self._money(group.get("quoted_shipping_amount")))
        grand_total = Decimal(self._money(response.get("grand_total") or response.get("total_amount")))
        if grand_total <= 0:
            grand_total = Decimal(self._money(response.get("platform_grand_total")))
        if grand_total <= 0:
            grand_total = product_amount + shipping_amount
        sources = list(group.get("source_orders") or [])
        quantities = [max(1, int(item.get("quantity") or 1)) for item in sources]
        product_parts = self._split_money(product_amount, quantities)
        shipping_parts = self._split_money(allocated_shipping_amount, quantities, evenly=True)
        now = utc_now()
        logistics = response.get("logistic_info") if isinstance(response.get("logistic_info"), dict) else {}
        if isinstance(response.get("order_logistics"), dict):
            logistics = response["order_logistics"]
        paid_at = str(response.get("payment_time") or response.get("platform_pay_time") or "").strip()
        tracking_number = official_supplier_tracking_number(group.get("supplier"),
            logistics.get("shipping_track_number")
            or logistics.get("channel_tracking_number")
            or logistics.get("logisticss_track_number")
            or logistics.get("logistics_track_number")
            or ""
        )
        tracking_company = str(
            logistics.get("logistics_company_name")
            or response.get("logistics_platform_text")
            or ""
        ).strip()
        try:
            remote_status = int(response.get("status"))
        except (TypeError, ValueError):
            remote_status = None
        with self._lock, closing(self._connect()) as connection:
            for source, product_cost, shipping_cost in zip(sources, product_parts, shipping_parts):
                connection.execute(
                    """
                    UPDATE batch_order_sources
                    SET product_cost = ?, shipping_cost = ?, total_cost = ?
                    WHERE group_id = ? AND posting_number = ?
                    """,
                    (
                        format(product_cost, ".2f"), format(shipping_cost, ".2f"),
                        format(product_cost + shipping_cost, ".2f"), group_id,
                        str(source["order_key"]),
                    ),
                )
            connection.execute(
                """
                UPDATE batch_order_groups
                SET actual_product_amount = ?, actual_shipping_amount = ?,
                    actual_grand_total = ?, cost_recorded_at = ?, remote_status = ?,
                    paid_at = ?, tracking_number = ?, tracking_company = ?,
                    remote_checked_at = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    format(product_amount, ".2f"), format(shipping_amount, ".2f"),
                    format(grand_total, ".2f"), now, remote_status, paid_at,
                    tracking_number, tracking_company, now, now, group_id,
                ),
            )
            connection.commit()
        return self.get_batch_group(group_id)

    def record_merged_batch_costs(self, purchase_id: str, response: dict[str, Any]) -> list[dict[str, Any]]:
        groups = self.list_purchase_groups(purchase_id)
        if not groups:
            raise OrderError("合并采购任务不存在", status_code=404)
        items = [item for item in (response.get("item") or response.get("order_items") or []) if isinstance(item, dict)]
        item_amounts: dict[str, Decimal] = {}
        for item in items:
            out_item_id = str(item.get("out_item_id") or item.get("third_item_id") or "")
            amount = Decimal(self._money(item.get("total") or item.get("subtotal")))
            if amount <= 0:
                amount = Decimal(self._money(item.get("price"))) * max(1, int(item.get("qty") or item.get("quantity") or 1))
            for group in groups:
                if out_item_id.startswith(str(group["id"])):
                    item_amounts[str(group["id"])] = amount.quantize(Decimal("0.01"))
                    break
        total_product = Decimal(self._money(response.get("subtotal") or response.get("product_amount") or response.get("platform_subtotal")))
        if total_product <= 0:
            total_product = sum(item_amounts.values(), Decimal("0"))
        quantities = [max(1, int(group.get("quantity") or 1)) for group in groups]
        fallback_products = self._split_money(total_product, quantities)
        product_parts = [item_amounts.get(str(group["id"]), fallback_products[index]) for index, group in enumerate(groups)]
        if item_amounts and sum(product_parts) != total_product:
            product_parts = self._split_money(total_product, quantities)
        shipping = Decimal(self._money(response.get("shipping_amount") or response.get("platform_shipping_amount")))
        if shipping <= 0:
            shipping = sum((Decimal(self._money(group.get("quoted_shipping_amount"))) for group in groups), Decimal("0"))
        shipping_parts = self._split_money(shipping, quantities)
        grand_total = Decimal(self._money(response.get("grand_total") or response.get("total_amount") or response.get("platform_grand_total")))
        if grand_total <= 0:
            grand_total = total_product + shipping
        for group, product_part, shipping_part in zip(groups, product_parts, shipping_parts):
            group_response = dict(response)
            group_response.update({
                "subtotal": format(product_part, ".2f"),
                "shipping_amount": format(shipping_part, ".2f"),
                "grand_total": format(product_part + shipping_part, ".2f"),
            })
            self.record_batch_costs(str(group["id"]), group_response)
        return self.list_purchase_groups(purchase_id)

    def batch_payload(self, group_id: str, store_id: int) -> dict[str, Any]:
        group = self.get_batch_group(group_id)
        product_code = str(group.get("fulfillment_product_code") or "").strip()
        out_order_id = str(group.get("out_order_id") or "").strip()
        if not product_code or not out_order_id:
            raise OrderError("订单专用产品尚未生成")
        return {
            "out_order_id": out_order_id,
            "store_id": int(store_id),
            "order_time": hicustom_datetime(group.get("created_at")),
            "remark": f"批量下单 SKU {group['sku']}，来源 {group['source_order_count']} 个平台订单",
            "currency_code": 7,
            "logistics_company_name": group["logistics_company_name"],
            "order_items": [{
                "product_code": product_code,
                "stock_item_id": int(group["stock_item_id"]),
                "qty": int(group["quantity"]),
                "out_item_id": f"{group['id']}-R{int(group.get('product_revision') or 0)}",
                "product_name": group["product_name"],
                "sku": group["sku"],
                "third_size": group["size_name"],
                "third_color": group["color_name"],
            }],
            "address": group["address"],
        }

    def merged_batch_payload(self, purchase_id: str, store_id: int) -> dict[str, Any]:
        groups = self.list_purchase_groups(purchase_id)
        if not groups:
            raise OrderError("合并采购任务不存在", status_code=404)
        if any(not str(group.get("fulfillment_product_code") or "").strip() for group in groups):
            raise OrderError("合并订单中仍有商品尚未生成")
        first = groups[0]
        return {
            "out_order_id": f"BATCH-{first['batch_id'][:10]}-P-{purchase_id[:10]}",
            "store_id": int(store_id),
            "order_time": hicustom_datetime(first.get("created_at")),
            "remark": f"智能合并下单，共 {len(groups)} 个商品项、{sum(int(group['quantity']) for group in groups)} 件",
            "currency_code": 7,
            "logistics_company_name": first["logistics_company_name"],
            "order_items": [{
                "product_code": str(group["fulfillment_product_code"]),
                "stock_item_id": int(group["stock_item_id"]),
                "qty": int(group["quantity"]),
                "out_item_id": f"{group['id']}-R{int(group.get('product_revision') or 0)}",
                "product_name": group["product_name"],
                "sku": group["sku"],
                "third_size": group["size_name"],
                "third_color": group["color_name"],
            } for group in groups],
            "address": first["address"],
        }

    def s2b_batch_payload(self, group_id: str, store_id: int) -> dict[str, Any]:
        group = self.get_batch_group(group_id)
        address = group["address"]
        return {
            "third_order_id": str(group.get("out_order_id") or f"BATCH-{group_id}"),
            # These orders are delivered to a domestic consolidation warehouse,
            # so S2B must use the generic channel.  Ozon (14) only exposes
            # pickup methods for a mainland-China destination.
            "platform": 99,
            "store_id": int(store_id),
            "remark": f"Ozon 批量下单 SKU {group['sku']}，来源 {group['source_order_count']} 个订单",
            "logistics_id": int(group.get("logistics_id") or 0),
            "address": {
                "firstname": address.get("firstname", ""),
                "lastname": address.get("lastname", ""),
                "country": s2b_country_code(address.get("country", "")),
                "province": address.get("region", ""),
                "city": address.get("city", ""),
                "postcode": address.get("postcode", ""),
                "mobile_phone": address.get("mobile_phone", ""),
                "address": " ".join(filter(None, (address.get("area"), address.get("street")))),
            },
            "items": [{
                "third_product_id": group["sku"],
                "product_id": int(group["remote_product_id"]),
                "num": int(group["quantity"]),
                "size_id": int(group["size_id"]),
                "color_id": int(group["color_id"]),
            }],
        }

    def counts(self, *, platform: str = "", store_id: str = "") -> dict[str, int]:
        clauses: list[str] = []
        params: list[Any] = []
        if platform:
            clauses.append("platform = ?")
            params.append(platform)
        if store_id:
            clauses.append("store_id = ?")
            params.append(store_id)
        where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
        with self._lock, closing(self._connect()) as connection:
            row = connection.execute(
                f"""
                SELECT COUNT(*) AS total,
                    SUM(match_status = 'matched') AS matched,
                    SUM(match_status != 'matched') AS unmatched,
                    SUM(submission_status = 'queued') AS queued,
                    SUM(submission_status IN ('preparing_product', 'submitting', 'checking')) AS active,
                    SUM(submission_status = 'submitted') AS submitted,
                    SUM(submission_status IN ('failed', 'unknown') OR tms_status = 'failed') AS failed,
                    SUM(match_status != 'matched' AND NOT {_CANCELLED_SQL}) AS pending_match,
                    SUM(match_status = 'matched' AND submission_status NOT IN ('submitted') AND tms_status NOT IN ('submitted','completed','failed','manual_tracking') AND NOT {_CANCELLED_SQL}) AS pending_purchase,
                    SUM(platform = 'ozon' AND submission_status = 'submitted' AND tms_status NOT IN ('submitted','completed','failed') AND NOT {_CANCELLED_SQL}) AS purchased_pending_tms,
                    SUM(platform = 'wb' AND submission_status = 'submitted' AND tms_status NOT IN ('submitted','completed','failed') AND NOT {_CANCELLED_SQL}) AS purchased_in_transit,
                    SUM((submission_status = 'submitted' AND tms_status NOT IN ('submitted','completed','failed') OR tms_status = 'manual_tracking') AND NOT {_CANCELLED_SQL}) AS post_purchase_processing,
                    SUM({_CANCELLED_SQL}) AS cancelled,
                    SUM({_CANCELLED_SQL} AND submission_status = 'submitted') AS cancelled_purchased,
                    SUM(tms_status IN ('submitted','completed')) AS completed,
                    SUM(tms_status = 'failed') AS tms_error
                FROM ozon_orders o
                {where}
                """,
                params,
            ).fetchone()
        return {key: int(row[key] or 0) for key in row.keys()}

    def list_orders(
        self, *, search: str = "", match_status: str = "", submission_status: str = "",
        platform: str = "", store_id: str = "", workflow_status: str = "",
        tracking_status: str = "", limit: int = 200, offset: int = 0,
    ) -> dict[str, Any]:
        clauses: list[str] = []
        params: list[Any] = []
        if search.strip():
            clauses.append(
                "(o.posting_number LIKE ? OR o.tms_tail_tracking_numbers LIKE ? OR o.manual_tracking_number LIKE ? "
                "OR EXISTS (SELECT 1 FROM ozon_order_items i WHERE i.posting_number = o.posting_number "
                "AND (i.offer_id LIKE ? OR i.product_name LIKE ?)) "
                "OR EXISTS (SELECT 1 FROM batch_order_sources s JOIN batch_order_groups g ON g.id = s.group_id "
                "WHERE s.posting_number = o.posting_number AND g.tracking_number LIKE ?))"
            )
            needle = f"%{search.strip()}%"
            params.extend([needle] * 6)
        if match_status:
            clauses.append("o.match_status = ?")
            params.append(match_status)
        if submission_status:
            clauses.append("o.submission_status = ?")
            params.append(submission_status)
        if platform:
            clauses.append("o.platform = ?")
            params.append(platform)
        if store_id:
            clauses.append("o.store_id = ?")
            params.append(store_id)
        tracking_exists = """EXISTS (
            SELECT 1 FROM batch_order_sources ts
            JOIN batch_order_groups tg ON tg.id = ts.group_id
            WHERE ts.posting_number = o.posting_number
              AND TRIM(COALESCE(tg.tracking_number, '')) <> ''
        )"""
        if tracking_status == "ready":
            clauses.append(f"({tracking_exists} OR TRIM(COALESCE(o.manual_tracking_number, '')) <> '')")
        elif tracking_status == "waiting":
            clauses.append(f"NOT {tracking_exists} AND TRIM(COALESCE(o.manual_tracking_number, '')) = ''")
        workflow_clauses = {
            "pending_match": f"o.match_status != 'matched' AND NOT {_CANCELLED_SQL}",
            "pending_purchase": f"o.match_status = 'matched' AND o.submission_status != 'submitted' AND o.tms_status NOT IN ('submitted','completed','failed','manual_tracking') AND NOT {_CANCELLED_SQL}",
            "purchased_pending_tms": f"o.platform = 'ozon' AND o.submission_status = 'submitted' AND o.tms_status NOT IN ('submitted','completed','failed') AND NOT {_CANCELLED_SQL}",
            "purchased_in_transit": f"o.platform = 'wb' AND o.submission_status = 'submitted' AND o.tms_status NOT IN ('submitted','completed','failed') AND NOT {_CANCELLED_SQL}",
            "post_purchase_processing": f"(o.submission_status = 'submitted' AND o.tms_status NOT IN ('submitted','completed','failed') OR o.tms_status = 'manual_tracking') AND NOT {_CANCELLED_SQL}",
            "cancelled": _CANCELLED_SQL,
            "completed": "o.tms_status IN ('submitted','completed')",
            "tms_error": "(o.tms_status = 'failed' OR o.submission_status IN ('failed','unknown'))",
        }
        if workflow_status in workflow_clauses:
            clauses.append(workflow_clauses[workflow_status])
        where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
        with self._lock, closing(self._connect()) as connection:
            filtered_total = int(connection.execute(f"SELECT COUNT(*) FROM ozon_orders o {where}", params).fetchone()[0])
            rows = connection.execute(
                f"SELECT o.* FROM ozon_orders o {where} ORDER BY o.in_process_at DESC, o.updated_at DESC LIMIT ? OFFSET ?",
                (*params, max(1, min(500, limit)), max(0,offset)),
            ).fetchall()
            values: list[dict[str, Any]] = []
            for row in rows:
                item_rows = connection.execute(
                    "SELECT * FROM ozon_order_items WHERE posting_number = ? ORDER BY line_no",
                    (row["posting_number"],),
                ).fetchall()
                items = [dict(item) for item in item_rows]
                values.append(self._public_order(row, items))
        return {"orders": values, "total":filtered_total, "counts": self.counts(platform=platform, store_id=store_id)}

    def get_order(self, posting_number: str) -> dict[str, Any]:
        with self._lock, closing(self._connect()) as connection:
            row = connection.execute(
                "SELECT * FROM ozon_orders WHERE posting_number = ?", (posting_number,)
            ).fetchone()
            if row is None:
                raise OrderError("订单不存在", status_code=404)
            items = [
                dict(item)
                for item in connection.execute(
                    "SELECT * FROM ozon_order_items WHERE posting_number = ? ORDER BY line_no",
                    (posting_number,),
                ).fetchall()
            ]
            result = self._public_order(row, items)
            result["recipient"] = self._decrypt_recipient(row["recipient_enc"])
            return result

    def update_recipient(self, posting_number: str, recipient: dict[str, Any]) -> dict[str, Any]:
        clean = {key: str(recipient.get(key) or "").strip() for key in (
            "country", "region", "city", "area", "street", "postcode", "firstname",
            "lastname", "mobile_phone", "telephone", "email"
        )}
        with self._lock, closing(self._connect()) as connection:
            if connection.execute(
                "SELECT 1 FROM ozon_orders WHERE posting_number = ?", (posting_number,)
            ).fetchone() is None:
                raise OrderError("订单不存在", status_code=404)
            connection.execute(
                "UPDATE ozon_orders SET recipient_enc = ?, updated_at = ? WHERE posting_number = ?",
                (self._encrypt_recipient(clean), utc_now(), posting_number),
            )
            connection.commit()
        return self.get_order(posting_number)

    def queue_submission(self, posting_number: str) -> dict[str, Any]:
        order = self.get_order(posting_number)
        if order["submission_status"] == "submitted":
            raise OrderError("该订单已经提交到 HICUSTOM，不能重复下单", status_code=409)
        if order["submission_status"] in {"queued", "preparing_product", "checking", "submitting"}:
            raise OrderError("该订单已在下单队列中", status_code=409)
        if order["match_status"] != "matched":
            raise OrderError("订单中仍有商品未匹配产品库")
        if not order["recipient_ready"]:
            raise OrderError("收件信息不完整，请先补充地址")
        with self._lock, closing(self._connect()) as connection:
            connection.execute(
                "UPDATE ozon_orders SET submission_status = 'queued', submission_error = '', updated_at = ? WHERE posting_number = ?",
                (utc_now(), posting_number),
            )
            connection.commit()
        return self.get_order(posting_number)

    def next_queued(self) -> dict[str, Any] | None:
        with self._lock, closing(self._connect()) as connection:
            row = connection.execute(
                """
                SELECT o.posting_number FROM ozon_orders o
                WHERE o.submission_status = 'queued'
                  AND NOT EXISTS (
                    SELECT 1 FROM batch_order_sources s
                    JOIN batch_order_groups g ON g.id = s.group_id
                    WHERE s.posting_number = o.posting_number
                      AND g.submission_status IN ('queued', 'preparing_product', 'checking', 'submitting')
                  )
                ORDER BY o.updated_at LIMIT 1
                """
            ).fetchone()
        return self.get_order(str(row["posting_number"])) if row else None

    def update_submission(self, posting_number: str, **changes: Any) -> dict[str, Any]:
        allowed = {"submission_status", "hicustom_order_id", "submission_error"}
        values = {key: value for key, value in changes.items() if key in allowed}
        if not values:
            return self.get_order(posting_number)
        assignments = ", ".join(f"{key} = ?" for key in values)
        with self._lock, closing(self._connect()) as connection:
            connection.execute(
                f"UPDATE ozon_orders SET {assignments}, updated_at = ? WHERE posting_number = ?",
                (*values.values(), utc_now(), posting_number),
            )
            connection.commit()
        return self.get_order(posting_number)

    def submission_payload(self, posting_number: str, store_id: int) -> dict[str, Any]:
        order = self.get_order(posting_number)
        items = []
        for item in order["items"]:
            items.append(
                {
                    "product_code": item["product_code"],
                    "stock_item_id": int(item["stock_item_id"]),
                    "qty": int(item["quantity"]),
                    "out_item_id": f"{posting_number}:{item['line_no']}",
                    "product_name": item["product_name"],
                    "sku": item["matched_sku"],
                    "third_size": item["size_name"],
                    "third_color": item["color_name"],
                }
            )
        return {
            "out_order_id": str(order.get("order_key") or posting_number),
            "store_id": int(store_id),
            "order_time": hicustom_datetime(order.get("in_process_at")),
            "remark": f"{str(order.get('platform') or 'ozon').upper()} {order.get('fulfillment_mode') or ''} {order.get('posting_number') or posting_number}",
            "order_items": items,
            "address": order["recipient"],
        }


class HicustomOrderQueue:
    def __init__(
        self,
        orders: OrderManager,
        settings_loader: Callable[[], dict[str, Any]],
        client_loader: Callable[[], Any],
    ) -> None:
        self.orders = orders
        self.settings_loader = settings_loader
        self.client_loader = client_loader
        self._wake = asyncio.Event()
        self._worker: asyncio.Task[None] | None = None

    async def start(self) -> None:
        raise OrderError("旧采购 Worker 已禁用；不得自动重放历史采购，需使用经验证的 ERP 任务适配器", status_code=409)

    async def stop(self) -> None:
        if self._worker:
            self._worker.cancel()
            await asyncio.gather(self._worker, return_exceptions=True)
            self._worker = None

    def wake(self) -> None:
        self._wake.set()

    async def _client_for(self, supplier: str = "hicustom") -> Any:
        try:
            return await self.client_loader(supplier)
        except TypeError:
            return await self.client_loader()

    async def sync_remote_orders(
        self, *, limit: int = 500, supplier: str = "", remote_reader=None,
        groups_snapshot=None, strict_apply: bool = False
    ) -> dict[str, Any]:
        """Refresh payment and domestic tracking for submitted batch orders."""
        groups = groups_snapshot if groups_snapshot is not None else await asyncio.to_thread(self.orders.list_batch_groups, limit=limit)
        targets = [
            group for group in groups
            if str(group.get("submission_status") or "") in {"submitted", "unknown"}
            and str(group.get("out_order_id") or "").strip()
            and (not supplier or str(group.get("supplier") or "hicustom") == supplier)
        ]
        if not targets:
            return {
                "checked": 0, "updated": 0, "paid": 0, "tracking": 0,
                "missing": 0, "failed": 0, "errors": [],
            }
        result: dict[str, Any] = {
            "checked": 0, "updated": 0, "paid": 0, "tracking": 0,
            "missing": 0, "failed": 0, "errors": [],
        }
        seen_remote_orders: set[tuple[str, str]] = set()
        for group in targets:
            supplier = str(group.get("supplier") or "hicustom")
            group_id = str(group.get("id") or "")
            out_order_id = str(group.get("out_order_id") or "").strip()
            remote_key = (supplier, out_order_id)
            if remote_key in seen_remote_orders:
                continue
            seen_remote_orders.add(remote_key)
            result["checked"] += 1
            remote_loaded = False
            try:
                if remote_reader is None:
                    client = await self._client_for(supplier)
                    remote = await client.order_by_out_order_id(out_order_id)
                else:
                    # Shared-data hosts apply facts fetched on the initiating PC.
                    # All original cost, merged-group and tracking rules below
                    # are identical to the local ERP path.
                    remote = await remote_reader(supplier, out_order_id)
                remote_loaded = True
                if not remote:
                    result["missing"] += 1
                    continue
                remote_error = self._remote_order_error(remote)
                if remote_error:
                    purchase_id = str(group.get("purchase_id") or "")
                    if supplier == "hicustom" and purchase_id and len(await asyncio.to_thread(self.orders.list_purchase_groups, purchase_id)) > 1:
                        await asyncio.to_thread(
                            self.orders.update_purchase_groups, purchase_id,
                            submission_status="failed",
                            submission_error=f"HICUSTOM 订单异常：{remote_error}",
                        )
                    else:
                        await asyncio.to_thread(
                            self.orders.update_batch_group,
                            group_id,
                            submission_status="failed",
                            submission_error=f"{'S2B' if supplier == 's2b' else 'HICUSTOM'} 订单异常：{remote_error}",
                        )
                    result["failed"] += 1
                    result["errors"].append(f"{out_order_id}：{remote_error}")
                    continue
                purchase_id = str(group.get("purchase_id") or "")
                purchase_groups = (
                    await asyncio.to_thread(self.orders.list_purchase_groups, purchase_id)
                    if supplier == "hicustom" and purchase_id else []
                )
                if len(purchase_groups) > 1:
                    refreshed_groups = await asyncio.to_thread(
                        self.orders.record_merged_batch_costs, purchase_id, remote
                    )
                    await asyncio.to_thread(
                        self.orders.update_purchase_groups, purchase_id,
                        submission_status="submitted",
                        hicustom_order_id=str(
                            remote.get("order_id") or remote.get("id")
                            or group.get("hicustom_order_id") or ""
                        ),
                        submission_error="",
                    )
                    refreshed = refreshed_groups[0]
                else:
                    refreshed = await asyncio.to_thread(
                        self.orders.record_batch_costs, group_id, remote
                    )
                    await asyncio.to_thread(
                        self.orders.update_batch_group,
                        group_id,
                        submission_status="submitted",
                        hicustom_order_id=str(
                            remote.get("order_id") or remote.get("id")
                            or group.get("hicustom_order_id") or ""
                        ),
                        submission_error="",
                    )
                result["updated"] += 1
                result["paid"] += int(bool(str(refreshed.get("paid_at") or "").strip()))
                result["tracking"] += int(bool(str(refreshed.get("tracking_number") or "").strip()))
            except Exception as error:  # noqa: BLE001
                if strict_apply and remote_loaded:
                    raise
                result["failed"] += 1
                result["errors"].append(f"{out_order_id}：{error}")
        result["errors"] = result["errors"][:20]
        return result

    async def worker_loop(self) -> None:
        while True:
            batch_group = await asyncio.to_thread(self.orders.next_batch_queued)
            if batch_group is not None:
                group_id = str(batch_group["id"])
                try:
                    await self.process_batch(group_id)
                except asyncio.CancelledError:
                    raise
                except Exception as error:  # noqa: BLE001
                    await asyncio.to_thread(
                        self.orders.update_batch_group,
                        group_id,
                        submission_status="failed",
                        submission_error=f"{type(error).__name__}: {error}"[:1000],
                    )
                continue
            order = await asyncio.to_thread(self.orders.next_queued)
            if order is None:
                self._wake.clear()
                if await asyncio.to_thread(self.orders.next_batch_queued) or await asyncio.to_thread(self.orders.next_queued):
                    self._wake.set()
                    continue
                await self._wake.wait()
                continue
            number = str(order.get("order_key") or order["posting_number"])
            try:
                await self.process(number)
            except asyncio.CancelledError:
                raise
            except Exception as error:  # noqa: BLE001
                await asyncio.to_thread(
                    self.orders.update_submission,
                    number,
                    submission_status="failed",
                    submission_error=f"{type(error).__name__}: {error}"[:1000],
                )

    @staticmethod
    def _remote_order_error(order: dict[str, Any]) -> str:
        value = str(order.get("error_msg") or order.get("error_message") or "").strip()
        for source, replacement in (("<br/>", "；"), ("<br>", "；"), ("<br />", "；")):
            value = value.replace(source, replacement)
        return value[:1000]

    @staticmethod
    def _source_asset(product: dict[str, Any], view_id: int) -> dict[str, Any]:
        for asset in product.get("oss_assets") or []:
            if not isinstance(asset, dict) or asset.get("asset_type") != "source":
                continue
            if int(asset.get("view_id") or 0) == view_id:
                return asset
        return {}

    async def _source_content(
        self, product: dict[str, Any], side: dict[str, Any]
    ) -> tuple[bytes, str, str]:
        view_id = int(side.get("view_id") or 0)
        mime_type = str(side.get("mime_type") or "image/png").lower()
        filename = str(side.get("source_filename") or f"original-{view_id}.png")
        candidates: list[Path] = []
        source_path = str(side.get("source_path") or "").strip()
        if source_path:
            candidates.append(Path(source_path))
        output_folder = Path(str(product.get("output_folder") or ""))
        asset = self._source_asset(product, view_id)
        local_path = str(asset.get("local_path") or "").strip()
        if local_path and str(output_folder):
            candidates.append(output_folder / local_path)
        candidates.extend(
            output_folder / "originals" / f"view-{view_id}" / name
            for name in ("original.png", "original.jpg", "original.jpeg")
        )
        for path in candidates:
            if path.is_file():
                content = await asyncio.to_thread(path.read_bytes)
                if not content or len(content) > 30 * 1024 * 1024:
                    raise OrderError(f"SKU {product.get('sku')} 印刷面 {view_id} 的原图文件无效")
                suffix = path.suffix.lower()
                if suffix in {".jpg", ".jpeg"}:
                    mime_type = "image/jpeg"
                elif suffix == ".png":
                    mime_type = "image/png"
                return content, path.name, mime_type
        oss_url = str(asset.get("oss_url") or "").strip()
        if oss_url:
            try:
                async with httpx.AsyncClient(timeout=90, follow_redirects=True) as http:
                    response = await http.get(oss_url)
                    response.raise_for_status()
            except httpx.HTTPError as error:
                raise OrderError(f"SKU {product.get('sku')} 无法从 OSS 恢复印刷面 {view_id}") from error
            if not response.content or len(response.content) > 30 * 1024 * 1024:
                raise OrderError(f"SKU {product.get('sku')} 的 OSS 原图文件无效")
            return response.content, filename, str(asset.get("content_type") or mime_type)
        raise OrderError(f"SKU {product.get('sku')} 缺少印刷面 {view_id} 的本地或 OSS 原图")

    async def _prepare_batch_product(
        self, group_id: str, client: Any, settings: dict[str, Any]
    ) -> dict[str, Any]:
        group = await asyncio.to_thread(self.orders.get_batch_group, group_id)
        reusable = await asyncio.to_thread(
            self.orders.reusable_hicustom_product,
            str(group["sku"]),
            exclude_group_id=group_id,
        )
        if reusable:
            try:
                reusable_gallery_codes = [
                    str(code).strip()
                    for code in json.loads(str(reusable.get("fulfillment_gallery_codes") or "[]"))
                    if str(code).strip()
                ]
            except (TypeError, ValueError, json.JSONDecodeError):
                reusable_gallery_codes = []
            gallery_checks = [
                await client.gallery_exists(code) for code in reusable_gallery_codes
            ]
            if reusable_gallery_codes and all(gallery_checks):
                revision = max(1, int(reusable.get("product_revision") or 1))
                return await asyncio.to_thread(
                    self.orders.update_batch_group,
                    group_id,
                    fulfillment_product_code=str(reusable["fulfillment_product_code"]),
                    fulfillment_gallery_codes=json.dumps(reusable_gallery_codes, ensure_ascii=False),
                    product_revision=revision,
                    out_order_id=f"BATCH-{group['batch_id'][:10]}-{group_id[:10]}-R{revision}",
                    product_prepared_at=utc_now(),
                )
        product = await asyncio.to_thread(self.orders.library.get_product, str(group["sku"]))
        recipe = product.get("recipe") if isinstance(product.get("recipe"), dict) else {}
        if str(recipe.get("supplier") or "hicustom") != "hicustom":
            raise OrderError(f"SKU {group['sku']} 不是指纹科技产品，不能提交到 HICUSTOM")
        sides = [dict(side) for side in recipe.get("sides") or [] if isinstance(side, dict)]
        if not sides:
            raise OrderError(f"SKU {group['sku']} 没有保存可复建的印刷面")
        revision = int(group.get("product_revision") or 0) + 1
        external_id = f"order-{group_id[:12]}-r{revision}"
        gallery_codes: list[str] = []
        for side in sides:
            content, filename, mime_type = await self._source_content(product, side)
            gallery = await client.upload_gallery(
                filename=filename,
                content=content,
                mime_type=mime_type,
                external_id=f"{external_id}-view-{int(side.get('view_id') or 0)}",
                external_customer_id=str(settings.get("hicustom_external_customer_id") or ""),
            )
            gallery_code = str(gallery.get("code") or "").strip()
            if not gallery_code:
                raise OrderError("HICUSTOM 图库上传响应缺少图片编码", outcome_unknown=True)
            side["gallery_code"] = gallery_code
            side["mime_type"] = mime_type
            gallery_codes.append(gallery_code)
        job = {
            "external_id": external_id,
            "external_customer_id": str(settings.get("hicustom_external_customer_id") or ""),
            "product_type_id": int(recipe.get("product_type_id") or product.get("product_type_id") or 0),
            "default_color_id": int(recipe.get("default_color_id") or product.get("default_color_id") or 0),
            "default_view_id": int(recipe.get("default_view_id") or product.get("default_view_id") or 0),
            "sides": sides,
        }
        created = await client.create_product(job)
        product_code = str(created.get("code") or "").strip()
        if not product_code:
            raise OrderError("HICUSTOM 订单专用产品响应缺少产品编码", outcome_unknown=True)
        out_order_id = f"BATCH-{group['batch_id'][:10]}-{group_id[:10]}-R{revision}"
        return await asyncio.to_thread(
            self.orders.update_batch_group,
            group_id,
            fulfillment_product_code=product_code,
            fulfillment_gallery_codes=json.dumps(gallery_codes, ensure_ascii=False),
            product_revision=revision,
            out_order_id=out_order_id,
            product_prepared_at=utc_now(),
        )

    async def _prepare_s2b_batch_product(
        self, group_id: str, client: Any, settings: dict[str, Any]
    ) -> dict[str, Any]:
        group = await asyncio.to_thread(self.orders.get_batch_group, group_id)
        reusable = await asyncio.to_thread(
            self.orders.reusable_s2b_product,
            str(group["sku"]),
            exclude_group_id=group_id,
        )
        if reusable:
            remote_product_id = int(reusable.get("remote_product_id") or 0)
            try:
                remote_product = await client.custom_product(remote_product_id)
            except OrderError as error:
                if error.status_code != 404:
                    raise
                remote_product = {}
            if remote_product:
                try:
                    reusable_material_ids = [
                        str(value).strip()
                        for value in json.loads(str(reusable.get("fulfillment_gallery_codes") or "[]"))
                        if str(value).strip()
                    ]
                except (TypeError, ValueError, json.JSONDecodeError):
                    reusable_material_ids = []
                if reusable_material_ids:
                    revision = max(1, int(reusable.get("product_revision") or 1))
                    group = await asyncio.to_thread(
                        self.orders.update_batch_group, group_id,
                        remote_product_id=remote_product_id,
                        fulfillment_product_code=str(reusable["fulfillment_product_code"]),
                        fulfillment_gallery_codes=json.dumps(reusable_material_ids, ensure_ascii=False),
                        product_revision=revision,
                        out_order_id=f"BATCH-{group['batch_id'][:10]}-{group_id[:10]}-R{revision}",
                        product_prepared_at=utc_now(),
                    )
                    options = await client.calculate_products(
                        products=[{
                            "product_id": remote_product_id,
                            "stock_sku_item_id": int(group.get("stock_item_id") or 0),
                            "num": int(group.get("quantity") or 1),
                        }],
                        address=group["address"], platform=99,
                    )
                    if options:
                        return await self._apply_s2b_logistics(group_id, group, options)
        product = await asyncio.to_thread(self.orders.library.get_product, str(group["sku"]))
        recipe = product.get("recipe") if isinstance(product.get("recipe"), dict) else {}
        if str(recipe.get("supplier") or product.get("supplier") or "hicustom") != "s2b":
            raise OrderError(f"SKU {group['sku']} 不是 S2B 产品")
        sides = [dict(side) for side in recipe.get("sides") or [] if isinstance(side, dict)]
        if not sides:
            raise OrderError(f"SKU {group['sku']} 没有保存可复建的印刷面")
        revision = int(group.get("product_revision") or 0) + 1
        external_id = f"s2b-order-{group_id[:12]}-r{revision}"
        material_ids: list[str] = []
        for side in sides:
            content, filename, mime_type = await self._source_content(product, side)
            gallery = await client.upload_gallery(
                filename=filename, content=content, mime_type=mime_type,
                external_id=f"{external_id}-view-{int(side.get('view_id') or 0)}",
                external_customer_id="",
            )
            material_id = str(gallery.get("code") or "").strip()
            if not material_id:
                raise OrderError("S2B 素材上传响应缺少素材编号", outcome_unknown=True)
            side["gallery_code"] = material_id
            side["mime_type"] = mime_type
            material_ids.append(material_id)
        created = await client.create_product({
            "product_type_id": int(recipe.get("product_type_id") or product.get("product_type_id") or 0),
            "default_color_id": int(group.get("color_id") or recipe.get("default_color_id") or 0),
            "selected_size_ids": [int(group.get("size_id") or recipe.get("default_size_id") or 0)],
            "product_name": str(product.get("product_name") or recipe.get("product_name") or group.get("product_name") or group["sku"]),
            "sides": sides,
        })
        product_id = int(created.get("id") or 0)
        product_code = str(created.get("code") or "").strip()
        if product_id <= 0 or not product_code:
            raise OrderError("S2B 订单专用产品响应缺少产品编号", outcome_unknown=True)
        out_order_id = f"BATCH-{group['batch_id'][:10]}-{group_id[:10]}-R{revision}"
        group = await asyncio.to_thread(
            self.orders.update_batch_group, group_id,
            remote_product_id=product_id,
            fulfillment_product_code=product_code,
            fulfillment_gallery_codes=json.dumps(material_ids, ensure_ascii=False),
            product_revision=revision, out_order_id=out_order_id,
            product_prepared_at=utc_now(),
        )
        options = await client.calculate_products(
            products=[{
                "product_id": product_id,
                "stock_sku_item_id": int(group.get("stock_item_id") or 0),
                "num": int(group.get("quantity") or 1),
            }],
            address=group["address"], platform=99,
        )
        if not options:
            raise OrderError("S2B 没有返回可用物流渠道")
        return await self._apply_s2b_logistics(group_id, group, options)

    async def _apply_s2b_logistics(
        self, group_id: str, group: dict[str, Any], options: list[dict[str, Any]]
    ) -> dict[str, Any]:
        preferred = str(group.get("logistics_company_name") or "").strip().casefold()
        def option_name(item: dict[str, Any]) -> str:
            return str(item.get("name") or item.get("logistics_name") or item.get("en_name") or item.get("full_en_name") or "").strip()
        requested_logistics_id = int(group.get("logistics_id") or 0)
        def option_ids(item: dict[str, Any]) -> set[int]:
            result: set[int] = set()
            for key in ("logistics_id", "logistics_platform_id", "id"):
                try: result.add(int(item.get(key) or 0))
                except (TypeError, ValueError): continue
            return result
        matched = [item for item in options if requested_logistics_id > 0 and requested_logistics_id in {
            *option_ids(item),
        }]
        if not matched:
            matched = [item for item in options if preferred and preferred in option_name(item).casefold()]
        # calculateProducts returns delivery modes for a mainland consolidation
        # address (pickup/home delivery), not always the carrier catalogue. A
        # previously selected S2B carrier id is authoritative; e.g. 109 is YTO.
        if not matched and requested_logistics_id > 0:
            quoted = group.get("quoted_shipping_amount")
            if requested_logistics_id == 109 and Decimal(str(quoted or "0")) <= 0:
                quoted = "5.00"
            return await asyncio.to_thread(
                self.orders.update_batch_group, group_id,
                logistics_id=requested_logistics_id,
                logistics_company_name=str(group.get("logistics_company_name") or requested_logistics_id),
                quoted_shipping_amount=str(quoted or "0.00"),
            )
        if preferred and not matched:
            available = "、".join(filter(None, (option_name(item) for item in options)))
            raise OrderError(
                f"S2B 未返回指定物流“{group.get('logistics_company_name')}”"
                f"；当前可用：{available or '无'}"
            )
        candidates = matched or options
        def option_amount(item: dict[str, Any]) -> Decimal:
            for key in ("amount", "min_amount", "max_amount"):
                try:
                    return Decimal(str(item.get(key) or "0"))
                except InvalidOperation:
                    continue
            return Decimal("999999999")
        selected = min(candidates, key=option_amount)
        logistics_id = int(selected.get("logistics_id") or selected.get("id") or 0)
        if logistics_id <= 0:
            raise OrderError("S2B 物流报价缺少渠道编号")
        return await asyncio.to_thread(
            self.orders.update_batch_group, group_id,
            logistics_id=logistics_id,
            logistics_company_name=option_name(selected) or str(logistics_id),
            quoted_shipping_amount=format(option_amount(selected), ".2f"),
        )

    async def _process_merged_hicustom_batch(
        self, purchase_id: str, store_id: int, client: Any, settings: dict[str, Any]
    ) -> None:
        groups = await asyncio.to_thread(self.orders.list_purchase_groups, purchase_id)
        await asyncio.to_thread(self.orders.assert_purchase_sources_active, [str(group["id"]) for group in groups])
        await asyncio.to_thread(
            self.orders.update_purchase_groups, purchase_id,
            submission_status="preparing_product", submission_error="",
        )
        for group in groups:
            if not str(group.get("fulfillment_product_code") or "").strip():
                await self._prepare_batch_product(str(group["id"]), client, settings)
        payload = await asyncio.to_thread(self.orders.merged_batch_payload, purchase_id, store_id)
        out_order_id = str(payload["out_order_id"])
        await asyncio.to_thread(
            self.orders.update_purchase_groups, purchase_id,
            submission_status="checking", submission_error="", out_order_id=out_order_id,
        )
        existing = await client.order_by_out_order_id(out_order_id)
        if existing:
            remote_error = self._remote_order_error(existing)
            if remote_error:
                await asyncio.to_thread(
                    self.orders.update_purchase_groups, purchase_id,
                    submission_status="failed",
                    hicustom_order_id=str(existing.get("order_id") or existing.get("id") or ""),
                    submission_error=f"HICUSTOM 订单异常，无法付款：{remote_error}",
                )
                return
            await asyncio.to_thread(self.orders.record_merged_batch_costs, purchase_id, existing)
            await asyncio.to_thread(
                self.orders.update_purchase_groups, purchase_id,
                submission_status="submitted",
                hicustom_order_id=str(existing.get("order_id") or existing.get("id") or ""),
                submission_error="",
            )
            return
        await asyncio.to_thread(
            self.orders.update_purchase_groups, purchase_id, submission_status="submitting"
        )
        await asyncio.to_thread(self.orders.assert_purchase_sources_active, [str(group["id"]) for group in groups])
        try:
            created = await client.create_order(payload)
        except Exception as error:
            if getattr(error, "outcome_unknown", False):
                await asyncio.to_thread(
                    self.orders.update_purchase_groups, purchase_id,
                    submission_status="unknown", submission_error=str(error),
                )
                return
            raise
        remote_error = self._remote_order_error(created)
        if remote_error:
            await asyncio.to_thread(
                self.orders.update_purchase_groups, purchase_id,
                submission_status="failed",
                hicustom_order_id=str(created.get("order_id") or created.get("id") or ""),
                submission_error=f"HICUSTOM 订单异常，无法付款：{remote_error}",
            )
            return
        await asyncio.to_thread(self.orders.record_merged_batch_costs, purchase_id, created)
        await asyncio.to_thread(
            self.orders.update_purchase_groups, purchase_id,
            submission_status="submitted",
            hicustom_order_id=str(created.get("order_id") or created.get("id") or ""),
            submission_error="",
        )

    async def process_batch(self, group_id: str) -> None:
        settings = self.settings_loader()
        group = await asyncio.to_thread(self.orders.get_batch_group, group_id)
        await asyncio.to_thread(self.orders.assert_purchase_sources_active, [group_id])
        supplier = str(group.get("supplier") or "hicustom")
        provider_name = "S2B" if supplier == "s2b" else "HICUSTOM"
        try:
            store_id = int(settings.get("s2b_store_id") if supplier == "s2b" else settings.get("hicustom_store_id") or 0)
        except (TypeError, ValueError):
            store_id = 0
        if store_id <= 0:
            raise OrderError(f"请先在设置中填写 {'S2B' if supplier == 's2b' else 'HICUSTOM'} 店铺编号")
        client = await self._client_for(supplier)
        purchase_id = str(group.get("purchase_id") or "")
        if supplier == "hicustom" and purchase_id:
            purchase_groups = await asyncio.to_thread(self.orders.list_purchase_groups, purchase_id)
            if len(purchase_groups) > 1:
                await self._process_merged_hicustom_batch(purchase_id, store_id, client, settings)
                return
        if supplier == "s2b" and (
            not str(group.get("fulfillment_product_code") or "").strip()
            or int(group.get("remote_product_id") or 0) <= 0
        ):
            await asyncio.to_thread(
                self.orders.update_batch_group, group_id,
                submission_status="preparing_product", submission_error="",
            )
            group = await self._prepare_s2b_batch_product(group_id, client, settings)
        if supplier == "hicustom" and not str(group.get("fulfillment_product_code") or "").strip():
            await asyncio.to_thread(
                self.orders.update_batch_group,
                group_id,
                submission_status="preparing_product",
                submission_error="",
            )
            group = await self._prepare_batch_product(group_id, client, settings)
        group = await asyncio.to_thread(
            self.orders.update_batch_group, group_id, submission_status="checking", submission_error=""
        )
        out_order_id = str(group.get("out_order_id") or "")
        existing = await client.order_by_out_order_id(out_order_id)
        if existing:
            remote_error = self._remote_order_error(existing)
            if remote_error:
                await asyncio.to_thread(
                    self.orders.update_batch_group,
                    group_id,
                    submission_status="failed",
                    hicustom_order_id=str(existing.get("order_id") or existing.get("id") or ""),
                    submission_error=f"{provider_name} 订单异常，无法付款：{remote_error}",
                )
                return
            await asyncio.to_thread(self.orders.record_batch_costs, group_id, existing)
            await asyncio.to_thread(
                self.orders.update_batch_group,
                group_id,
                submission_status="submitted",
                hicustom_order_id=str(existing.get("order_id") or existing.get("id") or ""),
                submission_error="",
            )
            return
        payload = await asyncio.to_thread(
            self.orders.s2b_batch_payload if supplier == "s2b" else self.orders.batch_payload,
            group_id, store_id,
        )
        await asyncio.to_thread(
            self.orders.update_batch_group, group_id, submission_status="submitting"
        )
        await asyncio.to_thread(self.orders.assert_purchase_sources_active, [group_id])
        try:
            created = await client.create_order(payload)
        except Exception as error:
            if getattr(error, "outcome_unknown", False):
                await asyncio.to_thread(
                    self.orders.update_batch_group,
                    group_id,
                    submission_status="unknown",
                    submission_error=str(error),
                )
                return
            raise
        remote_error = self._remote_order_error(created)
        if remote_error:
            await asyncio.to_thread(
                self.orders.update_batch_group,
                group_id,
                submission_status="failed",
                hicustom_order_id=str(created.get("order_id") or created.get("id") or ""),
                submission_error=f"{provider_name} 订单异常，无法付款：{remote_error}",
            )
            return
        await asyncio.to_thread(self.orders.record_batch_costs, group_id, created)
        await asyncio.to_thread(
            self.orders.update_batch_group,
            group_id,
            submission_status="submitted",
            hicustom_order_id=str(created.get("order_id") or created.get("id") or ""),
            submission_error="",
        )

    async def process(self, posting_number: str) -> None:
        order = await asyncio.to_thread(self.orders.get_order, posting_number)
        if is_platform_order_cancelled(order) or order.get("tms_status") == "manual_tracking":
            raise OrderError("平台订单已取消或已人工录入物流单号，禁止提交供应商采购", status_code=409)
        settings = self.settings_loader()
        try:
            store_id = int(settings.get("hicustom_store_id") or 0)
        except (TypeError, ValueError):
            store_id = 0
        if store_id <= 0:
            raise OrderError("请先在设置中填写 HICUSTOM 店铺编号")
        client = await self._client_for("hicustom")
        await asyncio.to_thread(
            self.orders.update_submission,
            posting_number,
            submission_status="checking",
            submission_error="",
        )
        existing = await client.order_by_out_order_id(posting_number)
        if existing:
            order_id = str(existing.get("order_id") or existing.get("id") or "")
            await asyncio.to_thread(
                self.orders.update_submission,
                posting_number,
                submission_status="submitted",
                hicustom_order_id=order_id,
                submission_error="",
            )
            return
        payload = await asyncio.to_thread(self.orders.submission_payload, posting_number, store_id)
        await asyncio.to_thread(
            self.orders.update_submission,
            posting_number,
            submission_status="submitting",
        )
        try:
            created = await client.create_order(payload)
        except Exception as error:
            if getattr(error, "outcome_unknown", False):
                await asyncio.to_thread(
                    self.orders.update_submission,
                    posting_number,
                    submission_status="unknown",
                    submission_error=str(error),
                )
                return
            raise
        order_id = str(created.get("order_id") or created.get("id") or "")
        await asyncio.to_thread(
            self.orders.update_submission,
            posting_number,
            submission_status="submitted",
            hicustom_order_id=order_id,
            submission_error="",
        )
