from __future__ import annotations

from erp.core.ai_concurrency import post as ai_post

import hashlib
import json
import math
import os
import re
import time
import uuid
import base64
from io import BytesIO
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from PIL import Image, ImageOps
from sqlalchemy import case, func, or_, select, update
from sqlalchemy.orm import Session, selectinload

from erp.core.config import get_settings
from erp.core.errors import AppError, ConflictError, NotFoundError
from erp.core.jsonutil import dumps, loads
from erp.core.security import vault
from erp.models import Asset, CatalogProduct, CatalogVariant, IntegrationConnection, Job, JobArtifact, JobEvent, Store, WbPublicationItem
from erp.services.jobs import JobContext, enqueue_job, get_job, job_handler


# Optional compliance/name fields that must stay blank for POD products.
# AI otherwise tends to copy the product title into them.
WB_OMITTED_CHARACTERISTIC_IDS={15000785,62945173}


def _roots() -> tuple[Path, Path]:
    root = get_settings().data_dir / "pod" / "wb"
    templates, exports = root / "templates", root / "exports"
    templates.mkdir(parents=True, exist_ok=True); exports.mkdir(parents=True, exist_ok=True)
    return templates, exports


def _manifest_path() -> Path: return _roots()[0] / "manifest.json"


def _manifest() -> list[dict[str, Any]]:
    path = _manifest_path()
    try: value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError): value = []
    return [dict(item) for item in value if isinstance(item, dict)] if isinstance(value, list) else []


def _save_manifest(items: list[dict[str, Any]]) -> None:
    path = _manifest_path(); temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8"); os.replace(temporary, path)


def _defined_columns(workbook: Any, sheet_name: str) -> dict[str, int]:
    result: dict[str, int] = {}
    for name in workbook.defined_names.values():
        try: destinations = list(name.destinations)
        except (AttributeError, TypeError, ValueError): continue
        for target_sheet, coordinate in destinations:
            if str(target_sheet).strip("'") != sheet_name: continue
            match = re.search(r"\$?([A-Z]+)\$?3", str(coordinate))
            if match: result[str(name.name)] = column_index_from_string(match.group(1))
    return result


def _column(columns: dict[str, int], prefix: str) -> int | None:
    return next((value for name, value in columns.items() if name.startswith(prefix)), None)


_FULL_COLUMN_REFERENCE = re.compile(r"^(\$?[A-Z]{1,3}):(\$?[A-Z]{1,3})$")


def _normalize_official_template(content: bytes) -> bytes:
    """Make WB's full-column validation ranges readable by openpyxl.

    Current WB exports contain validation targets such as ``S:S``. Excel accepts
    those targets, but openpyxl's MultiCellRange parser requires explicit row
    numbers and refuses to open the whole workbook. Only worksheet XML is
    adjusted; workbook values, styles, names and validation rules are retained.
    """
    source = BytesIO(content); target = BytesIO(); changed = False
    try:
        with ZipFile(source, "r") as incoming, ZipFile(target, "w", ZIP_DEFLATED) as outgoing:
            for info in incoming.infolist():
                value = incoming.read(info.filename)
                if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                    text = value.decode("utf-8")

                    def expand(match: re.Match[str]) -> str:
                        nonlocal changed
                        references = match.group(1).split()
                        normalized = []
                        for reference in references:
                            full_column = _FULL_COLUMN_REFERENCE.fullmatch(reference)
                            if full_column:
                                left, right = (part.replace("$", "") for part in full_column.groups())
                                reference = f"{left}1:{right}1048576"; changed = True
                            normalized.append(reference)
                        return f'sqref="{" ".join(normalized)}"'

                    value = re.sub(r'sqref="([^"]*)"', expand, text).encode("utf-8")
                outgoing.writestr(info, value)
    except Exception as error:
        raise AppError("WB 模板文件损坏或不是有效的 XLSX", code="wb_template_invalid") from error
    return target.getvalue() if changed else content


def template_schema(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[workbook.sheetnames[0]]; columns = _defined_columns(workbook, sheet.title)
        required = ("static_groupID_", "static_vendorCode_", "static_subject_", "static_photo_", "static_price_")
        missing = [value for value in required if _column(columns, value) is None]
        if missing: raise AppError("WB 模板缺少必要字段，请从 WB 后台下载最新模板", code="wb_template_invalid", details={"missing": missing})
        characteristics = []
        for name, column in columns.items():
            match = re.fullmatch(r"characteristic_(\d+)_\d+", name)
            if match:
                labels = [str(sheet.cell(row, column).value or "").strip() for row in (1, 2, 3, 4)]
                # Row 3 is the actual WB field header. Row 1 is often only a
                # broad section title (for example, packaging), which is not
                # useful to the AI when selecting a characteristic value.
                label = labels[2] or next((x for x in labels if x and not x.startswith("characteristic_")), "")
                characteristics.append({"id": int(match.group(1)), "name": label})
        sheet_title = sheet.title
    finally: workbook.close()
    return {"sheet": sheet_title, "columns": len(columns), "characteristics": characteristics}


def save_template(content: bytes, filename: str, subject_name: str, subject_id: int) -> dict[str, Any]:
    if not content or len(content) > 30 * 1024 * 1024: raise AppError("WB 模板为空或超过 30MB", code="wb_template_size")
    if not subject_name.strip() or int(subject_id) <= 0: raise AppError("请填写 WB 类目名称和类目 ID", code="wb_template_subject_missing")
    content = _normalize_official_template(content)
    identity = uuid.uuid4().hex; path = _roots()[0] / f"{identity}.xlsx"; path.write_bytes(content)
    try: schema = template_schema(path)
    except Exception: path.unlink(missing_ok=True); raise
    item = {"id": identity, "filename": Path(filename).name, "path": str(path), "subject_name": subject_name.strip(), "subject_id": int(subject_id), "sha256": hashlib.sha256(content).hexdigest(), "schema": schema, "created_at": datetime.now().astimezone().isoformat()}
    items = _manifest(); items.insert(0, item); _save_manifest(items); return item


def list_templates() -> list[dict[str, Any]]:
    return [{key: value for key, value in item.items() if key != "path"} for item in _manifest() if Path(str(item.get("path") or "")).is_file()]


def _template(template_id: str) -> dict[str, Any]:
    item = next((row for row in _manifest() if row.get("id") == template_id), None)
    if not item or not Path(str(item.get("path") or "")).is_file(): raise NotFoundError("WB 类目模板不存在")
    return item


def _snapshot_connection(row: IntegrationConnection | None) -> dict[str, Any]:
    if not row: return {}
    return {"base_url": row.base_url, "config": loads(row.config_json, {}), "secrets": vault.decrypt(row.secrets_encrypted)}


_WB_BLOCKING_STATUSES={
    "queued","running","exported_pending_confirmation","submitted",
    "pending_stock","stock_processing","pending_price","price_processing",
    "completed","paused","stock_failed","price_failed","sync_conflict","sync_unknown",
}


def _effective_ledger_status(session: Session, item: WbPublicationItem) -> str:
    if item.status == "withdrawn":return "cancelled"
    if item.job_id and item.status in {"queued","running"}:
        job=session.get(Job,item.job_id)
        if job and job.status in {"failed","cancelled"}:return job.status
        if job and job.status in {"queued","running","paused"}:return "running" if job.status=="running" else job.status
    return item.status


def publication_availability(session: Session, *, store_id: str, product_ids: list[str]) -> dict[str, Any]:
    ids=list(dict.fromkeys(product_ids))
    rows=session.scalars(select(WbPublicationItem).where(WbPublicationItem.store_id==store_id,WbPublicationItem.product_id.in_(ids))).all() if ids else []
    products:dict[str,Any]={identity:{"status":"unlisted","blocked":False,"items":[]} for identity in ids}
    rank={
        "completed":10,"price_processing":9,"pending_price":8,
        "stock_processing":7,"pending_stock":6,"submitted":5,
        "exported_pending_confirmation":4,"running":3,"queued":2,
        "price_failed":1,"stock_failed":1,"failed":1,"cancelled":0,"unlisted":-1,
    }
    for row in rows:
        status=_effective_ledger_status(session,row);value={"variant_id":row.variant_id,"offer_id":row.offer_id,"status":status,"job_id":row.job_id,"external_product_id":row.external_product_id,"last_error":row.last_error}
        target=products.setdefault(row.product_id,{"status":"unlisted","blocked":False,"items":[]});target["items"].append(value)
        if rank.get(status,0)>rank.get(target["status"],-1):target["status"]=status
        target["blocked"]=target["blocked"] or status in _WB_BLOCKING_STATUSES or bool(row.external_product_id)
    return {"store_id":store_id,"products":products}


def enqueue_history_reconcile(session:Session,store_id:str,*,stock_workflow:bool=False)->Job:
    from erp.core.task_identity import task_identity
    actor=task_identity.get() or {}
    store=session.get(Store,store_id)
    if not store or store.platform.casefold() not in {"wb","wildberries"} or not store.enabled:raise ConflictError("请选择已启用的 WB 店铺")
    session.execute(update(Store).where(Store.id==store.id).values(enabled=Store.enabled))
    for pending in session.scalars(select(Job).where(Job.type=='pod_wb_history_reconcile',
            Job.status.in_({'queued','running','paused'}),Job.payload_json.contains(store_id)).order_by(Job.created_at.desc()).limit(20)):
        previous=loads(pending.payload_json,{})
        if previous.get('store_id')==store_id:
            if not previous.get('shop_sync') and bool(previous.get('stock_workflow'))==stock_workflow and pending.created_by==actor.get('id','') and pending.target_worker_id==actor.get('worker_id',''):
                session.commit();return pending
            raise ConflictError('该 WB 店铺已有同步任务，请等待原电脑完成，避免重复读取平台')
    secrets=vault.decrypt(store.secrets_encrypted) if store.secrets_encrypted else {}
    if store.connection and store.connection.secrets_encrypted:secrets={**vault.decrypt(store.connection.secrets_encrypted),**secrets}
    if not str(secrets.get("token") or secrets.get("api_key") or secrets.get("wildberries_token") or "").strip():raise AppError("当前 WB 店铺未配置 Content API Token",code="wb_api_token_missing")
    prefix=str(store.offer_prefix or "").strip().upper()
    if not prefix:raise AppError("请先为 WB 店铺设置 SKU 前缀",code="wb_offer_prefix_missing")
    snapshot=vault.encrypt({"store":{"id":store.id,"name":store.name,"offer_prefix":prefix,"secrets":secrets}})
    return enqueue_job(session,job_type="pod_wb_history_reconcile",payload={"store_id":store.id,"store_name":store.name,"offer_prefix":prefix,"stock_workflow":stock_workflow},config_snapshot=snapshot,idempotency_key=f"pod-wb-history:{store.id}:{uuid.uuid4().hex}",max_attempts=1,steps=[("download","读取 WB 店铺商品卡"),("match","匹配 ERP POD SKU"),("ledger","写入历史上架台账")])


def create_export(session: Session, *, product_ids: list[str], store_id: str, template_id: str, brand: str, requirement: str, price_cny: float, pricing: dict[str, Any] | None = None, image_rules: list[dict[str, Any]] | None = None, variant_selection: str, variant_ids: list[str] | None = None, mode: str = "excel", generation_provider: str = "model", codex_request_id: str = "") -> Job:
    mode=str(mode or "excel").casefold()
    if mode not in {"excel","api"}:raise AppError("不支持的 WB 发布方式",code="wb_publish_mode_invalid")
    if generation_provider not in {"model", "codex"}:raise ConflictError("不支持的文案生成方式")
    if generation_provider == "codex":
        from erp.core.task_identity import task_identity
        from erp.models import WorkerNode
        actor=task_identity.get() or {};device=session.get(WorkerNode,actor.get("worker_id", ""))
        if not device or device.owner_user_id!=actor.get("id") or "codex_wb_v1" not in loads(device.capabilities_json,[]):
            raise ConflictError("请在原执行电脑安装并启动 0.6.72 或更新的客户端")
    ids = list(dict.fromkeys(product_ids))
    products = session.scalars(select(CatalogProduct).where(CatalogProduct.id.in_(ids), CatalogProduct.source_type == "pod")).all()
    if len(products) != len(ids): raise NotFoundError("部分 POD 产品不存在")
    store = session.get(Store, store_id)
    if not store or store.platform.casefold() not in {"wb", "wildberries"} or not store.enabled: raise ConflictError("请选择已启用的 WB 店铺")
    # Serialize duplicate-SKU checks with other submissions for this store.
    session.execute(update(Store).where(Store.id == store_id).values(enabled=Store.enabled))
    prefix=str(store.offer_prefix or "").strip().upper()
    if not prefix:raise AppError("请先为 WB 店铺设置 SKU 前缀",code="wb_offer_prefix_missing")
    if not re.fullmatch(r"[A-Za-z0-9_-]+",prefix):raise AppError("WB 店铺 SKU 前缀只能包含英文字母、数字、下划线和短横线",code="wb_offer_prefix_invalid")
    duplicate_prefix=session.scalar(select(Store).where(Store.id!=store.id,Store.enabled.is_(True),Store.platform.in_(["wb","wildberries"]),func.lower(Store.offer_prefix)==prefix.casefold()))
    if duplicate_prefix:raise ConflictError(f"SKU 前缀 {prefix} 已被店铺“{duplicate_prefix.name}”使用")
    template = _template(template_id)
    ai = session.scalar(select(IntegrationConnection).where(IntegrationConnection.provider == "ai_chat", IntegrationConnection.enabled.is_(True)))
    if not ai and generation_provider == "model": raise ConflictError("请先配置 micuapi 文案接口")
    store_secrets=vault.decrypt(store.secrets_encrypted) if store.secrets_encrypted else {}
    if store.connection and store.connection.secrets_encrypted:
        store_secrets={**vault.decrypt(store.connection.secrets_encrypted),**store_secrets}
    if mode=="api" and not str(store_secrets.get("token") or store_secrets.get("api_key") or store_secrets.get("wildberries_token") or "").strip():
        raise AppError("当前 WB 店铺未配置 Content API Token",code="wb_api_token_missing")
    snapshot = vault.encrypt({"ai": _snapshot_connection(ai) if ai and generation_provider=="model" else {}, "store": {"id": store.id, "name": store.name, "offer_prefix": store.offer_prefix,"secrets":store_secrets}})
    selected_variant_ids=list(dict.fromkeys(variant_ids or []))
    if variant_selection=="selected" and not selected_variant_ids: raise AppError("请至少勾选一个颜色/尺码 SKU",code="wb_variant_selection_missing")
    selected_variants=[variant for product in products for variant in _selected_variants(product,variant_selection,selected_variant_ids)]
    if not selected_variants:raise AppError("没有可发布的 SKU",code="wb_variant_selection_missing")
    offers={variant.id:f"{prefix.rstrip('-')}-{variant.sku}" for variant in selected_variants}
    too_long=[offer for offer in offers.values() if len(offer)>40]
    if too_long:raise AppError(f"WB 卖家 SKU 超过 40 字符：{too_long[0]}",code="wb_vendor_code_long")
    existing={row.variant_id:row for row in session.scalars(select(WbPublicationItem).where(WbPublicationItem.store_id==store.id,WbPublicationItem.variant_id.in_(list(offers)))).all()}
    blocked=[]
    for variant in selected_variants:
        row=existing.get(variant.id);status=_effective_ledger_status(session,row) if row else "unlisted"
        if status in _WB_BLOCKING_STATUSES or row and row.external_product_id:blocked.append({"sku":variant.sku,"offer_id":row.offer_id,"status":status})
    if blocked:raise AppError(f"当前店铺已有 {len(blocked)} 个 SKU 正在处理或已上架，请取消重复项",code="wb_duplicate_publication",details={"items":blocked})
    pricing_value=dict(pricing or {});pricing_value.setdefault("mode","manual");pricing_value.setdefault("manual_price_rub",float(price_cny))
    normalized_rules=[]
    for position,rule in enumerate(image_rules or []):
        key=Path(str(rule.get("key") or "")).name.strip();action=str(rule.get("action") or "additional")
        if key and action in {"main","additional","exclude"}:normalized_rules.append({"key":key,"action":action,"position":position})
    if normalized_rules and sum(rule["action"]=="main" for rule in normalized_rules)!=1:raise AppError("图片排序必须且只能指定一张主图",code="wb_image_main_invalid")
    payload = {"product_ids": ids, "store_id": store.id, "store_name": store.name, "template_id": template_id, "template_sha256": template["sha256"], "subject_name": template["subject_name"], "subject_id": template["subject_id"], "brand": brand.strip(), "requirement": requirement.strip(), "price_cny": float(price_cny), "pricing":pricing_value,"image_rules":normalized_rules, "variant_selection": variant_selection, "variant_ids": selected_variant_ids,"mode":mode}
    job_type="pod_wb_excel_export" if mode=="excel" else "pod_wb_api_publish"
    final_steps=[("excel","生成 WB 批量上传表格")] if mode=="excel" else [("barcodes","申请并保存 WB 条码"),("submit","创建 WB 商品卡"),("verify","核对商品编号并上传图片")]
    # Publish the queue entry and its SKU ledger in the caller's transaction.
    # Otherwise a personal worker can claim a job before its ledger exists.
    job=enqueue_job(session, job_type=job_type, payload=payload, config_snapshot=snapshot, idempotency_key=f"pod-wb-{mode}:{uuid.uuid4().hex}", max_attempts=1, steps=[("validate", "校验产品与 WB 模板"), ("copy", "Codex 生成本次俄文文案与属性" if generation_provider=="codex" else "AI 生成本次俄文文案与属性"),*final_steps], commit=False)
    if generation_provider=="codex":
        payload.update(generation_provider="codex",batch_id=job.id,codex_publish_approved=False,codex_products=[{"id":p.id,"sku":p.sku,"name":p.name} for p in products])
        payload['codex_request_id']=codex_request_id or job.id
        job.payload_json=dumps(payload);job.status="paused";job.pause_requested=True
        job.error_code="waiting_codex";job.error_message="等待 Codex 完成本次 WB 文案与属性"
    product_by_variant={variant.id:variant.product_id for variant in selected_variants}
    for variant_id,offer_id in offers.items():
        row=existing.get(variant_id)
        if row:
            row.job_id=job.id;row.template_id=template_id;row.offer_id=offer_id;row.mode=mode;row.status="queued";row.external_product_id="";row.last_error="";row.confirmed_at=None
        else:session.add(WbPublicationItem(store_id=store.id,product_id=product_by_variant[variant_id],variant_id=variant_id,job_id=job.id,template_id=template_id,offer_id=offer_id,mode=mode,status="queued"))
    return job


def confirm_export(session: Session, job_id: str, *, uploaded: bool) -> dict[str, Any]:
    job=get_job(session,job_id)
    if job.type!="pod_wb_excel_export":raise ConflictError("该任务不是 WB Excel 发布任务")
    rows=session.scalars(select(WbPublicationItem).where(WbPublicationItem.job_id==job.id)).all()
    if not rows:raise NotFoundError("该任务没有 WB 上架台账")
    export_rows=[row for row in rows if row.status!="failed"]
    if not export_rows or any(row.status != "exported_pending_confirmation" for row in export_rows):
        raise ConflictError("该批次已确认或状态已变化，请刷新后核对")
    status="submitted" if uploaded else "cancelled"
    for row in export_rows:row.status=status;row.confirmed_at=None
    session.commit()
    return {"job_id":job.id,"status":status,"items":len(export_rows),"failed_items":len(rows)-len(export_rows)}


def reconcile_excel_export(session: Session, job_id: str) -> dict[str, Any]:
    job=get_job(session,job_id)
    if job.type!="pod_wb_excel_export":raise ConflictError("该任务不是 WB Excel 发布任务")
    rows=[row for row in session.scalars(select(WbPublicationItem).where(WbPublicationItem.job_id==job.id)).all() if row.status!="failed"]
    if not rows or any(row.status not in {"submitted","pending_stock"} and not (row.status=="completed" and not row.external_product_id) for row in rows):
        raise ConflictError("请先确认已上传 Excel，再核对 WB 商品卡")
    store,snapshot=_wb_store_snapshot(session,rows[0].store_id)
    client=_WbContentClient(vault.decrypt(snapshot))
    try:found=client.find_cards_exact({row.offer_id for row in rows})
    finally:client.close()
    matched=0
    for row in rows:
        card=found.get(row.offer_id)
        nm_id=int(card.get("nmID") or 0) if card else 0
        if nm_id<=0:continue
        if row.external_product_id and row.external_product_id!=str(nm_id):
            raise ConflictError(f"{row.offer_id} 的 WB 商品编号与上次核对结果冲突")
        row.external_product_id=str(nm_id);row.status="pending_stock";row.last_error="";row.confirmed_at=datetime.now().astimezone();matched+=1
    session.commit()
    return {"job_id":job.id,"status":"pending_stock" if matched==len(rows) else "submitted","matched":matched,"total":len(rows),"pending":len(rows)-matched}


def enqueue_excel_reconcile(session: Session, job_id: str) -> Job:
    original = get_job(session, job_id)
    if original.type != 'pod_wb_excel_export' or original.status != 'completed':
        raise ConflictError('只有已完成的 WB Excel 任务可以核对')
    from erp.core.task_identity import task_identity
    actor = task_identity.get() or {}
    active = session.scalars(select(Job).where(
        Job.type == 'pod_wb_excel_reconcile', Job.status.in_({'queued', 'running'})))
    for existing in active:
        if (loads(existing.payload_json, {}).get('export_job_id') == job_id
                and (not get_settings().multiuser_enabled or (
                    existing.created_by == actor.get('id')
                    and existing.target_worker_id == actor.get('worker_id')))):
            return existing
    rows = list(session.scalars(select(WbPublicationItem).where(
        WbPublicationItem.job_id == job_id,
        WbPublicationItem.status.in_({'submitted', 'pending_stock'}))))
    unresolved = [row for row in rows if row.status == 'submitted']
    if unresolved:
        rows = unresolved
    if not rows or any(row.status not in {'submitted', 'pending_stock'} for row in rows):
        raise ConflictError('请先确认 Excel 已上传，并刷新商品状态')
    if len({row.store_id for row in rows}) != 1:
        raise ConflictError('该任务包含不同 WB 店铺')
    _, snapshot = _wb_store_snapshot(session, rows[0].store_id)
    job = enqueue_job(session, job_type='pod_wb_excel_reconcile',
                      payload={'export_job_id': job_id, 'store_id': rows[0].store_id,
                               'offer_ids': sorted({row.offer_id for row in rows})},
                      config_snapshot=snapshot,
                      idempotency_key=f'pod-wb-excel-reconcile:{job_id}:{uuid.uuid4().hex}',
                      max_attempts=1, steps=[('lookup', '在本人电脑核对 WB 商品卡')])
    return job


def list_exports(session: Session, limit: int = 100, *, include_codex_siblings: bool = False) -> list[dict[str, Any]]:
    jobs = session.scalars(select(Job).where(Job.type.in_(["pod_wb_excel_export","pod_wb_api_publish","pod_wb_history_reconcile"])).order_by(Job.created_at.desc()).limit(limit)).all()
    if include_codex_siblings:
        scopes={(loads(j.payload_json,{}).get('codex_request_id'),j.created_by,j.target_worker_id)
                for j in jobs if loads(j.payload_json,{}).get('generation_provider')=='codex' and loads(j.payload_json,{}).get('codex_request_id')}
        if scopes:
            from sqlalchemy import or_
            siblings=session.scalars(select(Job).where(Job.type.in_(['pod_wb_excel_export','pod_wb_api_publish']),
                or_(*(Job.payload_json.contains(scope[0]) for scope in scopes))))
            all_jobs={j.id:j for j in jobs}
            for job in siblings:
                if (loads(job.payload_json,{}).get('codex_request_id'),job.created_by,job.target_worker_id) in scopes:
                    all_jobs[job.id]=job
            jobs=sorted(all_jobs.values(),key=lambda j:j.created_at,reverse=True)
    result=[]
    for job in jobs:
        value={"id":job.id,"created_by":job.created_by,"target_worker_id":job.target_worker_id,"status":job.status,"progress":job.progress,"error_message":job.error_message,"created_at":job.created_at.isoformat(),"payload":loads(job.payload_json,{}),"mode":"history" if job.type=="pod_wb_history_reconcile" else "api" if job.type=="pod_wb_api_publish" else "excel"}
        from erp.core.task_identity import task_identity
        actor=task_identity.get() or {}
        if get_settings().multiuser_enabled and job.created_by!=actor.get("id"):
            value["payload"].pop("codex_outputs",None)
        ledger=session.scalars(select(WbPublicationItem).where(WbPublicationItem.job_id==job.id)).all()
        value["download_ready"]=any(item.role=="wb_excel" for item in job.artifacts)
        value["history_report_asset_id"]=next((item.asset_id for item in job.artifacts if item.role=="wb_reconcile_report"),None)
        ledger_statuses={item.status for item in ledger}
        legacy_unverified=job.type=="pod_wb_excel_export" and any(item.status=="completed" and not item.external_product_id for item in ledger)
        value["publication_status"]=("submitted" if "submitted" in ledger_statuses or legacy_unverified else "pending_stock" if "pending_stock" in ledger_statuses else "exported_pending_confirmation" if "exported_pending_confirmation" in ledger_statuses else ledger[0].status if ledger else "")
        value["copy_failed_items"]=sum(item.status=="failed" for item in ledger)
        value["pending_stock_items"]=sum(item.status=="pending_stock" for item in ledger)
        external_ids=list(dict.fromkeys(item.external_product_id for item in ledger if item.external_product_id))
        value["external_product_count"]=len(external_ids);value["external_product_ids"]=external_ids[:5]
        value["publication_items"]=len(ledger);result.append(value)
    return result


def export_file(session: Session, job_id: str) -> Path:
    job=get_job(session,job_id)
    artifact=next((item for item in job.artifacts if item.role=="wb_excel"),None)
    asset=session.get(Asset,artifact.asset_id) if artifact else None
    if not asset or not asset.local_path: raise ConflictError("WB 表格尚未生成完成")
    path=Path(asset.local_path)
    if not path.is_file(): raise NotFoundError("WB 表格文件不存在")
    return path


def _json_object(text: str) -> dict[str, Any]:
    cleaned=text.strip().removeprefix("```json").removesuffix("```").strip()
    try: value=json.loads(cleaned)
    except ValueError:
        match=re.search(r"\{.*\}",cleaned,re.S)
        if not match: raise AppError("AI 未返回有效 JSON",code="wb_ai_invalid")
        value=json.loads(match.group(0))
    if not isinstance(value,dict): raise AppError("AI 文案格式无效",code="wb_ai_invalid")
    return value


def _ai_image_data_url(url: str, maximum: int = 768) -> str:
    response=httpx.get(url,timeout=45,follow_redirects=True)
    response.raise_for_status()
    if len(response.content)>30*1024*1024:raise ValueError("image too large")
    with Image.open(BytesIO(response.content)) as opened:
        image=ImageOps.exif_transpose(opened).convert("RGB");image.thumbnail((maximum,maximum),Image.Resampling.LANCZOS)
        buffer=BytesIO();image.save(buffer,"JPEG",quality=82,optimize=True)
    return "data:image/jpeg;base64,"+base64.b64encode(buffer.getvalue()).decode("ascii")


def _generate_copy(config: dict[str, Any], product: CatalogProduct, template: dict[str, Any], payload: dict[str, Any]) -> dict[str, Any]:
    if payload.get("generation_provider")=="codex":
        from erp.services.wb_codex import validate_copy
        value=(payload.get("_codex_outputs_by_product") or {}).get(product.id)
        if value is None:raise AppError("等待原电脑 Codex 输出，不会调用外部模型",code="codex_output_missing")
        return validate_copy(value,template,payload.get("_copy_references") or [])
    from erp.core.endpoints import ai_endpoint
    ai=config.get("ai") or {}; token=str((ai.get("secrets") or {}).get("api_key") or "").strip()
    endpoint=str(ai.get("base_url") or "").rstrip("/")
    if not token or not endpoint: raise AppError("AI 文案接口配置不完整",code="wb_ai_not_configured")
    endpoint = ai_endpoint(endpoint)
    model=str((ai.get("config") or {}).get("default_model") or "gpt-5.6-terra")
    facts={"name":product.name,"supplier":product.supplier,"category":product.category,"production":loads(product.production_json,{}),"supplier_facts":loads(product.source_snapshot_json,{}),"variants":[{"sku":v.sku,"color":v.color,"size":v.size,"specification":v.specification} for v in product.variants if v.active]}
    schema=template.get("schema") or {}; prompt=("为 Wildberries 商品生成俄文资料。只返回 JSON：title(最多60字符)、description(最多5000字符)、characteristics数组，每项含id和value。不得猜测尺寸、重量、成本、SKU。只能使用给定事实。"
        f"\n类目：{payload['subject_name']}\n品牌：{payload['brand']}\n本次要求：{payload['requirement']}\n可填写属性：{json.dumps(schema.get('characteristics') or [],ensure_ascii=False)}\n产品事实：{json.dumps(facts,ensure_ascii=False)}")
    from erp.core.ai_usage import scope
    job_id=scope.get().get("job_id","")
    prompt+=f"\n本次上架：{job_id}；店铺：{payload.get('store_name') or payload.get('store_id') or ''}。必须为本次上架重新构思俄文标题和描述，与已给出的历史文案使用不同句式和段落结构，不得改动商品事实。不要把任务编号或店铺名称写入文案。"
    references=payload.get("_copy_references") or []
    if references:
        prompt+="\n以下历史版本仅用于避免重复，不是商品事实："+json.dumps(references[:5],ensure_ascii=False)
    images=[link.asset.oss_url for link in product.asset_links if link.is_current and str(link.asset.oss_url).startswith("https://")][:3]
    prepared=[]
    for url in images:
        try:prepared.append(_ai_image_data_url(str(url)))
        except (httpx.HTTPError,OSError,ValueError):continue
    content=[{"type":"text","text":prompt}]+[{"type":"image_url","image_url":{"url":url,"detail":"low"}} for url in prepared]
    from erp.core.ai_reuse import fingerprint, read_json, write_json, result_cache
    # New listing jobs ALWAYS generate fresh copy. Only the same job can
    # recover a paid success that preceded writing its batch checkpoint.
    if not job_id:
        return _distinct_copy(endpoint,token,model,content,references)
    key=fingerprint({"version":2,"job_id":job_id,"store":payload.get("store_id"),"product":product.id,"endpoint":endpoint,"account":fingerprint(token),"model":model,"content":content})
    with result_cache("wb-copy",key) as cached_path:
        cached=read_json(cached_path)
        if cached and str(cached.get("title") or "").strip() and str(cached.get("description") or "").strip():
            return cached
        result=_distinct_copy(endpoint,token,model,content,references)
        write_json(cached_path,result)
        return result


def _distinct_copy(endpoint,token,model,content,references):
    import re
    from difflib import SequenceMatcher
    def normalize(value):
        return re.sub(r"[\W_]+","",str(value or "").casefold())
    messages=list(content)
    for attempt in range(2):
        result=_request_copy(endpoint,token,model,messages)
        duplicates=[]
        for old in references:
            for key,threshold in (("title",.82),("description",.70)):
                left,right=normalize(result.get(key)),normalize(old.get(key))
                if left and right and SequenceMatcher(None,left,right).ratio()>=threshold:
                    duplicates.append(key)
        if not duplicates:
            return result
        messages=[*content,{"type":"text","text":"上一版与历史版本过于相似。请重新组织标题句式和描述段落，不能只替换同义词。"}]
    raise AppError("本次文案与历史店铺版本过于相似，已停止使用，请调整创作要求",code="wb_copy_not_distinct")


def _previous_copy_references(session,product_id,job_id):
    references=[]
    rows=session.execute(select(Job.id,Job.payload_json).where(Job.type.in_(["pod_wb_api_publish","pod_wb_excel_export"]),Job.id!=job_id,Job.payload_json.contains(product_id)).order_by(Job.created_at.desc()).limit(200))
    for identity,raw in rows:
        if product_id not in loads(raw,{}).get("product_ids",[]):continue
        try:value=json.loads((_roots()[1]/identity/"copy.json").read_text(encoding="utf-8")).get(product_id)
        except (OSError,ValueError,AttributeError):continue
        if isinstance(value,dict) and value.get("title") and value.get("description"):
            references.append({"title":value["title"],"description":value["description"]})
    return references


def _previous_copy_references_bulk(session, product_ids, job_id):
    """Read prior copy for a publication in batches, keeping per-product recency."""
    from sqlalchemy import or_

    identities=list(dict.fromkeys(product_ids))
    result={identity:[] for identity in identities}
    counts={identity:0 for identity in identities}
    requested=set(identities)
    copy_cache={}
    for start in range(0,len(identities),40):
        chunk=set(identities[start:start+40])
        rows=session.execute(select(Job.id,Job.payload_json).where(
            Job.type.in_(['pod_wb_api_publish','pod_wb_excel_export']),
            Job.id!=job_id,
            or_(*(Job.payload_json.contains(identity) for identity in chunk)),
        ).order_by(Job.created_at.desc()))
        for prior_id,raw in rows:
            matched=chunk.intersection(loads(raw,{}).get('product_ids',[]))
            if not matched:continue
            if prior_id not in copy_cache:
                try:copy_cache[prior_id]=json.loads((_roots()[1]/prior_id/'copy.json').read_text(encoding='utf-8'))
                except (OSError,ValueError,AttributeError):copy_cache[prior_id]={}
            copies=copy_cache[prior_id]
            if not isinstance(copies,dict):continue
            for identity in matched:
                if identity not in requested or counts[identity]>=200:continue
                counts[identity]+=1
                value=copies.get(identity)
                if isinstance(value,dict) and value.get('title') and value.get('description'):
                    result[identity].append({'title':value['title'],'description':value['description']})
    return result


def _request_copy(endpoint,token,model,content):
    try: response=ai_post("text", httpx.post,endpoint,headers={"Authorization":f"Bearer {token}"},json={"model":model,"messages":[{"role":"user","content":content}],"response_format":{"type":"json_object"}},timeout=180)
    except httpx.HTTPError as error: raise AppError(f"AI 文案请求失败：{error}",code="wb_ai_network") from error
    if not response.is_success: raise AppError(f"AI 文案接口 HTTP {response.status_code}",code="wb_ai_http")
    choices=response.json().get("choices") or []
    if not choices: raise AppError("AI 未返回文案结果",code="wb_ai_invalid")
    result=_json_object(str(choices[0].get("message",{}).get("content") or ""))
    if not str(result.get("title") or "").strip() or not str(result.get("description") or "").strip(): raise AppError("AI 未生成完整俄文标题和描述",code="wb_ai_incomplete")
    return result


def _wb_characteristics(config: dict[str, Any], subject_id: int) -> list[dict[str, Any]]:
    secrets=(config.get("store") or {}).get("secrets") or {};token=str(secrets.get("token") or secrets.get("api_key") or secrets.get("wildberries_token") or "").strip()
    if not token:return []
    try: response=httpx.get(f"https://content-api.wildberries.ru/content/v2/object/charcs/{subject_id}",params={"locale":"ru"},headers={"Authorization":token},timeout=45)
    except httpx.HTTPError as error: raise AppError(f"读取 WB 实时类目要求失败：{error}",code="wb_schema_network") from error
    if not response.is_success: raise AppError(f"读取 WB 实时类目要求失败：HTTP {response.status_code}",code="wb_schema_http")
    value=response.json();rows=value.get("data") if isinstance(value,dict) else value
    if not isinstance(rows,list): raise AppError("WB 返回的类目要求格式无效",code="wb_schema_invalid")
    return [dict(item) for item in rows if isinstance(item,dict)]


def _validated_characteristics(raw_values: list[dict[str, Any]], schema: list[dict[str, Any]]) -> list[dict[str, Any]]:
    raw={int(item.get("id") or 0):item.get("value") for item in raw_values if isinstance(item,dict) and int(item.get("id") or 0)>0};result=[];missing=[]
    for item in schema:
        identity=int(item.get("charcID") or 0);kind=int(item.get("charcType") or 0)
        if identity<=0 or kind==0 or identity in WB_OMITTED_CHARACTERISTIC_IDS:continue
        value=raw.get(identity)
        if kind==4:
            try: normalized=float(value) if value not in (None,"") else None
            except (TypeError,ValueError): normalized=None
            if normalized is not None and normalized.is_integer():normalized=int(normalized)
        else:
            normalized=[str(entry).strip() for entry in (value if isinstance(value,list) else re.split(r"[,;，；\n]+",str(value or ""))) if str(entry).strip()]
            maximum=int(item.get("maxCount") or 0)
            if maximum>0:normalized=normalized[:maximum]
        if bool(item.get("required")) and normalized in (None,[]):missing.append(str(item.get("name") or identity))
        if normalized not in (None,[]):result.append({"id":identity,"value":normalized})
    if missing:raise AppError("缺少 WB 实时类目必填属性："+"、".join(missing),code="wb_characteristics_missing")
    return result


def _template_value_limits(template_path: Path) -> dict[int, int]:
    """Read per-cell value counts printed in the seller's Excel template."""
    workbook = load_workbook(template_path, read_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        columns = _defined_columns(workbook, sheet.title)
        limits = {}
        for name, column in columns.items():
            match = re.fullmatch(r"characteristic_(\d+)_\d+", name)
            if not match:
                continue
            hint = str(sheet.cell(4, column).value or "")
            count = re.search(r"最大数量\s*[:：]\s*(\d+)", hint)
            if count and int(count.group(1)) > 0:
                limits[int(match.group(1))] = int(count.group(1))
        return limits
    finally:
        workbook.close()


def _cm(value: Any) -> int: return max(1, math.ceil(float(value or 0)/10))


def _variant_color_id(variant: Any) -> int:
    try: return int(loads(variant.properties_json, {}).get("color_id") or 0)
    except (TypeError, ValueError): return 0


def _apply_image_rules(candidates:list[dict[str,str]],image_rules:list[dict[str,Any]]|None)->list[str]:
    rules={str(item.get("key") or ""):{**item,"index":index} for index,item in enumerate(image_rules or [])}
    if not rules:return []
    matched=[(rules[item["key"]],item) for item in candidates if item["key"] in rules and rules[item["key"]].get("action")!="exclude"]
    # The numbered order shown in the UI is authoritative.  WB uses the first
    # URL as the actual main image, so silently moving the row merely labelled
    # ``main`` to the front changes the user's chosen gallery order.
    matched.sort(key=lambda pair:int(pair[0].get("position",pair[0]["index"])))
    return list(dict.fromkeys(item["url"] for _,item in matched))[:30]


def _listing_images(product: CatalogProduct, variant: Any, image_rules: list[dict[str, Any]] | None = None) -> list[str]:
    color_id=_variant_color_id(variant)
    preferred=[];generic=[]
    for link in sorted(product.asset_links,key=lambda item:(item.position,item.id)):
        if not link.is_current or link.role not in {"main","mockup","additional","sku"} or link.asset.kind not in {"image","rendering"}: continue
        url=str(link.asset.oss_url or "")
        if not url.startswith("https://"): continue
        metadata=loads(link.metadata_json,{})
        try: link_color=int(metadata.get("color_id") or loads(link.asset.metadata_json,{}).get("color_id") or 0)
        except (TypeError,ValueError): link_color=0
        if link.variant_id==variant.id or (color_id and link_color==color_id): preferred.append(url)
        elif not link.variant_id and not link_color: generic.append(url)
    candidates=[]
    for link in sorted(product.asset_links,key=lambda item:(item.position,item.id)):
        if not link.is_current or link.role not in {"main","mockup","additional","sku"} or link.asset.kind not in {"image","rendering"}:continue
        url=str(link.asset.oss_url or "")
        if not url.startswith("https://"):continue
        metadata=loads(link.metadata_json,{})
        try:link_color=int(metadata.get("color_id") or loads(link.asset.metadata_json,{}).get("color_id") or 0)
        except (TypeError,ValueError):link_color=0
        if link.variant_id==variant.id or (color_id and link_color==color_id) or (not link.variant_id and not link_color):
            candidates.append({"key":Path(str(link.asset.name or link.asset.oss_key or "")).name,"url":url})
    if image_rules:return _apply_image_rules(candidates,image_rules)
    return list(dict.fromkeys([*preferred,*generic]))[:30]


def _size_characteristics(subject_id: int, variant: Any) -> list[dict[str, Any]]:
    """Map exact supplier/ERP size text to the three known WB category fields."""
    text=" ".join(value for value in (str(variant.size or ""),str(variant.specification or "")) if value)
    values=[float(value.replace(",",".")) for value in re.findall(r"\d+(?:[.,]\d+)?",text)]
    if len(values)<2: return []
    width,height=values[:2]
    display=f"{width:g}x{height:g} см"
    size_field={321:244146,1479:234706,231:15003353}.get(int(subject_id))
    result=[{"id":90673,"value":width},{"id":90630,"value":height}]
    if size_field: result.insert(0,{"id":size_field,"value":display})
    return result


def _variant_color_characteristic(variant: Any) -> dict[str, Any] | None:
    """Prefer the factual ERP variant colour over an AI-inferred print colour."""
    value=str(getattr(variant,"color","") or "").strip().casefold()
    colors={
        "白色":"белый","白":"белый","white":"белый","белый":"белый",
        "黑色":"черный","黑":"черный","black":"черный","черный":"черный",
        "红色":"красный","红":"красный","red":"красный","красный":"красный",
        "蓝色":"синий","蓝":"синий","blue":"синий","синий":"синий",
        "绿色":"зеленый","绿":"зеленый","green":"зеленый","зелёный":"зеленый","зеленый":"зеленый",
        "黄色":"желтый","黄":"желтый","yellow":"желтый","желтый":"желтый",
        "粉色":"розовый","粉":"розовый","pink":"розовый","розовый":"розовый",
        "紫色":"фиолетовый","紫":"фиолетовый","purple":"фиолетовый","фиолетовый":"фиолетовый",
        "灰色":"серый","灰":"серый","gray":"серый","grey":"серый","серый":"серый",
        "棕色":"коричневый","棕":"коричневый","brown":"коричневый","коричневый":"коричневый",
        "米色":"бежевый","beige":"бежевый","бежевый":"бежевый",
        "橙色":"оранжевый","橙":"оранжевый","orange":"оранжевый","оранжевый":"оранжевый",
        "голубой":"голубой",
    }
    normalized=colors.get(value)
    return {"id":14177449,"value":[normalized]} if normalized else None


def _sku_sort_key(value: Any) -> tuple[Any, ...]:
    """Case-insensitive natural order so SKU-2 is before SKU-10."""
    text = str(value or "").strip().casefold()
    return tuple((0, int(part)) if part.isdigit() else (1, part) for part in re.split(r"(\d+)", text) if part)


def _card_sku_sort_key(card: dict[str, Any]) -> tuple[Any, ...]:
    sizes = card.get("sizes") or []
    variant_sku = str(sizes[0].get("sku") or "") if sizes and isinstance(sizes[0], dict) else ""
    primary = variant_sku or str(card.get("product_sku") or "") or str(card.get("vendorCode") or "")
    return (_sku_sort_key(primary), _sku_sort_key(card.get("vendorCode")), str(card.get("product_id") or ""))


def _selected_variants(product: CatalogProduct, selection: str, variant_ids: list[str] | None) -> list[Any]:
    variants=[variant for variant in product.variants if variant.active]
    if selection=="selected":
        allowed=set(variant_ids or []);variants=[variant for variant in variants if variant.id in allowed]
    elif selection!="all":
        count={"largest_one":1,"largest_two":2,"largest_three":3}.get(selection)
        if count:
            variants=sorted(variants,key=lambda variant:(float(variant.package_length_mm or product.package_length_mm or 0)*float(variant.package_width_mm or product.package_width_mm or 0),float(variant.weight_grams or product.package_weight_g or 0)),reverse=True)[:count]
    return sorted(variants,key=lambda variant:(_sku_sort_key(variant.sku),variant.id))


def _shipping_cost_cny(weight_kg: float) -> float:
    if not math.isfinite(weight_kg) or weight_kg <= 0: raise AppError("产品重量必须大于 0 公斤",code="wb_weight_invalid")
    billed=math.ceil(weight_kg*10-1e-9)/10
    if billed<=0.3: return round(2+58*billed,2)
    if billed<=20: return round(8+43*billed,2)
    raise AppError("产品重量超过 20 公斤，无法按当前运费规则计算售价",code="wb_weight_invalid")


def _variant_price(product: CatalogProduct, variant: Any, pricing: dict[str, Any]) -> dict[str, Any]:
    mode=str(pricing.get("mode") or "manual")
    if mode=="manual":
        value=float(pricing.get("manual_price_rub") or pricing.get("price_rub") or 0)
        if value<=0: return {"error":"请填写 WB 划线价"}
        return {"list_price_rub":value,"mode":mode,"source":"本批统一售价"}
    cost=float(variant.purchase_price_cny if variant.purchase_price_cny is not None else product.purchase_price_cny or 0)
    cost+=float(product.domestic_delivery_cny or 0)
    weight=float(variant.weight_grams or product.package_weight_g or 0)/1000
    if cost<=0: return {"error":"缺少准确拿货价"}
    if mode=="multiplier":
        multiplier=float(pricing.get("multiplier") or 0)
        if multiplier<=0:return {"error":"售价倍率必须大于 0"}
        sale_cny=cost*multiplier
        discount=float(pricing.get("discount_percent") or 0)/100
        price=math.ceil(sale_cny/max(0.01,1-discount))
        return {"list_price_rub":price,"sale_price_cny":round(price*(1-discount),2),"purchase_cost_cny":round(cost,2),"mode":mode,"source":"成本倍率（人民币数值原样提交）"}
    try: shipping=_shipping_cost_cny(weight)
    except AppError as error:return {"error":str(error)}
    discount=float(pricing.get("discount_percent",30))/100;target=float(pricing.get("target_profit_percent",30))/100
    commission=float(pricing.get("commission_percent",20))/100;advertising=float(pricing.get("advertising_percent",10))/100;returns=float(pricing.get("return_percent",10))/100
    label=float(pricing.get("label_fee_cny",2.5));other=float(pricing.get("other_cost_cny",4))
    denominator=1-commission-advertising-returns*advertising
    if denominator<=0 or discount>=1:return {"error":"佣金、广告、退货或折扣参数组合无效"}
    target_sale=(cost+label+shipping+other+returns*(cost+label+shipping)+target*(cost+label+other))/denominator
    price=math.ceil(target_sale/(1-discount))
    return {"list_price_rub":price,"sale_price_cny":round(price*(1-discount),2),"purchase_cost_cny":round(cost,2),"shipping_cny":shipping,"mode":"profit","source":"目标利润反算（人民币数值原样提交）"}


def preview_export(session: Session, *, product_ids: list[str], subject_id: int, variant_selection: str, variant_ids: list[str] | None = None, pricing: dict[str, Any] | None = None) -> dict[str, Any]:
    ids=list(dict.fromkeys(product_ids));products=session.scalars(select(CatalogProduct).where(CatalogProduct.id.in_(ids),CatalogProduct.source_type=="pod")).all()
    if len(products)!=len(ids): raise NotFoundError("部分 POD 产品不存在")
    selected=set(variant_ids or []);rows=[];card_total=0;variant_total=0;split_sizes=int(subject_id) in {321}
    for product in sorted(products,key=lambda item:(item.sku,item.id)):
        variants=_selected_variants(product,variant_selection,list(selected))
        if not variants: raise AppError(f"{product.sku} 没有选中的可用变体",code="wb_variant_missing")
        # 毛毯尺寸属于商品卡级属性，必须按尺寸拆卡；其余 POD 产品
        # 第一版按一个印花产品一张卡，所选规格写入该卡 sizes。
        cards=len(variants) if split_sizes else 1
        card_total+=cards;variant_total+=len(variants)
        rows.append({"product_id":product.id,"sku":product.sku,"name":product.name,"variant_count":len(variants),"card_count":cards,"variants":[{"id":v.id,"sku":v.sku,"color":v.color,"size":v.size,"specification":v.specification,"pricing":_variant_price(product,v,pricing or {})} for v in variants]})
    groups=[]
    flattened=[{"product_id":row["product_id"],"sku":row["sku"],"card_index":index+1} for row in rows for index in range(row["card_count"])]
    for offset in range(0,len(flattened),30): groups.append({"index":len(groups)+1,"card_count":len(flattened[offset:offset+30]),"cards":flattened[offset:offset+30]})
    errors=[{"sku":variant["sku"],"error":variant["pricing"]["error"]} for row in rows for variant in row["variants"] if variant["pricing"].get("error")]
    return {"product_count":len(rows),"variant_count":variant_total,"card_count":card_total,"group_count":len(groups),"split_size_cards":split_sizes,"products":rows,"groups":groups,"pricing_errors":errors}


def _write_excel(template_path: Path, output: Path, groups: list[dict[str, Any]], subject_name: str) -> int:
    template_limits=_template_value_limits(template_path)
    workbook=load_workbook(template_path);sheet=workbook[workbook.sheetnames[0]];columns=_defined_columns(workbook,sheet.title)
    char_columns={int(m.group(1)):col for name,col in columns.items() if (m:=re.fullmatch(r"characteristic_(\d+)_\d+",name))}
    prefix=datetime.now().strftime("%y%m%d%H%M%S")
    cards=sorted((card for group in groups for card in group["variants"]),key=_card_sku_sort_key)
    rows=[(int(f"{prefix}{offset//30+1:03d}"),card,card["images"]) for offset,card in enumerate(cards)]
    if not rows or len(rows)>1000: raise AppError("单个 WB 表格必须包含 1–1000 个商品变体",code="wb_export_size")
    maximum=max(columns.values());template_row=5
    for offset,(group_id,card,images) in enumerate(rows):
        row=template_row+offset
        if row!=template_row:
            for col in range(1,maximum+1):
                source,target=sheet.cell(template_row,col),sheet.cell(row,col)
                if source.has_style: target._style=copy(source._style)
                target.number_format=source.number_format;target.alignment=copy(source.alignment);target.protection=copy(source.protection)
        def write(prefix_name:str,value:Any):
            col=_column(columns,prefix_name)
            if col is not None: sheet.cell(row,col,value)
        if not 1<=len(images)<=30 or any(urlparse(url).scheme!="https" for url in images): raise AppError(f"{card['vendorCode']} 缺少 1–30 张 HTTPS OSS 图片",code="wb_images_invalid")
        dimensions=card["dimensions"]
        write("static_groupID_",group_id);write("static_vendorCode_",card["vendorCode"]);write("static_subject_",subject_name);write("static_photo_",";".join(images));write("static_price_",card["price"]);write("static_weightBrutto_",dimensions["weightBrutto"]);write("static_needKiz_","不需要");write("static_kizMarked_","否");write("static_isAdult_","不");write("static_minOrderQuantity_",1)
        fixed={15000000:card["title"],14177446:card["brand"],14177452:card["description"],90745:dimensions["width"],88953:dimensions["weightBrutto"],90849:dimensions["length"],90846:dimensions["height"],90702:1}
        for item in card.get("characteristics") or []:
            try: fixed[int(item.get("id"))]=item.get("value")
            except (TypeError,ValueError): continue
        for identity,value in fixed.items():
            if identity in char_columns and value not in (None,"",[]):
                limit=template_limits.get(identity)
                if limit:
                    entries=value if isinstance(value,list) else re.split(r"[;；\n]+",str(value))
                    value=[str(entry).strip() for entry in entries if str(entry).strip()][:limit]
                    if not value: continue
                sheet.cell(row,char_columns[identity],";".join(map(str,value)) if isinstance(value,list) else value)
    output.parent.mkdir(parents=True,exist_ok=True);workbook.save(output);return len(rows)


class _WbContentClient:
    """Synchronous Content API client used only by the background POD worker."""

    def __init__(self, config: dict[str, Any]) -> None:
        store=config.get("store") or {};secrets=store.get("secrets") or {}
        self.token=str(secrets.get("token") or secrets.get("api_key") or secrets.get("wildberries_token") or "").strip()
        if not self.token:raise AppError("WB 店铺未配置 Content API Token",code="wb_api_token_missing")
        base=str(store.get("base_url") or secrets.get("content_api_url") or "https://content-api.wildberries.ru").rstrip("/")
        self.http=httpx.Client(base_url=base,headers={"Authorization":self.token},timeout=60,follow_redirects=True)
        self._last_request=0.0
        self._last_media_request=0.0

    def close(self)->None:self.http.close()

    @staticmethod
    def _error(payload: Any, fallback: str)->str:
        if not isinstance(payload,dict):return fallback
        values=[str(payload.get(key) or "").strip() for key in ("errorText","message","detail","title")]
        additional=payload.get("additionalErrors")
        if isinstance(additional,dict):values.extend(f"{key}: {value}" for key,value in additional.items() if value)
        elif additional:values.append(str(additional))
        return "；".join(dict.fromkeys(value for value in values if value)) or fallback

    def request(self,method:str,path:str,*,params:dict[str,Any]|None=None,json_body:Any=None,retry_transient:bool=True)->dict[str,Any]:
        response: httpx.Response | None=None
        for attempt in range(6):
            elapsed=time.monotonic()-self._last_request
            if elapsed<0.65:time.sleep(0.65-elapsed)
            try:response=self.http.request(method,path,params=params,json=json_body)
            except httpx.HTTPError as error:
                if not retry_transient or attempt>=5:raise AppError(f"连接 WB Content API 失败：{error}",code="wb_api_network") from error
                time.sleep(min(8,2**attempt));continue
            self._last_request=time.monotonic()
            if response.status_code not in {429,500,502,503,504} or not retry_transient or attempt>=5:break
            try:delay=float(response.headers.get("X-RateLimit-Retry") or response.headers.get("Retry-After") or 2**attempt)
            except (TypeError,ValueError):delay=float(2**attempt)
            time.sleep(max(.7,min(30,delay)))
        assert response is not None
        try:value=response.json() if response.content else {}
        except ValueError as error:raise AppError(f"WB Content API 返回非 JSON 内容（HTTP {response.status_code}）",code="wb_api_invalid_response") from error
        if not response.is_success or (isinstance(value,dict) and value.get("error") is True):
            raise AppError(self._error(value,f"WB Content API HTTP {response.status_code}"),code="wb_api_rejected")
        return value if isinstance(value,dict) else {"data":value}

    def generate_barcodes(self,count:int)->list[str]:
        value=self.request("POST","/content/v2/barcodes",json_body={"count":count});rows=value.get("data") or []
        result=[str(item).strip() for item in rows if str(item).strip()]
        if len(result)!=count:raise AppError(f"WB 只返回 {len(result)}/{count} 个条码",code="wb_barcode_count")
        return result

    def create_cards(self,groups:list[dict[str,Any]])->None:
        # A timeout can happen after WB accepted this asynchronous request.
        # Do not repeat it here: the job retry first reconciles by vendorCode
        # and submits only cards that WB still does not have.
        self.request("POST","/content/v2/cards/upload",json_body=groups,retry_transient=False)

    def _search_cards(self,text:str)->list[dict[str,Any]]:
        value=self.request("POST","/content/v2/get/cards/list",json_body={"settings":{"filter":{"textSearch":text,"withPhoto":-1},"cursor":{"limit":100}}})
        return [dict(card) for card in value.get("cards") or [] if isinstance(card,dict)]

    def find_card(self,vendor_code:str,aliases:list[str]|None=None)->dict[str,Any]|None:
        """Find the current offer, then legacy offers containing an ERP SKU.

        Old stores used several prefixes and separators.  We deliberately do
        not guess when more than one WB card contains the same ERP identifier.
        """
        exact=self._search_cards(vendor_code)
        for card in exact:
            if str(card.get("vendorCode") or "").casefold()==vendor_code.casefold():
                card["_erp_match"]="current_offer";return card
        matches:dict[tuple[int,str],tuple[int,str,dict[str,Any]]]={}
        for alias in sorted({str(value).strip() for value in aliases or [] if str(value).strip()},key=len,reverse=True):
            needle=alias.casefold()
            for card in self._search_cards(alias):
                old_offer=str(card.get("vendorCode") or "").strip()
                if needle not in old_offer.casefold():continue
                identity=(int(card.get("nmID") or 0),old_offer.casefold())
                previous=matches.get(identity)
                if previous is None or len(alias)>previous[0]:matches[identity]=(len(alias),alias,card)
        if not matches:return None
        longest=max(value[0] for value in matches.values());best=[value for value in matches.values() if value[0]==longest]
        if len(best)!=1:
            raise AppError(f"ERP SKU {best[0][1]} 在 WB 匹配到多个历史商品卡，请人工核对后再上架",code="wb_legacy_sku_ambiguous")
        _,alias,card=best[0];card["_erp_match"]="legacy_contains_erp_sku";card["_erp_alias"]=alias;return card

    def list_all_cards(self,max_pages:int=500)->list[dict[str,Any]]:
        result=[];cursor={"limit":100}
        for _ in range(max_pages):
            value=self.request("POST","/content/v2/get/cards/list",json_body={"settings":{"sort":{"ascending":False},"filter":{"withPhoto":-1},"cursor":cursor}})
            page=[dict(card) for card in value.get("cards") or [] if isinstance(card,dict)];result.extend(page)
            next_cursor=value.get("cursor") if isinstance(value.get("cursor"),dict) else {}
            if len(page)<100 or not next_cursor.get("updatedAt") or not next_cursor.get("nmID"):break
            cursor={"limit":100,"updatedAt":next_cursor["updatedAt"],"nmID":next_cursor["nmID"]}
        return result

    def find_cards_exact(self,vendor_codes:set[str],max_pages:int=500)->dict[str,dict[str,Any]]:
        """Resolve exact seller SKUs through the cursor list.

        Newly created WB cards can appear in the seller cabinet and cursor
        listing before the text-search index sees them.  Walking newest-first
        also lets a retry recover the nmID without submitting duplicate cards.
        """
        wanted={str(value).strip().casefold():str(value).strip() for value in vendor_codes if str(value).strip()}
        found:dict[str,dict[str,Any]]={};cursor={"limit":100}
        for _ in range(max_pages):
            value=self.request("POST","/content/v2/get/cards/list",json_body={"settings":{"sort":{"ascending":False},"filter":{"withPhoto":-1},"cursor":cursor}})
            page=[dict(card) for card in value.get("cards") or [] if isinstance(card,dict)]
            for card in page:
                normalized=str(card.get("vendorCode") or "").strip().casefold()
                if normalized in wanted:
                    original=wanted[normalized];card["_erp_match"]="current_offer";found[original]=card
            if len(found)==len(wanted):break
            next_cursor=value.get("cursor") if isinstance(value.get("cursor"),dict) else {}
            if len(page)<100 or not next_cursor.get("updatedAt") or not next_cursor.get("nmID"):break
            cursor={"limit":100,"updatedAt":next_cursor["updatedAt"],"nmID":next_cursor["nmID"]}
        return found

    def errors(self)->dict[str,list[str]]:
        value=self.request("POST","/content/v2/cards/error/list",params={"locale":"ru"},json_body={"cursor":{"limit":100},"order":{"ascending":False}})
        data=value.get("data") if isinstance(value.get("data"),dict) else {};result={}
        for item in data.get("items") or []:
            errors=item.get("errors") if isinstance(item,dict) else None
            if not isinstance(errors,dict):continue
            for vendor,messages in errors.items():result[str(vendor)]=[str(message) for message in (messages if isinstance(messages,list) else [messages]) if message]
        return result

    def save_media(self,nm_id:int,urls:list[str])->None:
        # WB applies a stricter bucket to media replacement than to ordinary
        # Content API reads.  Keep enough distance between cards so a batch
        # does not exhaust that bucket after the first successful upload.
        elapsed=time.monotonic()-self._last_media_request
        if elapsed<3.1:time.sleep(3.1-elapsed)
        self.request("POST","/content/v3/media/save",json_body={"nmId":int(nm_id),"data":urls[:30]})
        self._last_media_request=time.monotonic()


def _api_card(card:dict[str,Any],barcodes:dict[str,str])->dict[str,Any]:
    # WB accepts package measurements only through ``dimensions``.  Some
    # category schemas still expose legacy characteristics for the same
    # values; sending both makes /content/v2/cards/upload reject the card.
    dimension_characteristic_ids={88952,88953,90745,90846,90849}
    characteristics=[
        item for item in (card.get("characteristics") or [])
        if int(item.get("id") or 0) not in dimension_characteristic_ids
        and int(item.get("id") or 0) not in WB_OMITTED_CHARACTERISTIC_IDS
    ]
    return {
        "vendorCode":card["vendorCode"],"kizMarked":False,"title":card["title"],"description":card["description"],
        "brand":card["brand"],"dimensions":card["dimensions"],"characteristics":characteristics,
        # The current POD subjects are dimensioned through category
        # characteristics, not WB apparel sizes.  Supplying techSize makes WB
        # treat the otherwise sizeless card as apparel and leaves it in Drafts.
        "sizes":[{"price":int(math.ceil(float(size["price"]))),"skus":[barcodes[str(size["variant_id"])]]} for size in card["sizes"]],
    }


def _publish_api(context:JobContext,config:dict[str,Any],root:Path,groups:list[dict[str,Any]],merged_groups:list[dict[str,Any]])->None:
    client=_WbContentClient(config)
    try:
        cards=[card for group in groups for card in group["variants"]]
        barcode_path=root/"barcodes.json"
        try:barcodes=json.loads(barcode_path.read_text(encoding="utf-8"))
        except (OSError,ValueError):barcodes={}
        variant_ids=[str(size["variant_id"]) for card in cards for size in card["sizes"]]
        missing_ids=[identity for identity in variant_ids if not str(barcodes.get(identity) or "")]
        context.start_step("barcodes")
        if missing_ids:
            generated=client.generate_barcodes(len(missing_ids));barcodes.update(dict(zip(missing_ids,generated,strict=True)))
            temporary=barcode_path.with_suffix(".tmp");temporary.write_text(json.dumps(barcodes,ensure_ascii=False,indent=2),encoding="utf-8");os.replace(temporary,barcode_path)
        context.complete_step("barcodes",details={"count":len(variant_ids),"reused":len(variant_ids)-len(missing_ids)})
        card_by_vendor={card["vendorCode"]:card for card in cards};found=client.find_cards_exact(set(card_by_vendor));preexisting:set[str]=set(found)
        context.start_step("submit")
        for index,vendor in enumerate(card_by_vendor,1):
            context.refresh();source=card_by_vendor[vendor]
            if vendor not in found:
                aliases=[source.get("product_sku"),*(size.get("sku") for size in source.get("sizes") or [])]
                existing=client.find_card(vendor,[str(value) for value in aliases if value])
                if existing and int(existing.get("nmID") or 0)>0:found[vendor]=existing;preexisting.add(vendor)
            context.set_progress("submit",index*35/max(1,len(card_by_vendor)),f"已核对 {index}/{len(card_by_vendor)} 个卖家 SKU")
        missing=set(card_by_vendor)-set(found)
        if missing:
            # The cards/upload request is asynchronous. A disconnected client
            # cannot prove whether WB accepted it, and search indexing can lag.
            # Persist the exact write intent before the call; after a restart,
            # only reconcile it and never submit the same intent again.
            intent_path=root/'api-submit-intent.json'
            if intent_path.exists():
                try: intent=json.loads(intent_path.read_text(encoding='utf-8'))
                except (OSError,ValueError): raise AppError('WB 建品请求记录损坏，请人工核对',code='wb_cards_outcome_unknown')
                if set(intent.get('vendors') or [])!=missing:
                    raise AppError('WB 建品结果尚未全部可见，请核对后恢复，禁止重复提交',code='wb_cards_outcome_unknown')
                # The same unresolved set may mean accepted-but-not-indexed;
                # defer instead of repeating the platform write.
                raise AppError('WB 建品结果尚未可见，请稍后核对原任务，禁止重复提交',code='wb_cards_outcome_unknown')
            request_groups=[]
            for group in merged_groups:
                values=[_api_card(card,barcodes) for card in group["variants"] if card["vendorCode"] in missing]
                if values:request_groups.append({"subjectID":int(context.payload["subject_id"]),"variants":values})
            temporary=intent_path.with_suffix('.tmp')
            temporary.write_text(json.dumps({'vendors':sorted(missing),'state':'pending'},ensure_ascii=False),encoding='utf-8')
            os.replace(temporary,intent_path)
            client.create_cards(request_groups)
            temporary.write_text(json.dumps({'vendors':sorted(missing),'state':'accepted'},ensure_ascii=False),encoding='utf-8')
            os.replace(temporary,intent_path)
            for ledger in context.session.scalars(select(WbPublicationItem).where(WbPublicationItem.job_id==context.job.id)).all():
                if ledger.status!="failed":ledger.status="submitted"
            context.session.commit();context.event(f"WB 已接收 {len(missing)} 张商品卡，开始等待商品编号")
        context.complete_step("submit",details={"cards":len(cards),"already_existed":len(found),"submitted":len(missing)})
        context.start_step("verify");unresolved=set(missing)
        for attempt in range(18):
            if not unresolved:break
            if attempt:time.sleep(10)
            context.refresh()
            resolved=client.find_cards_exact(unresolved)
            for vendor,card in resolved.items():
                if int(card.get("nmID") or 0)>0:found[vendor]=card;unresolved.remove(vendor)
            context.set_progress("verify",len(found)*60/max(1,len(cards)),f"已取得 {len(found)}/{len(cards)} 个 WB 商品编号")
        if unresolved:
            try:errors=client.errors()
            except AppError:errors={}
            messages=[f"{vendor}: {'；'.join(errors[vendor])}" for vendor in sorted(unresolved) if vendor in errors]
            raise AppError("WB 仍未生成部分商品卡"+("："+"；".join(messages[:5]) if messages else "，任务会从核对阶段自动重试"),code="wb_cards_pending")
        ledgers=context.session.scalars(select(WbPublicationItem).where(WbPublicationItem.job_id==context.job.id)).all();ledger_by_variant={row.variant_id:row for row in ledgers if row.status!="failed"}
        media_checkpoint_path=root/"media-checkpoint.json"
        try:media_checkpoint=json.loads(media_checkpoint_path.read_text(encoding="utf-8"))
        except (OSError,ValueError):media_checkpoint={}
        uploaded=0;media_reused=0
        for index,(vendor,card) in enumerate(card_by_vendor.items(),1):
            context.refresh();nm_id=int(found[vendor]["nmID"])
            # The media endpoint replaces the card gallery in the supplied
            # order.  Always apply the current task snapshot, including when
            # the WB card already exists; otherwise a changed main/additional
            # image order is silently ignored and WB keeps the old gallery.
            media_signature=hashlib.sha256(json.dumps({"nm_id":nm_id,"images":card["images"]},ensure_ascii=False,separators=(",",":")).encode("utf-8")).hexdigest()
            if media_checkpoint.get(vendor)!=media_signature:
                client.save_media(nm_id,card["images"]);uploaded+=1
                media_checkpoint[vendor]=media_signature
                temporary=media_checkpoint_path.with_suffix(".tmp");temporary.write_text(json.dumps(media_checkpoint,ensure_ascii=False,indent=2),encoding="utf-8");os.replace(temporary,media_checkpoint_path)
            else:media_reused+=1
            for size in card["sizes"]:
                ledger=ledger_by_variant.get(str(size["variant_id"]))
                if ledger:ledger.status="pending_stock";ledger.external_product_id=str(nm_id);ledger.last_error="";ledger.confirmed_at=None
            context.session.commit();context.set_progress("verify",60+index*40/max(1,len(cards)),f"已完成 {index}/{len(cards)} 张商品卡的图片核对")
        context.complete_step("verify",details={"cards":len(cards),"media_uploaded":uploaded,"media_reused":media_reused,"nm_ids":list(dict.fromkeys(int(item["nmID"]) for item in found.values()))})
    finally:client.close()


def _queue_removal_reason(item: WbPublicationItem, job: Job | None) -> str:
    if not job or job.status not in {"cancelled", "failed"}:
        return "请先在任务中心停止任务，并等待任务结束"
    if item.external_product_id or item.status not in {"queued", "running", "failed", "cancelled", "withdrawn"}:
        return "可能已提交 WB，请先核对平台结果"
    if any(step.key in {"submit", "verify", "excel"} and (step.started_at or step.status != "pending") for step in job.steps):
        return "已进入提交或导出步骤，请先核对平台结果"
    return ""


def remove_publication_from_queue(session: Session, ledger_id: str) -> dict[str, Any]:
    item = session.get(WbPublicationItem, ledger_id)
    if not item:
        raise NotFoundError("WB 上架记录不存在")
    if item.status == "withdrawn":
        return {"id": item.id, "status": "withdrawn"}
    job = session.get(Job, item.job_id) if item.job_id else None
    if job:
        session.execute(update(Job).where(Job.id == job.id).values(status=Job.status))
        session.refresh(job)
        session.refresh(item)
    reason = _queue_removal_reason(item, job)
    if reason:
        raise ConflictError(reason)
    item.status = "withdrawn"
    item.last_error = "用户移除已停止的上架记录；未删除产品库或 WB 商品"
    payload = loads(job.payload_json, {})
    payload["wb_queue_removed"] = True
    job.payload_json = dumps(payload)
    session.add(JobEvent(job_id=job.id, message=f"已从 WB 上架队列移除 {item.offer_id}；原任务禁止重试，请重新创建所需上架任务"))
    session.commit()
    return {"id": item.id, "status": item.status}


def list_publications(session:Session,stage:str,*,page:int=1,page_size:int=30,search:str="",store_id:str="")->dict[str,Any]:
    allowed={"processing","pending_price","price_processing","pending_stock","stock_processing","completed","exceptions"}
    if stage not in allowed:raise AppError("未知的 WB 流程阶段",code="wb_stage_invalid")
    statuses={
        "processing":{"queued","running","submitted"},
        "pending_price":{"pending_price","price_processing"},
        "price_processing":{"price_processing"},
        "pending_stock":{"pending_stock","stock_processing"},
        "stock_processing":{"stock_processing"},
        "completed":{"completed"},
        "exceptions":{"failed","price_failed","stock_failed","sync_conflict","sync_unknown"},
    }[stage]
    condition=WbPublicationItem.status.in_(statuses)
    query=select(WbPublicationItem).join(CatalogProduct,CatalogProduct.id==WbPublicationItem.product_id).join(CatalogVariant,CatalogVariant.id==WbPublicationItem.variant_id).join(Store,Store.id==WbPublicationItem.store_id).where(condition)
    if store_id:
        query=query.where(WbPublicationItem.store_id==store_id)
    term=search.strip()
    if term:
        pattern=f"%{term}%"
        query=query.where(or_(CatalogProduct.name.ilike(pattern),CatalogProduct.sku.ilike(pattern),CatalogVariant.sku.ilike(pattern),WbPublicationItem.offer_id.ilike(pattern),WbPublicationItem.external_product_id.ilike(pattern),Store.name.ilike(pattern)))
    total=int(session.scalar(select(func.count()).select_from(query.subquery())) or 0)
    page=max(1,int(page));page_size=max(1,min(200,int(page_size)))
    rows=session.execute(
        select(WbPublicationItem,CatalogProduct,CatalogVariant,Store,Job)
        .join(CatalogProduct,CatalogProduct.id==WbPublicationItem.product_id)
        .join(CatalogVariant,CatalogVariant.id==WbPublicationItem.variant_id)
        .join(Store,Store.id==WbPublicationItem.store_id)
        .outerjoin(Job,Job.id==WbPublicationItem.job_id)
        .where(WbPublicationItem.id.in_(query.with_only_columns(WbPublicationItem.id)))
        .order_by(*([case((WbPublicationItem.status == stage, 0), else_=1)]
                    if stage in {'pending_price','pending_stock'} else []),
                  WbPublicationItem.updated_at.desc(), WbPublicationItem.id)
        .offset((page-1)*page_size)
        .limit(page_size)
    ).all()
    from erp.domains.pod.wb_price_state import latest_operations, blocked_reason
    operations=latest_operations(session,[row[0] for row in rows]) if stage in {'pending_price','price_processing','exceptions'} else {}
    items=[]
    for ledger,product,variant,store,job in rows:
        operation=operations.get(ledger.id)
        payload=loads(job.payload_json,{}) if job else {}
        pricing=_variant_price(product,variant,payload['pricing']) if payload.get('pricing') else {}
        planned_row=next((r for r in payload.get('rows',[]) if r.get('id')==ledger.id),{})
        items.append({
            "id":ledger.id,"status":_effective_ledger_status(session,ledger),"last_error":ledger.last_error,"updated_at":ledger.updated_at,
            "can_remove":not bool(_queue_removal_reason(ledger,job)),"remove_blocked_reason":_queue_removal_reason(ledger,job),
            "operation_job_id":operation.id if operation else "",
            "operation_job_status":operation.status if operation else "",
            "price_submit_blocked_reason":blocked_reason(ledger,operation),
            "job_id":ledger.job_id,"job_status":job.status if job else "","store_id":store.id,"store_name":store.name,
            "product_id":product.id,"product_name":product.name,"product_sku":product.sku,"variant_id":variant.id,
            "sku":variant.sku,"size":variant.size or variant.specification or "单一规格","offer_id":ledger.offer_id,
            "nm_id":ledger.external_product_id,"price":planned_row.get('price',pricing.get("list_price_rub")),
            "discount_percent":int(planned_row.get('discount',(payload.get("pricing") or {}).get("discount_percent")) or 0),
        })
    return {"stage":stage,"total":total,"page":page,"page_size":page_size,"items":items}


def _wb_store_snapshot(session:Session,store_id:str)->tuple[Store,str]:
    store=session.get(Store,store_id)
    if not store or store.platform.casefold() not in {"wb","wildberries"} or not store.enabled:raise ConflictError("请选择已启用的 WB 店铺")
    secrets=vault.decrypt(store.secrets_encrypted) if store.secrets_encrypted else {}
    if store.connection and store.connection.secrets_encrypted:secrets={**vault.decrypt(store.connection.secrets_encrypted),**secrets}
    if not str(secrets.get("token") or secrets.get("api_key") or secrets.get("wildberries_token") or "").strip():raise AppError("当前 WB 店铺未配置 API Token",code="wb_api_token_missing")
    return store,vault.encrypt({"store":{"id":store.id,"name":store.name,"offer_prefix":store.offer_prefix,"secrets":secrets}})


def wb_warehouses(session:Session,store_id:str)->dict[str,Any]:
    store,snapshot=_wb_store_snapshot(session,store_id);config=vault.decrypt(snapshot);secrets=config["store"]["secrets"]
    token=str(secrets.get("token") or secrets.get("api_key") or secrets.get("wildberries_token"))
    response=httpx.get("https://marketplace-api.wildberries.ru/api/v3/warehouses",headers={"Authorization":token},timeout=60)
    if response.status_code!=200:raise AppError(f"读取 WB 仓库失败（HTTP {response.status_code}）",code="wb_warehouses_failed")
    return {"store_id":store.id,"items":response.json()}


def enqueue_stock(session:Session,ledger_ids:list[str],warehouse_id:int,amount:int)->Job:
    ids=list(dict.fromkeys(ledger_ids));rows=session.scalars(select(WbPublicationItem).where(WbPublicationItem.id.in_(ids))).all()
    if len(rows)!=len(ids) or not rows:raise NotFoundError("部分 WB 流程商品不存在")
    if len({row.store_id for row in rows})!=1:raise AppError("一次库存任务只能处理同一个 WB 店铺",code="wb_store_mixed")
    if any(row.status not in {"pending_stock","stock_failed"} for row in rows):raise ConflictError("所选商品不在待添加库存阶段")
    store,snapshot=_wb_store_snapshot(session,rows[0].store_id)
    job=enqueue_job(session,job_type="pod_wb_stock",payload={"ledger_ids":ids,"store_id":store.id,"warehouse_id":int(warehouse_id),"amount":int(amount)},config_snapshot=snapshot,idempotency_key=f"pod-wb-stock:{uuid.uuid4().hex}",max_attempts=1,steps=[("resolve","读取 WB 规格 ID"),("stock","批量设置库存")],commit=False)
    for row in rows:row.status="stock_processing";row.last_error=""
    session.commit();return job


def enqueue_price(session:Session,ledger_ids:list[str],price:float|None,discount_percent:int,*,verify_only:bool=False)->Job:
    ids=list(dict.fromkeys(ledger_ids));rows=session.scalars(select(WbPublicationItem).where(WbPublicationItem.id.in_(ids)).with_for_update().execution_options(populate_existing=True)).all()
    if len(rows)!=len(ids) or not rows:raise NotFoundError("部分 WB 流程商品不存在")
    if len({row.store_id for row in rows})!=1:raise AppError("一次价格任务只能处理同一个 WB 店铺",code="wb_store_mixed")
    allowed={"price_failed","pending_price"} if verify_only else {"pending_price","price_failed"}
    if any(row.status not in allowed for row in rows):raise ConflictError("所选商品不在可核对的价格阶段")
    if not verify_only and any(row.status == 'price_failed' and 'WB 实际价格或折扣与计划不一致；本次未写入' not in row.last_error for row in rows):
        raise ConflictError('原价格请求结果未核清，先只读核对后才可人工重新提交')
    if not verify_only and any(row.last_error.startswith('WB 上传编号 ') for row in rows):
        raise ConflictError('原 WB 价格上传已受理；请先使用只读核对原价格，禁止直接重复提交')
    store,snapshot=_wb_store_snapshot(session,rows[0].store_id)
    job=enqueue_job(session,job_type="pod_wb_price",payload={"ledger_ids":ids,"store_id":store.id,"price":price,"discount_percent":int(discount_percent),"verify_only":verify_only},config_snapshot=snapshot,idempotency_key=f"pod-wb-price:{uuid.uuid4().hex}",max_attempts=1,steps=[("verify","只读核对 WB 价格与折扣")] if verify_only else [("submit","提交价格和折扣"),("verify","等待 WB 处理完成")],commit=False)
    from erp.domains.pod.wb_price_recovery import freeze_plan
    freeze_plan(session, job, rows)
    for row in rows:row.status="price_processing";row.last_error=""
    session.commit();return job


@job_handler("pod_wb_stock")
def wb_stock_job(context:JobContext)->None:
    rows=context.session.scalars(select(WbPublicationItem).where(WbPublicationItem.id.in_(context.payload["ledger_ids"]))).all();client=_WbContentClient(vault.decrypt(context.job.config_snapshot_json))
    try:
        context.start_step("resolve");stocks=[]
        cards=client.find_cards_exact({row.offer_id for row in rows})
        for row in rows:
            card=cards.get(row.offer_id)
            if not card:raise AppError(f"未找到 WB 商品卡：{row.offer_id}",code="wb_card_missing")
            sizes=card.get("sizes") or []
            if len(sizes)!=1 or not int(sizes[0].get("chrtID") or sizes[0].get("chrtId") or 0):raise AppError(f"无法确定库存规格 ID：{row.offer_id}",code="wb_chrt_missing")
            stocks.append({"chrtId":int(sizes[0].get("chrtID") or sizes[0].get("chrtId")),"amount":int(context.payload["amount"])})
        context.complete_step("resolve",details={"items":len(stocks)});context.start_step("stock")
        token=client.token;response=httpx.put(f"https://marketplace-api.wildberries.ru/api/v3/stocks/{int(context.payload['warehouse_id'])}",headers={"Authorization":token},json={"stocks":stocks},timeout=60)
        if response.status_code!=204:raise AppError(f"WB 库存设置失败（HTTP {response.status_code}）：{response.text[:500]}",code="wb_stock_failed")
        for row in rows:row.status="pending_price";row.last_error=""
        context.session.commit();context.complete_step("stock",details={"items":len(stocks),"warehouse_id":context.payload["warehouse_id"],"amount":context.payload["amount"]})
    except Exception as error:
        for row in rows:row.status="stock_failed";row.last_error=str(error)
        context.session.commit();raise
    finally:client.close()


@job_handler("pod_wb_price")
def wb_price_job(context:JobContext)->None:
    rows=context.session.scalars(select(WbPublicationItem).where(WbPublicationItem.id.in_(context.payload["ledger_ids"]))).all();config=vault.decrypt(context.job.config_snapshot_json);secrets=config["store"]["secrets"]
    token=str(secrets.get("token") or secrets.get("api_key") or secrets.get("wildberries_token"));data=[]
    try:
        for row in rows:
            source=context.session.get(Job,row.job_id);payload=loads(source.payload_json,{}) if source else {};product=context.session.get(CatalogProduct,row.product_id);variant=context.session.get(CatalogVariant,row.variant_id)
            value=float(context.payload.get("price") or _variant_price(product,variant,payload.get("pricing") or {}).get("list_price_rub") or 0)
            if value<=0:raise AppError(f"{row.offer_id} 无法计算价格",code="wb_pricing_invalid")
            data.append({"nmID":int(row.external_product_id),"price":int(math.ceil(value)),"discount":int(context.payload["discount_percent"])})
        context.start_step("submit");http=httpx.Client(base_url="https://discounts-prices-api.wildberries.ru",headers={"Authorization":token},timeout=60);response=http.post("/api/v2/upload/task",json={"data":data})
        if response.status_code not in {200,208}:raise AppError(f"WB 价格提交失败（HTTP {response.status_code}）：{response.text[:500]}",code="wb_price_failed")
        upload_id=int((response.json().get("data") or {}).get("id") or 0);context.complete_step("submit",details={"items":len(data),"upload_id":upload_id});context.start_step("verify")
        verified=False
        for attempt in range(18):
            if attempt:time.sleep(5)
            status=http.get("/api/v2/history/tasks",params={"uploadID":upload_id})
            if status.status_code==200:
                value=status.json().get("data") or {}
                if int(value.get("status") or 0)==3 and int(value.get("successGoodsNumber") or 0)==len(data):verified=True;break
                if int(value.get("status") or 0) in {4,5,6}:raise AppError(f"WB 价格任务失败：{value}",code="wb_price_failed")
        http.close()
        if not verified:raise AppError("WB 价格与折扣仍在处理，任务将自动重试",code="wb_price_pending")
        for row in rows:row.status="completed";row.last_error="";row.confirmed_at=datetime.now().astimezone()
        context.session.commit();context.complete_step("verify",details={"items":len(data),"upload_id":upload_id})
    except Exception as error:
        for row in rows:row.status="price_failed";row.last_error=str(error)
        context.session.commit();raise


def _contained_longest(value:str,rows:list[Any],attribute:str)->list[Any]:
    normalized=value.casefold();matched=[]
    for row in rows:
        identity=str(getattr(row,attribute,"") or "").strip()
        if identity and identity.casefold() in normalized:matched.append(row)
    if not matched:return []
    maximum=max(len(str(getattr(row,attribute) or "").strip()) for row in matched)
    return [row for row in matched if len(str(getattr(row,attribute) or "").strip())==maximum]


@job_handler("pod_wb_history_reconcile")
def wb_history_reconcile_job(context:JobContext)->None:
    if context.payload.get('shop_sync'):
        raise ConflictError('店铺完整同步仅允许原发起电脑执行')
    config=vault.decrypt(context.job.config_snapshot_json);client=_WbContentClient(config)
    try:
        context.start_step("download");cards=client.list_all_cards();context.complete_step("download",details={"cards":len(cards)})
    finally:client.close()
    context.start_step("match")
    summary=apply_history_cards(context.session,context.job,cards)
    context.complete_step("match",details=summary)
    context.complete_step("ledger",details=summary)


def apply_history_cards(session:Session,job:Job,cards:list[dict[str,Any]])->dict[str,int]:
    """Apply a completed shop scan; never call WB or execute on the server."""
    payload=loads(job.payload_json,{})
    products=session.scalars(select(CatalogProduct).where(CatalogProduct.source_type=="pod").options(selectinload(CatalogProduct.variants))).all()
    variants=[variant for product in products for variant in product.variants if variant.active]
    workflow=bool(payload.get("stock_workflow"))
    prefix=str(payload.get("offer_prefix") or "").rstrip("-")
    exact={f"{prefix}-{variant.sku}".casefold():variant for variant in variants}
    ledger_offers={row.offer_id.casefold():row.variant_id for row in session.scalars(select(WbPublicationItem).where(WbPublicationItem.store_id==payload["store_id"])).all()} if workflow else {}
    variant_by_id={variant.id:variant for variant in variants}
    proposals:dict[str,list[dict[str,Any]]]={};unmatched=[];ambiguous=[]
    for index,card in enumerate(cards,1):
        vendor=str(card.get("vendorCode") or "").strip();nm_id=int(card.get("nmID") or 0)
        if not vendor or nm_id<=0:continue
        chosen=[];reason=""
        if vendor.casefold() in ledger_offers and ledger_offers[vendor.casefold()] in variant_by_id:chosen=[variant_by_id[ledger_offers[vendor.casefold()]]];reason="existing_offer"
        elif vendor.casefold() in exact:chosen=[exact[vendor.casefold()]];reason="current_offer"
        else:
            variant_matches=_contained_longest(vendor,variants,"sku")
            if len(variant_matches)==1:chosen=variant_matches;reason="legacy_variant_sku"
            elif len(variant_matches)>1:ambiguous.append({"vendor_code":vendor,"nm_id":nm_id,"reason":"multiple_variant_skus","erp_skus":[row.sku for row in variant_matches]})
            else:
                product_matches=_contained_longest(vendor,products,"sku")
                if len(product_matches)==1:chosen=[row for row in product_matches[0].variants if row.active];reason="legacy_product_sku"
                elif len(product_matches)>1:ambiguous.append({"vendor_code":vendor,"nm_id":nm_id,"reason":"multiple_product_skus","erp_skus":[row.sku for row in product_matches]})
        if not chosen:
            if not any(row.get("vendor_code")==vendor for row in ambiguous):unmatched.append({"vendor_code":vendor,"nm_id":nm_id})
        else:
            for variant in chosen:proposals.setdefault(variant.id,[]).append({"card":card,"reason":reason})
    accepted={}
    for variant_id,items in proposals.items():
        unique={int(item["card"].get("nmID") or 0) for item in items}
        if len(unique)==1:accepted[variant_id]=items[0]
        else:ambiguous.append({"variant_id":variant_id,"reason":"one_erp_sku_multiple_cards","wb_cards":[{"vendor_code":item["card"].get("vendorCode"),"nm_id":item["card"].get("nmID")} for item in items]})
    existing={row.variant_id:row for row in session.scalars(select(WbPublicationItem).where(WbPublicationItem.store_id==payload["store_id"],WbPublicationItem.variant_id.in_(list(accepted)))).all()} if accepted else {}
    written=0
    for variant in variants:
        match=accepted.get(variant.id)
        if not match:continue
        nm_id=str(int(match["card"].get("nmID") or 0));row=existing.get(variant.id);offer=f"{prefix}-{variant.sku}"
        if row and row.external_product_id and row.external_product_id!=nm_id:
            ambiguous.append({"variant_id":variant.id,"erp_sku":variant.sku,"reason":"ledger_nm_id_conflict","ledger_nm_id":row.external_product_id,"wb_nm_id":nm_id});continue
        if row:
            if workflow:
                if row.status in {"exported_pending_confirmation","submitted"} or (row.status=="completed" and not row.external_product_id):row.status="pending_stock"
                if not row.job_id:row.job_id=job.id
                row.external_product_id=nm_id;row.last_error="";row.confirmed_at=datetime.now().astimezone()
            elif not row.external_product_id:
                # An existing publication keeps its own job, stage and offer.
                # Historical recognition must not steal an active listing.
                row.external_product_id=nm_id
        else:session.add(WbPublicationItem(store_id=payload["store_id"],product_id=variant.product_id,variant_id=variant.id,job_id=job.id,template_id="",offer_id=str(match["card"].get("vendorCode") or offer) if workflow else offer,mode="history",status="pending_stock" if workflow else "completed",external_product_id=nm_id,confirmed_at=datetime.now().astimezone()))
        written+=1
    root=_roots()[1]/job.id;root.mkdir(parents=True,exist_ok=True);report=root/"WB历史POD匹配报告.json"
    report.write_text(json.dumps({"store_id":payload["store_id"],"store_name":payload["store_name"],"summary":{"wb_cards":len(cards),"matched_variants":written,"unmatched_cards":len(unmatched),"ambiguous":len(ambiguous)},"ambiguous":ambiguous,"unmatched":unmatched},ensure_ascii=False,indent=2),encoding="utf-8")
    asset=Asset(kind="file",name=report.name,source="pod_wb_history_reconcile",local_path=str(report),sha256=hashlib.sha256(report.read_bytes()).hexdigest(),mime_type="application/json",size_bytes=report.stat().st_size,read_only=True);session.add(asset);session.flush();session.add(JobArtifact(job_id=job.id,asset_id=asset.id,role="wb_reconcile_report"))
    return {"wb_cards":len(cards),"matched_variants":written,"unmatched_cards":len(unmatched),"ambiguous":len(ambiguous)}


def _prepare_copies_parallel(context, products, config, template, payload, copies, checkpoint, failures=None, failure_checkpoint=None):
    from concurrent.futures import ThreadPoolExecutor, wait, FIRST_COMPLETED
    from threading import Event
    from erp.core.ai_concurrency import limits, cancellation
    from sqlalchemy.orm import sessionmaker
    # Each desktop task owns a private database. The process-global factory
    # may still point at a previous task or the full local ERP database.
    SessionLocal = sessionmaker(bind=context.session.get_bind())

    # JobContext.refresh() expires ORM state in the owning session. Reading
    # context.job from a copy thread can then issue SQL on that shared session.
    # Snapshot all identifiers before starting any threads.
    job_id, job_type, job_attempt = context.job.id, context.job.type, context.job.attempts
    product_ids = tuple(product.id for product in products)
    failures = failures if failures is not None else {}
    missing = iter(identity for identity in product_ids if not isinstance(copies.get(identity), dict) and identity not in failures)
    stop = Event()
    count = sum(isinstance(copies.get(identity), dict) or identity in failures for identity in product_ids)
    failure = None

    def generate(identity):
        # ORM objects and sessions never cross thread boundaries.
        from erp.core.ai_usage import task_scope
        with task_scope(job_id, job_type, job_attempt), cancellation(stop), SessionLocal() as session:
            if stop.is_set():
                raise RuntimeError("批次已停止")
            product = session.get(CatalogProduct, identity)
            if product is None:
                raise NotFoundError("文案任务产品不存在")
            references = (payload['_copy_references_by_product'].get(identity, [])
                          if '_copy_references_by_product' in payload
                          else _previous_copy_references(session,identity,job_id))
            for attempt in range(1, 4):
                if stop.is_set():
                    raise RuntimeError("批次已停止")
                try:
                    return _generate_copy(config, product, template, {**payload,"_copy_references":references})
                except Exception as error:
                    if getattr(error, 'code', '') == 'ai_write_outcome_unknown':
                        raise
                    if attempt == 3:
                        return {"sku": product.sku, "name": product.name, "error": str(error)[:1000], "attempts": attempt}
                    if stop.wait(attempt):
                        raise RuntimeError("批次已停止")
            raise AssertionError("文案重试次数异常")

    with ThreadPoolExecutor(max_workers=limits()["text"], thread_name_prefix="wb-copy") as pool:
        pending = {}
        def submit_next():
            identity = next(missing, None)
            if identity is not None:
                pending[pool.submit(generate, identity)] = identity
        for _ in range(limits()["text"]):
            submit_next()
        while pending:
            if failure is None:
                try:
                    context.refresh()
                except Exception as error:
                    failure = error; stop.set()
            done, _ = wait(pending, timeout=.2, return_when=FIRST_COMPLETED)
            for future in done:
                identity = pending.pop(future)
                try:
                    result = future.result()
                    if "error" in result:
                        failures[identity] = result
                        if failure_checkpoint is not None:
                            temporary = failure_checkpoint.with_suffix(".tmp")
                            temporary.write_text(json.dumps(failures, ensure_ascii=False, indent=2), encoding="utf-8")
                            os.replace(temporary, failure_checkpoint)
                    else:
                        copies[identity] = result
                        temporary = checkpoint.with_suffix(".tmp")
                        temporary.write_text(json.dumps(copies, ensure_ascii=False, indent=2), encoding="utf-8")
                        os.replace(temporary, checkpoint)
                    count += 1
                    if failure is None:
                        context.set_progress("copy", count * 100 / max(1, len(products)), f"已处理 {count}/{len(products)} 个产品文案，失败 {len(failures)} 个")
                except Exception as error:
                    if failure is None:
                        failure = error; stop.set()
            if failure is None:
                for _ in done:
                    submit_next()
        if failure is not None:
            raise failure


@job_handler("pod_wb_api_publish")
@job_handler("pod_wb_excel_export")
def wb_export_job(context: JobContext) -> None:
    payload=context.payload
    if payload.get("generation_provider")=="codex" and (not payload.get("_codex_outputs_by_product") or (payload.get("mode")=="api" and not payload.get("codex_publish_approved"))):
        raise AppError("请先完成 Codex 回填，并在原页面确认 API 发布",code="codex_output_missing")
    template=_template(str(payload["template_id"]));template_path=Path(template["path"])
    if hashlib.sha256(template_path.read_bytes()).hexdigest()!=payload["template_sha256"]: raise ConflictError("WB 模板已变化，请重新创建任务")
    products=context.session.scalars(select(CatalogProduct).where(CatalogProduct.id.in_(payload["product_ids"]),CatalogProduct.source_type=="pod")).all()
    if len(products)!=len(payload["product_ids"]): raise NotFoundError("任务中的部分 POD 产品不存在")
    products=sorted(products,key=lambda product:(_sku_sort_key(product.sku),product.id))
    for ledger in context.session.scalars(select(WbPublicationItem).where(WbPublicationItem.job_id==context.job.id)).all():
        if not (str(payload.get("mode") or "excel")=="api" and ledger.status=="completed"):ledger.status="running"
    context.session.commit()
    config=vault.decrypt(context.job.config_snapshot_json);live_schema=_wb_characteristics(config,int(payload["subject_id"]))
    if not live_schema:
        numeric_ids={90745,88953,88952,90849,90846,90673,90630,90602,90675,90702}
        live_schema=[{"charcID":item["id"],"name":item.get("name") or str(item["id"]),"charcType":4 if int(item["id"]) in numeric_ids else 1,"required":False,"maxCount":0} for item in (template.get("schema") or {}).get("characteristics") or []]
    template_limits=_template_value_limits(template_path)
    for item in live_schema:
        limit=template_limits.get(int(item.get("charcID") or 0))
        if limit:
            live_limit=int(item.get("maxCount") or 0)
            item["maxCount"]=min(live_limit,limit) if live_limit>0 else limit
    prompt_template={**template,"schema":{**(template.get("schema") or {}),"characteristics":[{"id":int(item.get("charcID") or 0),"name":item.get("name"),"required":bool(item.get("required")),"max_count":int(item.get("maxCount") or 0),"type":int(item.get("charcType") or 0)} for item in live_schema if int(item.get("charcID") or 0)>0 and int(item.get("charcID") or 0) not in WB_OMITTED_CHARACTERISTIC_IDS]}}
    context.start_step("validate");context.complete_step("validate",details={"products":len(products),"template":template["filename"],"live_characteristics":len(live_schema)})
    root=_roots()[1]/context.job.id;root.mkdir(parents=True,exist_ok=True);checkpoint=root/"copy.json";failure_checkpoint=root/"copy-failures.json"
    try: copies=json.loads(checkpoint.read_text(encoding="utf-8"))
    except (OSError,ValueError): copies={}
    try: copy_failures=json.loads(failure_checkpoint.read_text(encoding="utf-8"))
    except (OSError,ValueError): copy_failures={}
    groups=[];context.start_step("copy")
    _prepare_copies_parallel(context, products, config, prompt_template, payload, copies, checkpoint, copy_failures, failure_checkpoint)
    if copy_failures:
        failed_ids=set(copy_failures)
        for ledger in context.session.scalars(select(WbPublicationItem).where(WbPublicationItem.job_id==context.job.id)).all():
            if ledger.product_id in failed_ids:
                ledger.status="failed"
                ledger.last_error=f"文案生成尝试 3 次后失败：{copy_failures[ledger.product_id]['error']}"[:1000]
        context.session.commit()
        if not any(item.role=="wb_copy_failures" for item in context.job.artifacts):
            asset=Asset(kind="file",name=failure_checkpoint.name,source="pod_wb_export",local_path=str(failure_checkpoint),sha256=hashlib.sha256(failure_checkpoint.read_bytes()).hexdigest(),mime_type="application/json",size_bytes=failure_checkpoint.stat().st_size,read_only=True)
            context.session.add(asset);context.session.flush();context.session.add(JobArtifact(job_id=context.job.id,asset_id=asset.id,role="wb_copy_failures"));context.session.commit()
        context.event(f"{len(copy_failures)} 个产品文案生成尝试 3 次后仍失败，已跳过并继续处理其余产品；详情见 copy-failures.json",level="warning")
    ready_products=[product for product in products if product.id not in copy_failures]
    if not ready_products:
        raise AppError(f"全部 {len(products)} 个产品文案生成失败，未生成 WB 发布包；详情见 copy-failures.json",code="wb_copy_all_failed")
    for index,product in enumerate(ready_products,1):
        context.refresh();copy_value=copies.get(product.id)
        selection=str(payload.get("variant_selection") or "all")
        variants=_selected_variants(product,selection,payload.get("variant_ids") or [])
        if not variants: raise AppError(f"{product.sku} 没有启用的 SKU",code="wb_variant_missing")
        if len(variants)>30: raise AppError(f"{product.sku} 包含 {len(variants)} 个启用 SKU，单个 WB 合并组最多 30 个",code="wb_group_too_large")
        cards=[]
        split_sizes=int(payload["subject_id"]) in {321}
        variant_groups=[[variant] for variant in variants] if split_sizes else [variants]
        for card_variants in variant_groups:
            variant=card_variants[0]
            vendor="-".join(value for value in (str((config.get("store") or {}).get("offer_prefix") or "").strip("-"),variant.sku) if value)
            if len(vendor)>40: raise AppError(f"WB 卖家 SKU 超过40字符：{vendor}",code="wb_vendor_code_long")
            length=variant.package_length_mm or product.package_length_mm;width=variant.package_width_mm or product.package_width_mm;height=variant.package_height_mm or product.package_height_mm;weight=variant.weight_grams or product.package_weight_g
            if not all(float(value or 0)>0 for value in (length,width,height,weight)): raise AppError(f"{variant.sku} 缺少包装尺寸或重量",code="wb_package_missing")
            dimensions={"length":_cm(length),"width":_cm(width),"height":_cm(height),"weightBrutto":round(float(weight)/1000,3)}
            sides=[dimensions["length"],dimensions["width"],dimensions["height"]]
            if max(sides)>115 or sum(sides)>200: raise AppError(f"{variant.sku} 包装尺寸超过 WB 限制：{sides[0]}×{sides[1]}×{sides[2]} cm",code="wb_package_oversize")
            images=_listing_images(product,variant,payload.get("image_rules") or [])
            if not images: raise AppError(f"{variant.sku} 没有可公开访问的 HTTPS 效果图",code="wb_images_invalid")
            generated={int(item.get("id") or 0):item.get("value") for item in copy_value.get("characteristics") or [] if isinstance(item,dict)}
            for item in _size_characteristics(int(payload["subject_id"]),variant): generated[item["id"]]=item["value"]
            factual_color=_variant_color_characteristic(variant)
            if factual_color:generated[factual_color["id"]]=factual_color["value"]
            validated=_validated_characteristics([{"id":identity,"value":value} for identity,value in generated.items()],live_schema)
            generated={int(item["id"]):item["value"] for item in validated}
            prices=[_variant_price(product,item,payload.get("pricing") or {}) for item in card_variants]
            errors=[value.get("error") for value in prices if value.get("error")]
            if errors: raise AppError(f"{variant.sku} 定价失败：{errors[0]}",code="wb_pricing_invalid")
            sizes=[{"techSize":str(item.size or item.specification or "ONE SIZE")[:30],"variant_id":item.id,"sku":item.sku,"price":prices[position]["list_price_rub"]} for position,item in enumerate(card_variants)]
            card_price=max(float(value["list_price_rub"]) for value in prices)
            cards.append({"vendorCode":vendor,"title":str(copy_value["title"])[:60],"description":str(copy_value["description"])[:5000],"brand":payload["brand"],"price":card_price,"dimensions":dimensions,"characteristics":[{"id":identity,"value":value} for identity,value in generated.items()],"images":images,"sizes":sizes,"product_id":product.id,"product_sku":product.sku})
        groups.append({"product_id":product.id,"sku":product.sku,"copy":copy_value,"variants":cards})
        context.set_progress("copy",index*100/max(1,len(ready_products)),f"已生成 {index}/{len(ready_products)} 个可发布产品文案；跳过 {len(copy_failures)} 个失败产品")
    context.complete_step("copy",details={"products":len(groups),"failed_products":len(copy_failures),"checkpoint":str(checkpoint),"failure_report":str(failure_checkpoint) if copy_failures else ""})
    # A WB merged link contains at most 30 cards.  Flatten cards across all
    # selected POD products, then chunk once so Excel and API share grouping.
    all_cards=sorted((card for product_group in groups for card in product_group["variants"]),key=_card_sku_sort_key)
    merged_groups=[{"group_index":index//30+1,"variants":all_cards[index:index+30]} for index in range(0,len(all_cards),30)]
    package=root/"publication-package.json";package.write_text(json.dumps({"version":2,"platform":"wb","payload":payload,"product_groups":groups,"groups":merged_groups,"summary":{"products":len(groups),"failed_products":len(copy_failures),"variants":sum(len(_selected_variants(product,str(payload.get('variant_selection') or 'all'),payload.get('variant_ids') or [])) for product in ready_products),"cards":len(all_cards),"links":len(merged_groups)}},ensure_ascii=False,indent=2),encoding="utf-8")
    if str(payload.get("mode") or "excel")=="api":
        _publish_api(context,config,root,groups,merged_groups)
        return
    context.start_step("excel");output=root/f"WB批量上传_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx";row_count=_write_excel(template_path,output,merged_groups,payload["subject_name"])
    for path,role,mime in ((output,"wb_excel","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),(package,"wb_package","application/json")):
        asset=Asset(kind="file",name=path.name,source="pod_wb_export",local_path=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest(),mime_type=mime,size_bytes=path.stat().st_size,read_only=True);context.session.add(asset);context.session.flush();context.session.add(JobArtifact(job_id=context.job.id,asset_id=asset.id,role=role))
    for ledger in context.session.scalars(select(WbPublicationItem).where(WbPublicationItem.job_id==context.job.id)).all():
        if ledger.status!="failed":ledger.status="exported_pending_confirmation"
    context.session.commit();context.complete_step("excel",details={"rows":row_count,"groups":len(merged_groups),"file":output.name})
