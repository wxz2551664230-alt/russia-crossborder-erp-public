from __future__ import annotations

import math
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class SeerfarImportError(ValueError):
    pass


@dataclass(slots=True)
class ImportedResearchProduct:
    source_file: str
    source_row: int
    source_sku: str
    title: str
    rank: int | None = None
    ozon_url: str = ""
    image_url: str = ""
    image_urls: list[str] = field(default_factory=list)
    brand: str = ""
    category: str = ""
    selling_price_rub: float | None = None
    sales: int | None = None
    rating: float | None = None
    weight_grams: float | None = None
    raw: dict[str, Any] = field(default_factory=dict)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _identifier(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        return str(int(value)) if value.is_integer() else format(value, ".15g")
    text = str(value).strip()
    return text.split(".", 1)[0] if re.fullmatch(r"\d+\.0+", text) else text


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    cleaned = re.sub(r"[^0-9,.\-]", "", str(value)).replace(",", "")
    try:
        return float(cleaned) if cleaned not in {"", "-", "."} else None
    except ValueError:
        return None


def _integer(value: Any) -> int | None:
    number = _number(value)
    return None if number is None else int(round(number))


def _weight_grams(value: Any) -> float | None:
    number = _number(value)
    if number is None:
        return None
    text = str(value).lower()
    return number * 1000 if any(unit in text for unit in ("kg", "кг", "千克")) else number


def _headers(values: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for index, value in enumerate(values, 1):
        name = _text(value) or f"未命名列{index}"
        seen[name] = seen.get(name, 0) + 1
        result.append(name if seen[name] == 1 else f"{name}#{seen[name]}")
    return result


def _image_urls(row: dict[str, Any]) -> list[str]:
    urls: list[str] = []
    for key, value in row.items():
        label = str(key).casefold()
        if not any(token in label for token in ("主图", "附图", "图片", "image", "photo")):
            continue
        for part in re.split(r"[\r\n,;|\s]+", _text(value)):
            if part.startswith(("http://", "https://")) and part not in urls:
                urls.append(part)
    return urls


def import_seerfar_workbook(path_value: str | Path) -> list[ImportedResearchProduct]:
    path = Path(path_value)
    if not path.is_file():
        raise SeerfarImportError(f"文件不存在：{path}")
    if path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise SeerfarImportError("仅支持 Seerfar 导出的 .xlsx 或 .xlsm 文件")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise SeerfarImportError(f"无法读取 Excel：{error}") from error
    try:
        sheet = workbook["Data"] if "Data" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = _headers(list(next(rows)))
        except StopIteration as error:
            raise SeerfarImportError("Excel 中没有数据") from error
        missing = sorted({"标题", "详情页地址", "SKU"}.difference(headers))
        if missing:
            raise SeerfarImportError("缺少必要字段：" + "、".join(missing))
        seen: set[str] = set()
        products: list[ImportedResearchProduct] = []
        for row_number, values in enumerate(rows, 2):
            row = dict(zip(headers, values, strict=False))
            sku = _identifier(row.get("SKU"))
            title = _text(row.get("标题"))
            if not sku and not title:
                continue
            if not sku:
                raise SeerfarImportError(f"第 {row_number} 行缺少 SKU")
            if sku in seen:
                continue
            seen.add(sku)
            images = _image_urls(row)
            products.append(ImportedResearchProduct(
                source_file=str(path.resolve()),
                source_row=row_number,
                source_sku=sku,
                title=title,
                rank=_integer(row.get("排名")),
                ozon_url=_text(row.get("详情页地址")),
                image_url=images[0] if images else "",
                image_urls=images,
                brand=_text(row.get("品牌")),
                category=_text(row.get("类目")),
                selling_price_rub=_number(row.get("售价")),
                sales=_integer(row.get("销量")),
                rating=_number(row.get("评分")),
                weight_grams=_weight_grams(row.get("重量")),
                raw={key: _text(value) for key, value in row.items()},
            ))
        return products
    finally:
        workbook.close()

