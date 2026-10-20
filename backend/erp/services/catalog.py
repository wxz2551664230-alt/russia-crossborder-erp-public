from __future__ import annotations

import hashlib
import json
import re
import sqlite3
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlparse

from openpyxl import load_workbook
from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session, selectinload

from erp.core.config import get_settings
from erp.core.errors import AppError, ConflictError, NotFoundError
from erp.core.jsonutil import dumps, loads
from erp.models import (
    Asset,
    CatalogDuplicateSuggestion,
    CatalogImportRun,
    CatalogProduct,
    CatalogProductAsset,
    CatalogVariant,
    Listing,
    ResearchProduct,
    SourcingCandidate,
    StoreListing,
)
from erp.schemas import CatalogProductCreate, CatalogProductUpdate
from erp.services.assets import register_local_asset
from erp.services.jobs import JobContext, job_handler


POD_LEGACY_SOURCE = "pod_product_library"
SELF_SELECTED_LEGACY_SOURCE = "self_selected_workbook"
FOLLOW_LEGACY_SOURCE = "ozon_follow_listing"


def product_query():
    return select(CatalogProduct).options(
        selectinload(CatalogProduct.variants),
        selectinload(CatalogProduct.asset_links).selectinload(CatalogProductAsset.asset),
    )


def get_catalog_product(session: Session, product_id: str) -> CatalogProduct:
    item = session.scalar(product_query().where(CatalogProduct.id == product_id))
    if not item:
        raise NotFoundError("产品不存在")
    return item


def create_catalog_product(session: Session, value: CatalogProductCreate) -> CatalogProduct:
    item = CatalogProduct(
        sku=value.sku.strip(),
        group_code=(value.group_code or value.sku).strip(),
        source_type=value.source_type,
        name=value.name.strip(),
        category=value.category.strip(),
        source_url=value.source_url.strip(),
        supplier=value.supplier.strip(),
        supplier_product_id=value.supplier_product_id.strip(),
        purchase_price_cny=value.purchase_price_cny,
        domestic_delivery_cny=value.domestic_delivery_cny,
        currency=value.currency.upper(),
        package_length_mm=value.package_length_mm,
        package_width_mm=value.package_width_mm,
        package_height_mm=value.package_height_mm,
        package_weight_g=value.package_weight_g,
        description_category_id=value.description_category_id,
        type_id=value.type_id,
        category_display=value.category_display.strip(),
        notes=value.notes,
        status=value.status,
        tags_json=dumps(value.tags),
        production_json=dumps(value.production),
        source_snapshot_json=dumps(value.source_snapshot),
    )
    session.add(item)
    session.flush()
    for variant in value.variants:
        session.add(CatalogVariant(
            product_id=item.id,
            sku=variant.sku.strip(),
            supplier_product_id=variant.supplier_product_id.strip(),
            supplier_sku_id=variant.supplier_sku_id.strip(),
            specification=variant.specification.strip(),
            color=variant.color.strip(),
            size=variant.size.strip(),
            purchase_price_cny=variant.purchase_price_cny,
            domestic_delivery_cny=variant.domestic_delivery_cny,
            package_length_mm=variant.package_length_mm,
            package_width_mm=variant.package_width_mm,
            package_height_mm=variant.package_height_mm,
            weight_grams=variant.weight_grams,
            barcode=variant.barcode.strip(),
            active=variant.active,
            properties_json=dumps(variant.properties),
        ))
    session.commit()
    refresh_duplicate_suggestions(session, item.id)
    session.commit()
    return get_catalog_product(session, item.id)


def update_catalog_product(session: Session, item: CatalogProduct, value: CatalogProductUpdate) -> CatalogProduct:
    values = value.model_dump(exclude_unset=True)
    json_fields = {"tags": "tags_json", "production": "production_json", "source_snapshot": "source_snapshot_json"}
    for input_key, model_key in json_fields.items():
        if input_key in values:
            setattr(item, model_key, dumps(values.pop(input_key) or ([] if input_key == "tags" else {})))
    for key, raw in values.items():
        setattr(item, key, raw.upper() if key == "currency" and raw else raw)
    session.commit()
    refresh_duplicate_suggestions(session, item.id)
    session.commit()
    return get_catalog_product(session, item.id)


def link_catalog_asset(
    session: Session,
    *,
    product: CatalogProduct,
    asset: Asset,
    role: str,
    position: int,
    variant_id: str | None = None,
    source_kind: str = "import",
    set_current: bool = True,
    metadata: dict[str, Any] | None = None,
) -> CatalogProductAsset:
    if variant_id and not any(variant.id == variant_id for variant in product.variants):
        raise ConflictError("该 SKU 不属于当前产品")
    if role == "main" and set_current:
        for current in product.asset_links:
            if current.is_current and current.role == "main" and current.variant_id == variant_id:
                current.is_current = False
    link = next((row for row in product.asset_links if row.asset_id == asset.id and row.role == role and row.variant_id == variant_id), None)
    if not link:
        link = CatalogProductAsset(product_id=product.id, asset_id=asset.id, role=role, variant_id=variant_id)
        session.add(link)
        product.asset_links.append(link)
    link.position = position
    link.is_current = set_current
    link.source_kind = source_kind
    link.metadata_json = dumps(metadata or {})
    return link


def _normalize_url(value: str) -> str:
    return str(value or "").strip().rstrip("/").casefold()


def refresh_duplicate_suggestions(session: Session, product_id: str) -> int:
    product = get_catalog_product(session, product_id)
    candidates = session.scalars(product_query().where(
        CatalogProduct.id != product.id,
        CatalogProduct.source_type != product.source_type,
    )).all()
    own_skus = {variant.sku.casefold() for variant in product.variants if variant.sku}
    own_hashes = {link.asset.sha256 for link in product.asset_links if link.is_current and link.asset.sha256}
    created = 0
    for candidate in candidates:
        reasons: list[str] = []
        score = 0.0
        if product.source_url and _normalize_url(product.source_url) == _normalize_url(candidate.source_url):
            reasons.append("货源链接相同"); score += 0.55
        if product.supplier_product_id and product.supplier_product_id == candidate.supplier_product_id:
            reasons.append("供应商商品 ID 相同"); score += 0.65
        shared_skus = own_skus & {variant.sku.casefold() for variant in candidate.variants if variant.sku}
        if shared_skus:
            reasons.append(f"SKU 相同：{', '.join(sorted(shared_skus)[:3])}"); score += 0.35
        shared_hashes = own_hashes & {link.asset.sha256 for link in candidate.asset_links if link.is_current and link.asset.sha256}
        if shared_hashes:
            reasons.append(f"图片哈希相同：{len(shared_hashes)} 张"); score += min(0.5, len(shared_hashes) * 0.15)
        if not reasons:
            continue
        first_id, second_id = sorted((product.id, candidate.id))
        suggestion = session.scalar(select(CatalogDuplicateSuggestion).where(
            CatalogDuplicateSuggestion.product_id == first_id,
            CatalogDuplicateSuggestion.candidate_product_id == second_id,
        ))
        if not suggestion:
            suggestion = CatalogDuplicateSuggestion(product_id=first_id, candidate_product_id=second_id)
            session.add(suggestion); created += 1
        if suggestion.status != "merged":
            suggestion.score = min(1.0, score)
            suggestion.reasons_json = dumps(reasons)
    return created


def inspect_pod_database(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise NotFoundError(f"POD 历史数据库不存在：{path}")
    connection = sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)
    try:
        required = {"products", "product_versions", "product_variants", "product_oss_assets"}
        tables = {row[0] for row in connection.execute("select name from sqlite_master where type='table'")}
        missing = sorted(required - tables)
        if missing:
            raise AppError("POD 历史数据库结构不完整", code="pod_database_invalid", details=missing)
        counts = {table: int(connection.execute(f"select count(*) from {table}").fetchone()[0]) for table in sorted(required)}
        broken_versions = int(connection.execute(
            "select count(*) from products p left join product_versions v on v.id=p.current_version_id where p.current_version_id is not null and v.id is null"
        ).fetchone()[0])
        blank_skus = int(connection.execute("select count(*) from products where trim(sku)='' ").fetchone()[0])
        return {"counts": counts, "broken_current_versions": broken_versions, "blank_skus": blank_skus, "source_database": str(path)}
    finally:
        connection.close()


def _asset_role(asset_type: str, position: int) -> str:
    if asset_type.casefold() == "source":
        return "artwork"
    return "main" if position == 0 else "mockup"


def _pod_supplier_from_sku(sku: str) -> str:
    normalized = str(sku or "").strip().upper()
    if normalized.startswith("S2B"):
        return "S2B"
    if normalized.startswith("HC"):
        return "HICUSTOM"
    return ""


def import_pod_database(context: JobContext, source_path: Path, run: CatalogImportRun) -> None:
    report = inspect_pod_database(source_path)
    run.counts_json = dumps(report["counts"])
    run.report_json = dumps(report)
    if run.dry_run:
        run.status = "completed"
        context.session.commit()
        return
    connection = sqlite3.connect(f"file:{source_path.as_posix()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    inserted = updated = variant_count = asset_count = 0
    try:
        categories = {str(row["id"]): str(row["name"]) for row in connection.execute("select id,name from product_categories")}
        versions = {str(row["id"]): dict(row) for row in connection.execute("select * from product_versions")}
        products = list(connection.execute("select * from products order by created_at,id"))
        context.start_step("products")
        for index, raw in enumerate(products, 1):
            context.refresh()
            legacy_id = str(raw["id"])
            item = context.session.scalar(select(CatalogProduct).where(
                CatalogProduct.legacy_source == POD_LEGACY_SOURCE,
                CatalogProduct.legacy_id == legacy_id,
            ))
            if item:
                updated += 1
            else:
                item = CatalogProduct(legacy_source=POD_LEGACY_SOURCE, legacy_id=legacy_id, source_type="pod", sku=str(raw["sku"] or legacy_id))
                context.session.add(item); context.session.flush(); inserted += 1
            version = versions.get(str(raw["current_version_id"] or ""), {})
            item.group_code = str(raw["family_no"] or raw["sku"] or legacy_id)
            item.name = str(raw["name"] or raw["sku"] or "POD 产品")
            item.category = categories.get(str(raw["category_id"]), str(raw["category_id"] or ""))
            item.supplier = _pod_supplier_from_sku(item.sku) or str(item.supplier or "HICUSTOM/S2B")
            item.supplier_product_id = str(version.get("product_type_id") or "")
            item.notes = str(raw["notes"] or "")
            item.status = str(raw["status"] or "ready")
            item.tags_json = str(raw["tags_json"] or "[]")
            item.production_json = dumps({
                "version_id": version.get("id"), "version_no": version.get("version_no"),
                "job_id": version.get("job_id"), "job_status": version.get("job_status"),
                "product_type_id": version.get("product_type_id"), "default_color_id": version.get("default_color_id"),
                "default_view_id": version.get("default_view_id"), "product_code": version.get("product_code"),
                "product_code_status": version.get("product_code_status"), "output_folder": version.get("output_folder"),
                "recipe": loads(str(version.get("recipe_json") or "{}"), {}),
                "assets": loads(str(version.get("assets_json") or "{}"), {}),
            })
            item.source_snapshot_json = dumps({"legacy_product": dict(raw), "source_database": str(source_path)})
            existing_variants = {
                str(row.legacy_id): row for row in context.session.scalars(select(CatalogVariant).where(
                    CatalogVariant.product_id == item.id,
                    CatalogVariant.legacy_source == POD_LEGACY_SOURCE,
                )).all()
            }
            rows = connection.execute("select * from product_variants where product_id=? order by variant_no,id", (legacy_id,)).fetchall()
            for variant in rows:
                variant_id = str(variant["id"])
                target_variant = existing_variants.pop(variant_id, None)
                if not target_variant:
                    target_variant = CatalogVariant(product_id=item.id, legacy_source=POD_LEGACY_SOURCE, legacy_id=variant_id)
                    context.session.add(target_variant)
                target_variant.sku = str(variant["sku"])
                target_variant.supplier_product_id = str(version.get("product_type_id") or "")
                target_variant.supplier_sku_id = str(variant["stock_item_id"] or "")
                target_variant.specification = str(variant["stock_code"] or "")
                target_variant.color = str(variant["color_name"] or "")
                target_variant.size = str(variant["size_name"] or "")
                target_variant.package_length_mm = variant["package_length"]
                target_variant.package_width_mm = variant["package_width"]
                target_variant.package_height_mm = variant["package_height"]
                target_variant.weight_grams = variant["weight"]
                target_variant.active = bool(variant["active"])
                target_variant.properties_json = dumps({
                        "variant_no": variant["variant_no"], "color_id": variant["color_id"], "size_id": variant["size_id"],
                        "stock_code": variant["stock_code"], "package_volume": variant["package_volume"],
                        "size_attributes": loads(variant["size_attributes_json"], []), "prices": loads(variant["prices_json"], {}),
                        "image_group_id": variant["image_group_id"], "selected": bool(variant["selected"]),
                    })
                variant_count += 1
            for obsolete in existing_variants.values():
                obsolete.active = False
            if version:
                current_links = context.session.scalars(select(CatalogProductAsset).where(
                    CatalogProductAsset.product_id == item.id,
                    CatalogProductAsset.source_kind == POD_LEGACY_SOURCE,
                )).all()
                for link in current_links:
                    context.session.delete(link)
                media = connection.execute("select * from product_oss_assets where version_id=?", (version["id"],)).fetchall()
                archived_assets = loads(str(version.get("assets_json") or "{}"), {}).get("oss_assets") or []
                asset_order = {str(row.get("asset_id") or ""): index for index, row in enumerate(archived_assets) if isinstance(row, dict)}
                media = sorted(media, key=lambda row: (asset_order.get(str(row["asset_id"]), 1_000_000), str(row["asset_id"]), str(row["id"])))
                rendering_index = 0
                for position, media_row in enumerate(media):
                    media_id = str(media_row["id"])
                    asset = context.session.scalar(select(Asset).where(Asset.legacy_source == POD_LEGACY_SOURCE, Asset.legacy_id == media_id))
                    if not asset:
                        asset = Asset(legacy_source=POD_LEGACY_SOURCE, legacy_id=media_id)
                        context.session.add(asset); asset_count += 1
                    asset.kind = str(media_row["asset_type"] or "image")
                    asset.name = Path(str(media_row["local_path"] or media_row["oss_key"] or media_id)).name
                    asset.source = "pod"
                    raw_local_path = Path(str(media_row["local_path"] or ""))
                    if raw_local_path and not raw_local_path.is_absolute() and version.get("output_folder"):
                        raw_local_path = Path(str(version["output_folder"])) / raw_local_path
                    asset.local_path = str(raw_local_path) if str(raw_local_path) != "." else ""
                    asset.oss_bucket = str(media_row["oss_bucket"] or "")
                    asset.oss_key = str(media_row["oss_key"] or "")
                    asset.oss_url = str(media_row["oss_url"] or media_row["oss_uri"] or "")
                    asset.sha256 = str(media_row["sha256"] or "")
                    asset.mime_type = str(media_row["content_type"] or "application/octet-stream")
                    asset.size_bytes = int(media_row["size"] or 0)
                    asset.read_only = True
                    asset.mirror_status = "mirrored" if asset.oss_url else "not_requested"
                    asset.metadata_json = dumps({"color_id": media_row["color_id"], "color_name": media_row["color_name"], "view_id": media_row["view_id"], "etag": media_row["etag"], "crc64": media_row["crc64"]})
                    context.session.flush()
                    asset_type = str(media_row["asset_type"] or "")
                    role_position = rendering_index if asset_type == "rendering" else position
                    link_catalog_asset(context.session, product=item, asset=asset, role=_asset_role(asset_type, role_position), position=role_position, source_kind=POD_LEGACY_SOURCE, metadata={"color_id": media_row["color_id"], "color_name": media_row["color_name"], "view_id": media_row["view_id"]})
                    if asset_type == "rendering":
                        rendering_index += 1
            if index % 50 == 0:
                context.session.commit()
                context.set_progress("products", index * 100 / max(1, len(products)), f"已索引 {index}/{len(products)} 个 POD 产品")
        context.session.commit()
        context.complete_step("products", details={"inserted": inserted, "updated": updated, "variants": variant_count, "assets": asset_count})
        run.status = "completed"
        run.counts_json = dumps({"inserted": inserted, "updated": updated, "variants": variant_count, "assets": asset_count})
        context.session.commit()
    finally:
        connection.close()


@job_handler("catalog_pod_import")
def catalog_pod_import_job(context: JobContext) -> None:
    if get_settings().multiuser_enabled:
        raise ConflictError('POD 历史导入只能在发起人的电脑读取旧数据库')
    run = context.session.get(CatalogImportRun, str(context.payload.get("import_run_id") or ""))
    if not run:
        raise NotFoundError("导入记录不存在")
    run.status = "running"; context.session.commit()
    try:
        import_pod_database(context, Path(run.source_path), run)
    except Exception as error:
        run.status = "failed"; run.errors_json = dumps([str(error)]); context.session.commit(); raise


HEADER_ALIASES = {
    # This is intentionally the ERP whitelist from
    # product-library-import-template.xlsx. All other columns in the complete
    # user workbook are ignored.
    "source_url": ("来源链接",), "name": ("中文名称",), "sku": ("SKU",),
    "specification": ("规格型号",), "color": ("颜色",), "size": ("尺码",),
    "material": ("材质",), "purchase_price_cny": ("拿货价",),
    "item_length": ("长",), "item_width": ("宽",), "item_height": ("高",),
    "package_length_mm": ("长(cm)",), "package_width_mm": ("宽(cm)",),
    "package_height_mm": ("高（cm)", "高(cm)"), "package_weight_g": ("重量(g)",),
    "main_oss_url": ("主图链接",), "additional_oss_urls": ("附图链接",),
    "title_ru": ("产品名称", "商品名称"), "notes": ("备注",),
}
REQUIRED_HEADERS = {"source_url", "name", "sku", "purchase_price_cny", "package_length_mm", "package_width_mm", "package_height_mm", "package_weight_g", "main_oss_url", "additional_oss_urls", "title_ru"}
IMAGE_ROLES = {"main", "additional", "sku", "artwork", "mockup"}


def _split_values(value: Any) -> list[str]:
    return [part.strip() for part in re.split(r"[\r\n,;；]+", str(value or "")) if part.strip()]


def _http_url(value: str) -> bool:
    parsed = urlparse(value)
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def parse_self_selected_workbook(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    normalized = {alias.casefold(): key for key, aliases in HEADER_ALIASES.items() for alias in aliases}
    sheet = None; header_row = 0; column_map: dict[int, str] = {}
    # The complete user template has grouped headings on row 1 and actual
    # field names on row 2. Users may remove row 1, so search by field names
    # rather than relying on a sheet name, row number, or column order.
    for candidate in workbook.worksheets:
        for row_index, values in enumerate(candidate.iter_rows(min_row=1, max_row=min(10, candidate.max_row), values_only=True), 1):
            mapped = {index: normalized[str(value or "").strip().casefold()] for index, value in enumerate(values) if str(value or "").strip().casefold() in normalized}
            if len(mapped) > len(column_map): sheet, header_row, column_map = candidate, row_index, mapped
    if sheet is None:
        workbook.close(); raise AppError("没有找到产品字段表头", code="catalog_import_headers_missing", details=sorted(REQUIRED_HEADERS))
    missing = sorted(REQUIRED_HEADERS - set(column_map.values()))
    if missing:
        workbook.close()
        raise AppError("导入表缺少必填列", code="catalog_import_headers_missing", details=missing)
    records: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    seen_skus: dict[str, int] = {}
    for row_index, raw_row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        record = {key: raw_row[index] if index < len(raw_row) else None for index, key in column_map.items()}
        if not any(str(value or "").strip() for value in record.values()):
            continue
        record["group_code"] = str(record.get("sku") or "").strip()
        record["notes"] = str(record.get("notes") or "").strip()
        for key in ("group_code", "sku", "name", "source_url", "title_ru"):
            record[key] = str(record.get(key) or "").strip()
            if not record[key]: errors.append({"row": row_index, "field": key, "message": "不能为空"})
        for key in ("purchase_price_cny", "package_length_mm", "package_width_mm", "package_height_mm", "package_weight_g"):
            try:
                record[key] = float(record.get(key))
                if record[key] <= 0: raise ValueError
            except (TypeError, ValueError): errors.append({"row": row_index, "field": key, "message": "必须是大于 0 的数字"})
        for key in ("package_length_mm", "package_width_mm", "package_height_mm"):
            if isinstance(record.get(key), (int, float)): record[key] *= 10
        for key in ("item_length", "item_width", "item_height"):
            try: record[key] = float(record[key]) if str(record.get(key) or "").strip() else None
            except (TypeError, ValueError): errors.append({"row": row_index, "field": key, "message": "必须是数字"})
        if record["source_url"] and not _http_url(record["source_url"]): errors.append({"row": row_index, "field": "source_url", "message": "货源链接格式不正确"})
        record["main_oss_url"] = str(record.get("main_oss_url") or "").strip()
        if record["main_oss_url"] and not _http_url(record["main_oss_url"]): errors.append({"row": row_index, "field": "main_oss_url", "message": "主图 OSS 链接格式不正确"})
        record["additional_oss_urls"] = _split_values(record.get("additional_oss_urls"))
        if not record["additional_oss_urls"]: errors.append({"row": row_index, "field": "additional_oss_urls", "message": "至少需要一张附图"})
        for url in record["additional_oss_urls"]:
            if not _http_url(url): errors.append({"row": row_index, "field": "additional_oss_urls", "message": f"链接格式不正确：{url}"})
        sku_key = record["sku"].casefold()
        if sku_key in seen_skus: errors.append({"row": row_index, "field": "sku", "message": f"与第 {seen_skus[sku_key]} 行 SKU 重复"})
        else: seen_skus[sku_key] = row_index
        record["row"] = row_index
        records.append(record)
    workbook.close()
    return records, errors


def parse_self_selected_images(path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "图片清单" not in workbook.sheetnames:
        workbook.close()
        return [], []
    sheet = workbook["图片清单"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows, ())]
    aliases = {"产品组编号": "group_code", "SKU（可留空）": "sku", "SKU": "sku", "图片相对路径": "path", "图片类型": "role", "排序": "position"}
    columns = {index: aliases[value] for index, value in enumerate(headers) if value in aliases}
    errors: list[dict[str, Any]] = []
    images: list[dict[str, Any]] = []
    root = path.parent.resolve()
    for row_index, raw in enumerate(rows, 2):
        item = {key: raw[index] if index < len(raw) else None for index, key in columns.items()}
        if not any(str(value or "").strip() for value in item.values()): continue
        group_code = str(item.get("group_code") or "").strip()
        relative = str(item.get("path") or "").strip().replace("\\", "/")
        role = str(item.get("role") or "additional").strip().casefold()
        try: position = max(0, int(item.get("position") or 0))
        except (TypeError, ValueError): position = 0; errors.append({"row": row_index, "field": "position", "message": "排序必须是整数"})
        candidate = (root / relative).resolve()
        if not group_code: errors.append({"row": row_index, "field": "group_code", "message": "不能为空"})
        if not relative or root not in candidate.parents: errors.append({"row": row_index, "field": "path", "message": "图片路径无效或超出资料包"})
        elif not candidate.is_file(): errors.append({"row": row_index, "field": "path", "message": f"找不到图片：{relative}"})
        if role not in IMAGE_ROLES: errors.append({"row": row_index, "field": "role", "message": f"不支持的图片类型：{role}"})
        images.append({"row": row_index, "group_code": group_code, "sku": str(item.get("sku") or "").strip(), "path": candidate, "relative_path": relative, "role": role, "position": position})
    workbook.close()
    return images, errors


def _url_asset(session: Session, url: str, *, legacy_id: str, source: str, name: str,
               sha256: str = '', size_bytes: int = 0) -> Asset:
    asset = session.scalar(select(Asset).where(Asset.legacy_source == SELF_SELECTED_LEGACY_SOURCE, Asset.legacy_id == legacy_id))
    if not asset:
        asset = Asset(legacy_source=SELF_SELECTED_LEGACY_SOURCE, legacy_id=legacy_id)
        session.add(asset)
    asset.kind = "image"; asset.name = name; asset.source = source; asset.oss_url = url
    asset.local_path = ""; asset.read_only = True; asset.mirror_status = "mirrored"; asset.mime_type = "image/*"
    if sha256: asset.sha256 = sha256
    if size_bytes: asset.size_bytes = size_bytes
    return asset


def apply_self_selected_records(session: Session, run: CatalogImportRun,
                                records: list[dict[str, Any]], errors: list[dict[str, Any]],
                                counts: dict[str, int], local_images: list[dict[str, Any]],
                                *, on_refresh=None, on_progress=None) -> dict[str, int]:
    """Consolidate already validated PC rows using the original catalog mapping."""
    category_snapshot = loads(run.report_json, {}).get("ozon_category") or {}
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records: groups[str(record["group_code"])].append(record)
    inserted = updated = 0
    for index, (group_code, rows) in enumerate(groups.items(), 1):
        if on_refresh: on_refresh()
        first = rows[0]; legacy_id = group_code.casefold()
        product = session.scalar(select(CatalogProduct).where(CatalogProduct.legacy_source == SELF_SELECTED_LEGACY_SOURCE, CatalogProduct.legacy_id == legacy_id))
        if not product:
            product = session.scalar(select(CatalogProduct).where(CatalogProduct.source_type == "self_selected", CatalogProduct.group_code == group_code))
        if not product:
            product = CatalogProduct(legacy_source=SELF_SELECTED_LEGACY_SOURCE, legacy_id=legacy_id, source_type="self_selected", sku=str(first["sku"])); session.add(product); session.flush(); inserted += 1
        else:
            updated += 1
            if product.legacy_source == SELF_SELECTED_LEGACY_SOURCE: product.legacy_id = legacy_id
        product.group_code=group_code; product.name=str(first["name"]); product.source_url=str(first["source_url"]); product.supplier="1688"
        product.supplier_product_id=str(first.get("supplier_product_id") or ""); product.purchase_price_cny=float(first["purchase_price_cny"]); product.domestic_delivery_cny=float(first.get("domestic_delivery_cny") or 0) or None
        product.currency=str(first.get("currency") or "CNY").upper(); product.package_length_mm=float(first["package_length_mm"]); product.package_width_mm=float(first["package_width_mm"]); product.package_height_mm=float(first["package_height_mm"]); product.package_weight_g=float(first["package_weight_g"])
        product.description_category_id=int(category_snapshot["description_category_id"]); product.type_id=int(category_snapshot["type_id"]); product.category_display=str(category_snapshot["display"]); product.category=str(category_snapshot["display"])
        product.notes=str(first.get("notes") or ""); product.status="ready"; product.source_snapshot_json=dumps({"workbook": str(run.source_path), "rows": [row["row"] for row in rows]})
        product.production_json=dumps({"workflow":"self_listing_batch","copywriting":{"title_ru":str(first.get("title_ru") or "")},"images":{"main":str(first.get("main_oss_url") or ""),"additional":first.get("additional_oss_urls") or []}})
        existing_variants = {variant.sku.casefold(): variant for variant in product.variants}
        for row in rows:
            sku = str(row["sku"])
            variant = existing_variants.pop(sku.casefold(), None)
            if not variant:
                variant = CatalogVariant(product=product, sku=sku, legacy_source=SELF_SELECTED_LEGACY_SOURCE, legacy_id=f"{legacy_id}:{sku.casefold()}")
                session.add(variant)
            variant.sku=sku;variant.supplier_product_id="";variant.supplier_sku_id="";variant.specification=str(row.get("specification") or "");variant.color=str(row.get("color") or "");variant.size=str(row.get("size") or "");variant.purchase_price_cny=float(row["purchase_price_cny"]);variant.domestic_delivery_cny=None;variant.package_length_mm=float(row["package_length_mm"]);variant.package_width_mm=float(row["package_width_mm"]);variant.package_height_mm=float(row["package_height_mm"]);variant.weight_grams=float(row["package_weight_g"]);variant.barcode="";variant.active=True;variant.properties_json=dumps({"material":str(row.get("material") or ""),"item_length":row.get("item_length"),"item_width":row.get("item_width"),"item_height":row.get("item_height"),"copywriting":{"title_ru":str(row.get("title_ru") or "")}})
        for obsolete in existing_variants.values(): obsolete.active=False
        for link in list(product.asset_links):
            if link.source_kind == SELF_SELECTED_LEGACY_SOURCE: session.delete(link)
        session.flush()
        variant_ids = {variant.sku: variant.id for variant in session.scalars(select(CatalogVariant).where(CatalogVariant.product_id == product.id)).all()}
        for row in rows:
            original_urls = ([("main", str(row.get("original_main_oss_url") or ""))] if row.get("original_main_oss_url") else []) + [("additional", url) for url in row.get("original_additional_oss_urls") or []]
            for position, (role, url) in enumerate(original_urls):
                asset = _url_asset(session,url,legacy_id=f"{group_code.casefold()}:original:{row['sku'].casefold()}:{role}:{position}:{url}",source="self_selected_original",name=f"{group_code}-original-{role}-{position}.jpg")
                session.flush(); link_catalog_asset(session,product=product,asset=asset,role=role,position=position if role == "additional" else 0,variant_id=variant_ids.get(row["sku"]),source_kind=SELF_SELECTED_LEGACY_SOURCE,set_current=False,metadata={"asset_version":"original","sku":row["sku"]})
            urls = ([("main", str(row["main_oss_url"]))] if row["main_oss_url"] else []) + [("additional", url) for url in row["additional_oss_urls"]]
            for position, (role, url) in enumerate(urls):
                asset = _url_asset(session,url,legacy_id=f"{group_code.casefold()}:{role}:{position}:{url}",source="self_selected",name=f"{group_code}-{role}-{position}.jpg")
                session.flush(); link_catalog_asset(session,product=product,asset=asset,role=role,position=position if role == "additional" else 0,variant_id=variant_ids.get(row["sku"]),source_kind=SELF_SELECTED_LEGACY_SOURCE,metadata={"asset_version":"watermarked" if row.get("watermarked_main_oss_url") else "legacy","sku":row["sku"]})
        for image in (item for item in local_images if item["group_code"] == group_code):
            if image.get('oss_url'):
                image_hash = image['sha256']
                asset = _url_asset(session, image['oss_url'],
                                   legacy_id=f"{group_code.casefold()}:local:{image_hash}",
                                   source='self_selected_bundle', name=Path(image['relative_path']).name,
                                   sha256=image_hash, size_bytes=image['size_bytes'])
            else:
                image_hash = hashlib.sha256(image["path"].read_bytes()).hexdigest()
                asset = register_local_asset(session, image["path"], kind="image", source="self_selected_bundle", read_only=True, calculate_hash=True, metadata={"relative_path": image["relative_path"], "import_run_id": run.id}, legacy_source=SELF_SELECTED_LEGACY_SOURCE, legacy_id=f"{group_code.casefold()}:local:{image_hash}")
            session.flush()
            link_catalog_asset(session, product=product, asset=asset, role=image["role"], position=image["position"], variant_id=variant_ids.get(image["sku"]) if image["sku"] else None, source_kind=SELF_SELECTED_LEGACY_SOURCE, metadata={"relative_path": image["relative_path"]})
        if on_progress and index % 25 == 0: on_progress(index, len(groups))
    session.flush()
    for product in session.scalars(select(CatalogProduct).where(CatalogProduct.source_type == "self_selected", CatalogProduct.group_code.in_(list(groups)))).all(): refresh_duplicate_suggestions(session, product.id)
    run.status="completed_with_errors" if errors else "completed"
    run.counts_json=dumps({**counts,"inserted":inserted,"updated":updated})
    return {"inserted":inserted,"updated":updated,"rejected_groups":counts.get("rejected_groups",0)}


@job_handler("catalog_self_selected_import")
def catalog_self_selected_import_job(context: JobContext) -> None:
    run = context.session.get(CatalogImportRun, str(context.payload.get("import_run_id") or ""))
    if not run: raise NotFoundError("导入记录不存在")
    run.status = "running"; context.session.commit()
    try:
        context.start_step("validate")
        workbook_path = Path(run.source_path)
        records, errors = parse_self_selected_workbook(workbook_path)
        local_images, image_errors = parse_self_selected_images(workbook_path)
        errors.extend(image_errors)
        product_groups = {str(row.get("group_code")) for row in records}
        known_skus = {str(row.get("sku")) for row in records}
        for image in local_images:
            if image["group_code"] not in product_groups:
                errors.append({"row": image["row"], "field": "group_code", "message": "图片对应的产品组不在产品导入表中"})
            if image["sku"] and image["sku"] not in known_skus:
                errors.append({"row": image["row"], "field": "sku", "message": "图片对应的 SKU 不在产品导入表中"})
        counts = {"rows": len(records), "product_groups": len(product_groups), "local_images": len(local_images), "errors": len(errors)}
        run.counts_json = dumps(counts); run.errors_json = dumps(errors)
        context.session.commit(); context.complete_step("validate", details=counts)
        if run.dry_run:
            run.status = "completed" if not errors else "validation_failed"
            context.session.commit(); return
        if errors:
            invalid_rows = {int(error["row"]) for error in errors if isinstance(error, dict) and error.get("row")}
            invalid_groups = {str(row["group_code"]) for row in records if row["row"] in invalid_rows}
            invalid_groups.update(str(image["group_code"]) for image in local_images if image["row"] in invalid_rows)
            records = [row for row in records if str(row["group_code"]) not in invalid_groups]
            local_images = [image for image in local_images if str(image["group_code"]) not in invalid_groups]
            counts["rejected_groups"] = len(invalid_groups)
            counts["accepted_groups"] = len({str(row["group_code"]) for row in records})
            if not records:
                run.status = "validation_failed"; run.counts_json = dumps(counts); context.session.commit(); return
        context.start_step("import")
        def progress(index: int, total: int) -> None:
            context.session.commit()
            context.set_progress("import", index*100/max(1,total), f"已导入 {index}/{total} 个产品组")
        summary = apply_self_selected_records(context.session, run, records, errors, counts, local_images,
                                              on_refresh=context.refresh, on_progress=progress)
        context.session.commit()
        context.complete_step("import", details=summary)
    except Exception as error:
        run.status="failed"; run.errors_json=dumps([*loads(run.errors_json,[]),str(error)]); context.session.commit(); raise


def ensure_follow_catalog_product(session: Session, listing: Listing) -> CatalogProduct:
    existing = session.scalar(select(CatalogProduct).where(CatalogProduct.legacy_source == FOLLOW_LEGACY_SOURCE, CatalogProduct.legacy_id == listing.id))
    if existing:
        return existing
    source = listing.research_product
    if not source:
        raise AppError("跟卖商品缺少原始选品资料", code="follow_catalog_source_missing")
    selected = next((row for row in source.sourcing_candidates if row.selected), None)
    package = loads(listing.package_json, {})
    # Follow products currently have no reliable source for package dimensions.
    # Operational defaults used for Ozon submission/pricing must not be stored in
    # the catalog as product facts.
    dimensions_are_facts = package.get("dimensions_source") in {"measured", "supplier", "manual", "imported_fact"}
    primary_store_id = str(loads(listing.image_policy_json, {}).get("catalog_primary_store_id") or "")
    child = next((row for row in listing.store_listings if row.store_id == primary_store_id and loads(row.assets_json, {}).get("primary_image") and loads(row.assets_json, {}).get("images")), None)
    if child is None:
        child = next((row for row in listing.store_listings if loads(row.assets_json, {}).get("primary_image") and loads(row.assets_json, {}).get("images")), None)
    if child is None:
        child = next((row for row in listing.store_listings if loads(row.assets_json, {}).get("primary_image")), None)
    image_state = loads(child.assets_json, {}) if child else {}
    missing = []
    if not selected: missing.append("已确认货源")
    if not image_state.get("primary_image"): missing.append("自主生成主图")
    if not image_state.get("images"): missing.append("OZON 附图")
    if missing: raise AppError("跟卖商品尚未达到入库条件",code="follow_catalog_incomplete",details=missing)
    erp_sku = f"FOL-{listing.base_offer_id}"
    source_images = loads(source.image_urls_json, [])
    ozon_facts = {"research_product_id":source.id,"source_sku":source.source_sku,"ozon_url":source.ozon_url,"original_title":source.title,"original_category":source.category,"original_brand":source.brand,"original_image_url":source.image_url,"original_image_urls":source_images,"selling_price_rub":source.selling_price_rub,"sales":source.sales,"rating":source.rating,"weight_grams":source.weight_grams,"raw":loads(source.raw_json,{})}
    production = {"kind":"ozon_follow","fact_source":"ozon_front","ozon_product_sku":source.source_sku,"ozon_url":source.ozon_url,"reference_images":source_images or ([source.image_url] if source.image_url else []),"selected_supplier_product_id":selected.external_product_id,"selected_supplier_sku_id":selected.external_sku_id,"selected_supplier_specification":selected.sku_name,"catalog_primary_store_id":child.store_id if child else ""}
    product = CatalogProduct(sku=erp_sku,group_code=erp_sku,source_type="follow",name=source.title,category=source.category,source_url=selected.source_url,supplier="1688",supplier_product_id=selected.external_product_id,purchase_price_cny=selected.unit_price_cny,domestic_delivery_cny=selected.domestic_delivery_cny,currency="CNY",package_length_mm=float(package["depth_mm"]) if dimensions_are_facts and package.get("depth_mm") else None,package_width_mm=float(package["width_mm"]) if dimensions_are_facts and package.get("width_mm") else None,package_height_mm=float(package["height_mm"]) if dimensions_are_facts and package.get("height_mm") else None,package_weight_g=float(package["weight_g"]) if package.get("weight_g") else None,description_category_id=source.description_category_id,type_id=source.type_id,category_display=source.category_display,status="ready",legacy_source=FOLLOW_LEGACY_SOURCE,legacy_id=listing.id,production_json=dumps(production),source_snapshot_json=dumps(ozon_facts),metadata_json=dumps({"ozon_base_offer_id":listing.base_offer_id,"erp_sku_rule":"FOL-{base_offer_id}","package_dimensions_status":"unavailable" if not dimensions_are_facts else "fact"}))
    session.add(product); session.flush()
    session.add(CatalogVariant(product_id=product.id,sku=erp_sku,supplier_product_id=selected.external_product_id,supplier_sku_id=selected.external_sku_id,specification=selected.sku_name,purchase_price_cny=selected.unit_price_cny,domestic_delivery_cny=selected.domestic_delivery_cny,weight_grams=float(package["weight_g"]) if package.get("weight_g") else None,properties_json=selected.raw_json))
    urls = [("main", image_state["primary_image"]), *(("additional", url) for url in image_state.get("images") or [])]
    for position, (role, url) in enumerate(urls):
        asset = Asset(kind="image",name=f"{listing.base_offer_id}-{role}-{position}.jpg",source="follow",oss_url=str(url),read_only=True,mirror_status="mirrored",mime_type="image/*",legacy_source=FOLLOW_LEGACY_SOURCE,legacy_id=f"{listing.id}:{role}:{position}")
        session.add(asset); session.flush(); link_catalog_asset(session,product=product,asset=asset,role=role,position=position if role == "additional" else 0,source_kind=FOLLOW_LEGACY_SOURCE)
    listing.catalog_product_id=product.id
    session.commit(); refresh_duplicate_suggestions(session,product.id); session.commit()
    return product
